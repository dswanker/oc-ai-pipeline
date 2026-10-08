"""standards_match.py: customer standard form matching (docs/OC_STANDARD_FORM_MATCHING_PLAN.md, agreed rules 2026-10-09).

Each protocol form is matched to a customer standard form by CDASH domain. On a match the standard form is used
EXACTLY (every field, choice list, constraint, required flag, relevance, calculation, hint, appearance, item group
and setting); the protocol analysis keeps the visit placement and form-level scheduling. A field is added only when
the protocol clearly and directly specifies that data point and the standard lacks it (verified verbatim quote).
Nothing is ever removed from a standard form.

Sources, in priority order: uploaded XLSForm (single .xlsx or ZIP) > referenced OC studies (in listed order) >
uploaded ODM XML. Every file is either used or named as not usable.

  load_sources(files, referenced)   -> {"forms": [standard form], "files": [...], "fingerprint": "..."}
  apply(spec, sources)              -> spec with matched forms replaced (never mutates its input; unchanged on error)
  build_add_request / apply_additions   rule 3: fields the protocol specifies and the standard lacks
  proposals(spec) / apply_proposal / reject_proposal   rule 7/8: logic proposed on standard forms, applied on Approve

Kill switch: STANDARDS_MATCHING=0. Study-agnostic: no customer-, protocol- or form-specific code.
"""
import copy, difflib, hashlib, io, json, os, re, zipfile

VERSION = 1
STATUS = "CUSTOMER_STANDARD"
SRC_XLSFORM, SRC_ODM = "uploaded XLSForm", "uploaded ODM"
ADDED = "Added from protocol"

# XLSForm survey columns the form builder always writes (skills/edc-builder/scripts/build_xlsforms.py SURVEY_COLS).
SURVEY_COLS = ["type", "name", "label", "bind::oc:itemgroup", "hint", "appearance", "bind::oc:briefdescription",
               "bind::oc:description", "relevant", "required", "required_message", "constraint",
               "constraint_message", "default", "calculation", "trigger", "readonly", "image", "repeat_count",
               "bind::oc:external"]
# XLSForm column -> the key the Study Spec JSON uses for it (same table as the form builder).
_TO_JSON = {"bind::oc:itemgroup": "bind__oc_itemgroup", "bind::oc:external": "bind__oc_external",
            "bind::oc:briefdescription": "bind__oc_briefdescription", "bind::oc:description": "bind__oc_description",
            "bind::oc:constraint-type": "bind__oc_constraint_type", "bind::oc:required-type": "bind__oc_required_type"}
_NON_DATA = ("begin group", "end group", "begin repeat", "end repeat", "note", "begin_group", "end_group",
             "begin_repeat", "end_repeat")
# Keys the pipeline adds to a survey row that are not form content (never written to the XLSForm).
META_KEYS = {"completion_status", "library_source", "flag_reason", "concept", "concept_qualifier", "concept_source",
             "cdash", "qrs", "edit_checks", "edit_check_exprs", "edit_check_msgs", "edit_check_set_by",
             "edit_check_covered_by", "edit_check_details", "edit_checks_suppressed", "edit_check_suppressed_paths",
             "constraint_engine_only", "constraint_message_locked", "provenance", "protocol_quote", "sdtm"}
# Form-level keys that hold the form's content; everything else on a form is scheduling / metadata.
CONTENT_KEYS = ("survey", "choices", "settings", "extra_cols")

# CDASH / SDTM domain names, used only to recognise a domain from a form title.
DOMAIN_NAMES = {
    "AE": ("adverse event",), "CM": ("concomitant medication", "prior concomitant medication", "prior and concomitant medication"),
    "DM": ("demographic", "demography"), "DS": ("disposition",), "DV": ("protocol deviation", "deviation"),
    "EG": ("ecg", "electrocardiogram", "12 lead ecg"), "EX": ("exposure", "study drug administration"),
    "IE": ("inclusion exclusion criteria", "eligibility criteria", "eligibility", "inclusion exclusion"),
    "LB": ("laboratory", "laboratory test", "lab", "laboratory result"), "MH": ("medical history",),
    "PE": ("physical exam", "physical examination"), "PR": ("procedure", "concomitant procedure"),
    "SU": ("substance use",), "VS": ("vital sign",), "SC": ("subject characteristic",), "CE": ("clinical event",),
    "DA": ("drug accountability",), "DD": ("death detail", "death"), "HO": ("healthcare encounter", "hospitalization"),
    "CO": ("comment",), "PC": ("pharmacokinetic", "pk sampling", "pharmacokinetic sampling"),
    "RS": ("disease response",), "TU": ("tumor identification",), "TR": ("tumor result",),
    "MB": ("microbiology",), "MI": ("microscopic finding",), "QS": ("questionnaire",),
}


def enabled():
    return os.environ.get("STANDARDS_MATCHING", "1") != "0"


def _log(msg):
    print(f"[standards-match] {msg}", flush=True)


# ── File detection ───────────────────────────────────────────────────────────────

def _is_xlsx_container(names):
    return "[Content_Types].xml" in names and any(n.startswith("xl/") for n in names)


def _has_survey_sheet(data):
    """True when the workbook has a 'survey' sheet (decided from the workbook part, no full load)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            wb = zf.read("xl/workbook.xml").decode("utf-8", errors="ignore")
        return bool(re.search(r"<(?:\w+:)?sheet\b[^>]*\bname=\"survey\"", wb, re.I))
    except Exception:
        return False


def _is_odm(data):
    head = data[:4000].lstrip(b"\xef\xbb\xbf \t\r\n")
    if not head.startswith(b"<"):
        return False
    text = head.decode("utf-8", errors="ignore")
    return bool(re.search(r"<(?:\w+:)?ODM[\s>]", text))


def detect_kind(data):
    """'XLSFORM' (workbook with a survey sheet), 'XLSFORM_ZIP' (ZIP with .xlsx members), 'ODM_XML', or
    'UNSTRUCTURED' (everything else, including .docx and workbooks that are not XLSForms: text path)."""
    if not data:
        return "UNSTRUCTURED"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = zf.namelist()
        except Exception:
            return "UNSTRUCTURED"
        if _is_xlsx_container(names):
            return "XLSFORM" if _has_survey_sheet(data) else "UNSTRUCTURED"
        if "[Content_Types].xml" in names:
            return "UNSTRUCTURED"  # .docx / .pptx and other Office containers
        if any(_zip_member_ok(n) and n.lower().endswith((".xlsx", ".xml")) for n in names):
            return "XLSFORM_ZIP"
        return "UNSTRUCTURED"
    return "ODM_XML" if _is_odm(data) else "UNSTRUCTURED"


def _zip_member_ok(name):
    base = os.path.basename(name)
    return bool(base) and not name.startswith("__") and not base.startswith((".", "~$"))


# ── Parsers: one normalised standard-form model ──────────────────────────────────

def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = {}
        for h, v in zip(headers, r):
            s = _cell(v)
            if h and s != "":
                d[h] = s
        if d:
            out.append(d)
    return headers, out


def has_logic(survey):
    """A standard carries logic when any row has a constraint or a show-when (rule 8)."""
    return any(str(r.get("constraint") or "").strip() or str(r.get("relevant") or "").strip()
               for r in survey or [] if isinstance(r, dict))


def _model(form_oid, title, survey, choices, settings, extra_cols, source, source_file, fmt, logic=None,
           repeating=None):
    return {"form_oid": str(form_oid or "").strip(), "title": str(title or "").strip(), "domain": None,
            "domain_basis": "", "survey": survey, "choices": choices, "settings": settings,
            "extra_cols": extra_cols, "has_logic": has_logic(survey) if logic is None else logic,
            "repeating": repeating, "source": source, "source_file": source_file, "format": fmt}


def parse_xlsform(data, file_name, source=SRC_XLSFORM):
    """One XLSForm workbook -> standard form, or None when it has no survey sheet. Cell text is kept as written."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if "survey" not in wb.sheetnames:
        return None
    headers, rows = _sheet_rows(wb["survey"])
    survey = [{_TO_JSON.get(k, k): v for k, v in r.items()} for r in rows]
    extra = [h for h in dict.fromkeys(headers) if h and h not in SURVEY_COLS and any(h in r for r in rows)]
    choices = _sheet_rows(wb["choices"])[1] if "choices" in wb.sheetnames else []
    srows = _sheet_rows(wb["settings"])[1] if "settings" in wb.sheetnames else []
    settings = dict(srows[0]) if srows else {}
    stem = os.path.splitext(os.path.basename(file_name))[0]
    form_oid = settings.get("form_id") or stem
    rep = any(str(r.get("type") or "").strip().lower().replace("_", " ") == "begin repeat" for r in survey)
    return _model(form_oid, settings.get("form_title") or stem, survey, choices, settings, extra, source,
                  os.path.basename(file_name), "xlsform", repeating=rep)


_ODM_TYPES = {"text": "text", "string": "text", "integer": "integer", "float": "decimal", "double": "decimal",
              "decimal": "decimal", "date": "date", "partialdate": "text", "time": "text", "datetime": "text",
              "boolean": "text"}


def _l(el):
    return el.tag.split("}")[-1]


def _translated(el, child):
    for c in el:
        if _l(c) == child:
            for t in c.iter():
                if _l(t) == "TranslatedText" and (t.text or "").strip():
                    return t.text.strip()
    return ""


def _list_name(name, taken):
    base = re.sub(r"[^A-Za-z0-9_]", "_", str(name or "LIST")).strip("_") or "LIST"
    out, n = base, 2
    while out in taken:
        out, n = f"{base}_{n}", n + 1
    return out


def parse_odm(data, file_name, source=SRC_ODM):
    """ODM XML study metadata -> standard forms (FormDef, ItemGroupRef/Def, ItemRef/Def, CodeList,
    OpenClinica:MultiSelectList). ODM carries structure but no logic: has_logic is False."""
    import xml.etree.ElementTree as ET
    root = ET.fromstring(data)
    forms = []
    for mdv in (e for e in root.iter() if _l(e) == "MetaDataVersion"):
        form_defs = [c for c in mdv if _l(c) == "FormDef"]
        if not form_defs:
            continue
        groups = {c.get("OID"): c for c in mdv if _l(c) == "ItemGroupDef"}
        items = {c.get("OID"): c for c in mdv if _l(c) == "ItemDef"}
        code_lists = {c.get("OID"): c for c in mdv if _l(c) == "CodeList"}
        multi = {c.get("ID"): c for c in mdv.iter() if _l(c) == "MultiSelectList"}
        for fd in form_defs:
            survey, choices, lists, any_rep = [], [], {}, False
            for gref in (c for c in fd if _l(c) == "ItemGroupRef"):
                gd = groups.get(gref.get("ItemGroupOID"))
                if gd is None:
                    continue
                gname = gd.get("Name") or gd.get("OID")
                any_rep = any_rep or str(gd.get("Repeating") or "").lower() == "yes"
                refs = [c for c in gd if _l(c) == "ItemRef"]
                refs.sort(key=lambda r: int(r.get("OrderNumber")) if str(r.get("OrderNumber") or "").isdigit() else 10 ** 6)
                for ref in refs:
                    it = items.get(ref.get("ItemOID"))
                    if it is None:
                        continue
                    row = {"type": _ODM_TYPES.get(str(it.get("DataType") or "text").lower(), "text"),
                           "name": it.get("Name") or it.get("OID"), "bind__oc_itemgroup": gname}
                    label = _translated(it, "Question")
                    if label:
                        row["label"] = label
                    cl = next((c for c in it if _l(c) == "CodeListRef"), None)
                    ms = next((c for c in it.iter() if _l(c) == "MultiSelectListRef"), None)
                    src_list = (code_lists.get(cl.get("CodeListOID")) if cl is not None
                                else multi.get(ms.get("MultiSelectListID")) if ms is not None else None)
                    if src_list is not None:
                        key = ("CL" if cl is not None else "MSL", src_list.get("OID") or src_list.get("ID"))
                        if key not in lists:
                            lists[key] = _list_name(src_list.get("Name"), set(lists.values()))
                            for ci in src_list:
                                if _l(ci) in ("CodeListItem", "MultiSelectListItem", "EnumeratedItem"):
                                    code = ci.get("CodedValue") or ci.get("CodedOptionValue") or ""
                                    choices.append({"list_name": lists[key], "label": _translated(ci, "Decode") or code,
                                                    "name": code})
                        row["type"] = f"{'select_one' if cl is not None else 'select_multiple'} {lists[key]}"
                    if str(ref.get("Mandatory") or "").lower() == "yes":
                        row["required"] = "yes"
                    if it.get("BriefDescription"):
                        row["bind__oc_briefdescription"] = it.get("BriefDescription")
                    survey.append(row)
            oid = fd.get("OID") or ""
            fid = oid[2:] if oid.upper().startswith("F_") else oid
            forms.append(_model(fid, fd.get("Name") or fid, survey, choices,
                                {"form_title": fd.get("Name") or fid, "form_id": fid}, [], source,
                                os.path.basename(file_name), "odm", logic=False, repeating=any_rep))
        if forms:
            break  # the first MetaDataVersion with forms is the study; later ones are sites
    return forms


def parse_file(file_name, data, source_xlsform=SRC_XLSFORM, source_odm=SRC_ODM):
    """(forms, record) for one file. record = {"file", "kind", "forms", "usable", "note"}; never raises."""
    rec = {"file": os.path.basename(str(file_name or "file")), "kind": "UNSTRUCTURED", "forms": 0, "usable": False,
           "note": ""}
    forms = []
    try:
        kind = rec["kind"] = detect_kind(data)
        if kind == "XLSFORM":
            f = parse_xlsform(data, file_name, source_xlsform)
            forms = [f] if f else []
        elif kind == "ODM_XML":
            forms = parse_odm(data, file_name, source_odm)
        elif kind == "XLSFORM_ZIP":
            skipped = []
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                for name in zf.namelist():
                    if not _zip_member_ok(name) or name.endswith("/"):
                        continue
                    member = zf.read(name)
                    mk = detect_kind(member)
                    try:
                        if mk == "XLSFORM":
                            f = parse_xlsform(member, name, source_xlsform)
                            forms += [f] if f else []
                        elif mk == "ODM_XML":
                            forms += parse_odm(member, name, source_odm)
                        else:
                            skipped.append(os.path.basename(name))
                    except Exception as e:
                        skipped.append(f"{os.path.basename(name)} ({type(e).__name__})")
            if skipped:
                rec["note"] = "members not usable as forms: " + ", ".join(skipped[:20])
        else:
            rec["note"] = "not an XLSForm or ODM XML; used as reference text for the analysis only"
        forms = [f for f in forms if f and f["survey"]]
        rec["forms"], rec["usable"] = len(forms), bool(forms)
        if kind != "UNSTRUCTURED" and not forms:
            rec["note"] = rec["note"] or "no forms found in the file"
    except Exception as e:
        rec["note"] = f"could not be read ({type(e).__name__}: {e})"
        forms = []
    return forms, rec


# ── CDASH domain of a standard form ──────────────────────────────────────────────

_CDASH = {}


def _cdash():
    """(std, fields, by_var, domains) from the CDISC layer, or (None, {}, {}, set()) when it is unavailable."""
    if "v" not in _CDASH:
        try:
            import cdisc_ct, cdisc_cdash
            std = cdisc_ct.load_standards(None)
            fields = cdisc_cdash.load_cdashig_fields()
            by_var = {}
            for (_d, v), rec in fields.items():
                by_var.setdefault(v, []).append(rec)
            _CDASH["v"] = (std, fields, by_var, {d for d, _ in fields})
        except Exception as e:
            _log(f"CDISC layer unavailable for domain derivation ({e}); form OID/title only")
            _CDASH["v"] = (None, {}, {}, set())
    return _CDASH["v"]


def norm_id(form_id):
    s = re.sub(r"[^A-Z0-9]", "", str(form_id or "").upper())
    return s[1:] if str(form_id or "").upper().startswith("F_") else s


def _norm_title(t):
    words = re.sub(r"[^a-z0-9 ]", " ", str(t or "").lower()).split()
    return " ".join(w[:-1] if len(w) > 3 and w.endswith("s") else w for w in words
                    if w not in ("form", "log", "crf", "the", "and", "of"))


def is_data_row(row):
    if not isinstance(row, dict) or not row.get("name"):
        return False
    return str(row.get("type") or "").strip().lower() not in _NON_DATA


def _field_domains(row):
    """CDASH domains a field's name points to (deterministic: the name, or a part of it, is a CDASHIG variable
    that belongs to one domain)."""
    std, _fields, by_var, _dom = _cdash()
    if std is None:
        return set()
    import cdisc_ct
    name = str(row.get("name") or "").upper()
    out = set()
    for tok in dict.fromkeys([cdisc_ct.field_variable(name, std)] + [t for t in name.split("_") if len(t) >= 4]):
        cands = by_var.get(tok) or []
        doms = {c["domain"] for c in cands}
        own = {d for d in doms if tok.startswith(d)}
        if len(doms) == 1:
            out |= doms
        elif len(own) == 1:
            out |= own
    return out


def derive_domain(form_oid, title, survey):
    """(domain or None, basis, note). From the form OID/name and the fields' CDASH variables."""
    _std, _f, _bv, domains = _cdash()
    domains = domains or set(DOMAIN_NAMES)
    oid = norm_id(form_oid)
    votes = {}
    for r in survey or []:
        if is_data_row(r) and str(r.get("type") or "").strip().lower() != "calculate":
            for d in _field_domains(r):
                votes[d] = votes.get(d, 0) + 1
    ranked = sorted(votes.items(), key=lambda kv: (-kv[1], kv[0]))
    total = sum(votes.values())
    top, n = ranked[0] if ranked else (None, 0)
    if oid in domains or oid in DOMAIN_NAMES:
        note = (f"fields point to {top} ({n} of {total}); the form OID was used"
                if top and top != oid and n >= 3 and n / total >= 0.6 else "")
        return oid, "form OID", note
    n_data = sum(1 for r in survey or [] if is_data_row(r)) or 1
    if top and n >= 2 and n / total >= 0.5 and (n >= 3 or n / n_data >= 0.25):
        second = ranked[1] if len(ranked) > 1 else None
        note = f"also {second[0]} ({second[1]} fields)" if second and second[1] >= 0.8 * n else ""
        return top, f"field concepts ({n} of {total} CDASH fields)", note
    nt = _norm_title(title)
    for d, names in sorted(DOMAIN_NAMES.items()):
        if nt and any(nt == _norm_title(x) for x in names):
            return d, "form title", ""
    if len(oid) > 2 and oid[:2] in domains and votes.get(oid[:2]):
        return oid[:2], f"form OID prefix + {votes[oid[:2]]} CDASH field(s)", ""
    return None, "", ""


# ── Sources ──────────────────────────────────────────────────────────────────────

def _tier(source):
    return 0 if source == SRC_XLSFORM else 2 if source == SRC_ODM else 1


def load_sources(files=None, referenced=None):
    """files: [(file name, bytes)] from the standards column. referenced: [{"label": "referenced study <id>",
    "forms": [(file name, bytes)], "note": ""}] in listed order (Item 3).
    Returns {"forms", "files", "fingerprint"}. Never raises; unreadable files are reported, not fatal."""
    forms, records, parts = [], [], []
    for name, data in files or []:
        fs, rec = parse_file(name, data)
        forms += fs
        records.append(rec)
        parts.append(("file", rec["file"], hashlib.sha256(data or b"").hexdigest()))
    for i, ref in enumerate(referenced or []):
        label = ref.get("label") or f"referenced study {i + 1}"
        n = 0
        for name, data in ref.get("forms") or []:
            fs, rec = parse_file(name, data, source_xlsform=label, source_odm=label)
            for f in fs:
                f["ref_order"] = i
            forms += fs
            n += len(fs)
            parts.append((label, os.path.basename(str(name)), hashlib.sha256(data or b"").hexdigest()))
        records.append({"file": label, "kind": "REFERENCED_STUDY", "forms": n, "usable": n > 0,
                        "note": ref.get("note") or ("" if n else "no forms could be fetched")})
        parts.append((label, "", ""))
    # the same form from several sources: the highest-priority source wins (XLSForm > referenced > ODM)
    forms.sort(key=lambda f: (_tier(f["source"]), f.get("ref_order", 0)))
    kept, seen = [], {}
    for f in forms:
        k = norm_id(f["form_oid"])
        if k in seen:
            _log(f"standard form {f['form_oid']} from {f['source']} ({f['source_file']}) is superseded by the one "
                 f"from {seen[k]['source']} ({seen[k]['source_file']})")
            continue
        seen[k] = f
        f["domain"], f["domain_basis"], f["domain_note"] = derive_domain(f["form_oid"], f["title"], f["survey"])
        kept.append(f)
    fp = hashlib.sha256(json.dumps([VERSION] + sorted(parts)).encode()).hexdigest() if parts else ""
    return {"forms": kept, "files": records, "fingerprint": fp}


def catalog_text(sources, max_fields=400):
    """Compact catalog of the structured standard forms for the analysis context (title, domain, field names)."""
    lines = []
    for f in (sources or {}).get("forms") or []:
        names = [str(r.get("name")) for r in f["survey"] if is_data_row(r)]
        more = f" (+{len(names) - max_fields} more)" if len(names) > max_fields else ""
        lines.append(f"FORM {f['form_oid']} | {f['title']} | CDASH domain: {f['domain'] or 'none'} | source: {f['source']}"
                     f"\n  fields: {', '.join(names[:max_fields])}{more}")
    return "\n".join(lines)


def files_log(sources):
    """One line per file for the monday log: what was used and what was not usable."""
    out = []
    for r in (sources or {}).get("files") or []:
        if r["usable"]:
            out.append(f"{r['file']}: {r['forms']} standard form(s) read"
                       + (f" ({r['note']})" if r.get("note") else ""))
        else:
            out.append(f"{r['file']}: NOT usable for form matching ({r.get('note') or 'no forms'})")
    return out


# ── Matching ─────────────────────────────────────────────────────────────────────

def _title_score(a, b):
    ta, tb = _norm_title(a), _norm_title(b)
    if not ta or not tb:
        return 0.0
    wa, wb = set(ta.split()), set(tb.split())
    return max(difflib.SequenceMatcher(None, ta, tb).ratio(), len(wa & wb) / len(wa | wb))


def _similarity(pform, sform):
    """Name similarity of a protocol form and a standard form: 1.0 for the same form id, otherwise the titles
    (short form ids such as RT / PR share letters by chance, so partial id overlap only counts as containment)."""
    a, b = norm_id(pform.get("form_id")), norm_id(sform["form_oid"])
    if a and a == b:
        return 1.0
    contained = 0.6 if min(len(a), len(b)) >= 3 and (a in b or b in a) else 0.0
    return round(max(contained, _title_score(pform.get("form_title"), sform["title"])), 3)


def _protocol_domains(form):
    raw = str(form.get("cdash_domain") or "").upper()
    doms = [d for d in re.split(r"[^A-Z0-9]+", raw) if d]
    if not doms:
        d, _b, _n = derive_domain(form.get("form_id"), form.get("form_title"), form.get("survey"))
        doms = [d] if d else []
    return doms


def match_forms(pforms, sforms):
    """[(protocol form, standard form, basis, score, note)] one-to-one. Domain first (by source priority, then name
    similarity inside a domain), then identical form id / title for forms the domain pass left unmatched."""
    out, used_p, used_s = [], set(), set()
    tiers = sorted({(_tier(s["source"]), s.get("ref_order", 0)) for s in sforms})
    pdoms = {id(p): _protocol_domains(p) for p in pforms}

    def take(p, s, basis, score, note):
        used_p.add(id(p)); used_s.add(id(s))
        out.append((p, s, basis, score, note))

    for tier in tiers:
        tier_forms = [s for s in sforms if (_tier(s["source"]), s.get("ref_order", 0)) == tier]
        for dom in sorted({s["domain"] for s in tier_forms if s["domain"]}):
            cand_s = [s for s in tier_forms if s["domain"] == dom and id(s) not in used_s]
            cand_p = [p for p in pforms if dom in pdoms[id(p)] and id(p) not in used_p]
            pairs = sorted(((_similarity(p, s), pi, si) for pi, p in enumerate(cand_p) for si, s in enumerate(cand_s)),
                           key=lambda t: (-t[0], t[1], t[2]))
            for score, pi, si in pairs:
                p, s = cand_p[pi], cand_s[si]
                if id(p) in used_p or id(s) in used_s:
                    continue
                rivals = [(sc, cand_s[j]["form_oid"]) for sc, i, j in pairs if i == pi and j != si
                          and id(cand_s[j]) not in used_s]
                rivals_p = [(sc, cand_p[i].get("form_id")) for sc, i, j in pairs if j == si and i != pi
                            and id(cand_p[i]) not in used_p]
                note = ""
                if rivals or rivals_p:
                    close = [f"{n} ({sc})" for sc, n in rivals + rivals_p if score - sc <= 0.05]
                    note = (f"{len(cand_s)} standard and {len(cand_p)} protocol form(s) share domain {dom}; chosen by "
                            f"name similarity {score}" + (f"; AMBIGUOUS, close alternatives: {', '.join(close)}" if close else ""))
                take(p, s, f"CDASH domain {dom}", score, note)
    if os.environ.get("STANDARDS_MATCH_BY_NAME", "1") != "0":
        for tier in tiers:
            for s in [s for s in sforms if (_tier(s["source"]), s.get("ref_order", 0)) == tier and id(s) not in used_s]:
                best = None
                for p in pforms:
                    if id(p) in used_p:
                        continue
                    same_id = norm_id(p.get("form_id")) == norm_id(s["form_oid"]) != ""
                    tscore = _title_score(p.get("form_title"), s["title"])
                    both_dom = bool(pdoms[id(p)]) and bool(s["domain"])
                    if same_id or (tscore >= 0.9 and not both_dom):
                        cand = (1.0 if same_id else round(tscore, 3), "same form id" if same_id else "form title", p)
                        if best is None or cand[0] > best[0]:
                            best = cand
                if best:
                    p = best[2]
                    take(p, s, f"{best[1]} (no shared CDASH domain: protocol {'/'.join(pdoms[id(p)]) or 'none'}, "
                               f"standard {s['domain'] or 'none'})", best[0], "")
    return out


# ── Splice ───────────────────────────────────────────────────────────────────────

def core_row(row):
    return {k: v for k, v in row.items() if k not in META_KEYS}


def content_sha(survey, choices, settings):
    payload = {"survey": [core_row(r) for r in survey or [] if isinstance(r, dict)],
               "choices": [{k: v for k, v in c.items() if k != "source"} for c in choices or [] if isinstance(c, dict)],
               "settings": {k: v for k, v in (settings or {}).items() if k != "version"}}
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _field_summary(form):
    return [{"name": r.get("name"), "type": r.get("type"), "label": str(r.get("label") or "")[:160]}
            for r in form.get("survey") or [] if is_data_row(r)]


def _rename_refs(spec, old, new, skip):
    """Point everything that referenced the protocol form's id at the standard form's id."""
    if old == new or not old:
        return
    o_f, n_f = ("F_" + norm_id(old)), ("F_" + norm_id(new))
    for pl in ((spec.get("schedule_of_events") or {}).get("form_placements") or []) \
            if isinstance(spec.get("schedule_of_events"), dict) else []:
        if isinstance(pl, dict) and pl.get("form_id") == old:
            pl["form_id"] = new

    def fix(text):
        if not isinstance(text, str) or norm_id(old) not in text.upper():
            return text
        text = re.sub(r"(FormOID=')(?:F_)?" + re.escape(norm_id(old)) + r"(')", lambda m: m.group(1) + n_f + m.group(2),
                      text, flags=re.I)
        return text

    for f in spec.get("forms") or []:
        if not isinstance(f, dict) or f is skip:
            continue
        for dep in f.get("cross_form_dependencies") or []:
            if not isinstance(dep, dict):
                continue
            if dep.get("source_form") in (old, o_f):
                dep["source_form"] = new
            for k in ("source_item_oid", "source_itemgroup_oid"):
                v = dep.get(k)
                if isinstance(v, str) and v.split(".", 1)[0] in (old, o_f) and "." in v:
                    dep[k] = new + "." + v.split(".", 1)[1]
            if dep.get("xpath_expression"):
                dep["xpath_expression"] = fix(dep["xpath_expression"])
        if f.get("customer_standard"):
            continue  # a standard form's own logic is never rewritten
        for r in f.get("survey") or []:
            if isinstance(r, dict) and r.get("calculation"):
                r["calculation"] = fix(r["calculation"])


def _splice(spec, pform, sform, basis, score, note):
    """Replace the protocol form's content with the standard form, exactly. Returns the match record."""
    old_id = pform.get("form_id")
    original = copy.deepcopy(pform)
    survey = []
    for r in sform["survey"]:
        row = copy.deepcopy(r)
        row.update({"library_source": "CUSTOM", "completion_status": "COMPLETE", "flag_reason": ""})
        survey.append(row)
    choices = [{**copy.deepcopy(c), "source": STATUS} for c in sform["choices"]]
    settings = copy.deepcopy(sform["settings"])
    new_id = str(settings.get("form_id") or sform["form_oid"] or old_id)
    taken = {f.get("form_id") for f in spec.get("forms") or [] if isinstance(f, dict) and f is not pform}
    id_note = ""
    if new_id in taken:
        id_note = f"standard form id {new_id} is already used by another form; kept {old_id}"
        new_id = old_id
        settings["form_id"] = old_id
    n_data = sum(1 for r in survey if is_data_row(r))
    pform["survey"], pform["choices"], pform["settings"] = survey, choices, settings
    pform["extra_cols"] = list(sform.get("extra_cols") or [])
    pform["form_id"] = new_id
    pform["form_title"] = sform["title"] or pform.get("form_title")
    if not pform.get("cdash_domain") and sform.get("domain"):
        pform["cdash_domain"] = sform["domain"]
    if sform.get("repeating") is not None:
        pform["has_repeating_group"] = bool(sform["repeating"])
    pform["cross_form_dependencies"] = []
    pform["customer_form_name"] = sform["title"] or sform["form_oid"]
    pform["customer_standard"] = {
        "form_name": sform["title"] or sform["form_oid"], "form_oid": sform["form_oid"], "source": sform["source"],
        "source_file": sform["source_file"], "domain": sform.get("domain"), "domain_basis": sform.get("domain_basis"),
        "match_basis": basis, "score": score, "has_logic": bool(sform["has_logic"]), "replaced_form_id": old_id,
        # an original XLSForm is built verbatim; an ODM gives structure only, so the form builder lays it out
        "verbatim": sform.get("format") == "xlsform",
        "content_sha": content_sha(survey, choices, settings), "added_fields": [], "approved": []}
    pform["library_match"] = {"status": STATUS, "source_type": STATUS, "customer_form_name": pform["customer_form_name"],
                              "source": sform["source"], "fields_from_library": n_data,
                              "fields_extended_from_protocol": 0, "fields_from_cdash_default": 0}
    spec.setdefault("standards_originals", {})[new_id] = original
    _rename_refs(spec, old_id, new_id, pform)
    return {"protocol_form": old_id, "protocol_title": original.get("form_title"), "form_id": new_id,
            "standard_form": sform["form_oid"], "standard_title": sform["title"], "source": sform["source"],
            "source_file": sform["source_file"], "domain": sform.get("domain"),
            "protocol_domain": original.get("cdash_domain"), "basis": basis, "score": score,
            "note": "; ".join(x for x in (note, id_note, sform.get("domain_note")) if x),
            "has_logic": bool(sform["has_logic"]), "fields_kept": n_data, "fields_added": []}


def _restore(spec):
    """Undo an earlier match (sources changed): matched forms get their protocol-analysed content back, keeping
    the current scheduling."""
    originals = spec.get("standards_originals") or {}
    forms = spec.get("forms") or []
    for i, f in enumerate(forms):
        if not isinstance(f, dict) or not f.get("customer_standard"):
            continue
        orig = originals.get(f.get("form_id"))
        if not isinstance(orig, dict):
            continue
        restored = copy.deepcopy(orig)
        for k, v in f.items():
            if k not in CONTENT_KEYS and k not in ("form_id", "form_title", "customer_standard", "customer_form_name",
                                                   "library_match", "cross_form_dependencies", "cdash_domain",
                                                   "has_repeating_group"):
                restored[k] = v
        _rename_refs(spec, f.get("form_id"), restored.get("form_id"), f)
        forms[i] = restored
    spec.pop("standards_originals", None)
    sm = spec.get("study_meta") if isinstance(spec.get("study_meta"), dict) else {}
    prev = sm.get("standards_match") or {}
    prev["proposals"] = []


def state(spec):
    sm = (spec or {}).get("study_meta") if isinstance(spec, dict) else None
    return (sm or {}).get("standards_match") or {} if isinstance(sm, dict) else {}


def needs_match(spec, sources):
    """True when the spec has not been matched against these sources (rule 9)."""
    if not enabled() or not isinstance(spec, dict):
        return False
    fp = (sources or {}).get("fingerprint") or ""
    prev = state(spec)
    if not fp and not prev.get("fingerprint"):
        return False
    return prev.get("fingerprint") != fp or prev.get("version") != VERSION


def apply(spec, sources):
    """Match and splice. Returns a NEW spec (the input is never mutated); on any error the input is returned
    unchanged. A spec already matched against the same sources is returned as is."""
    if not needs_match(spec, sources):
        return spec
    try:
        out = copy.deepcopy(spec)
        prev = state(out)
        if prev.get("matched"):
            _restore(out)
        sforms = (sources or {}).get("forms") or []
        pforms = [f for f in out.get("forms") or [] if isinstance(f, dict)]
        matched = [_splice(out, p, s, basis, score, note) for p, s, basis, score, note in match_forms(pforms, sforms)]
        used = {m["standard_form"] for m in matched}
        got = {m["form_id"] for m in matched}
        warnings = _dangling_refs(out, got)
        sm = out.setdefault("study_meta", {})
        sm["standards_match"] = {
            "version": VERSION, "fingerprint": (sources or {}).get("fingerprint") or "",
            "files": (sources or {}).get("files") or [], "matched": matched,
            "protocol_forms_without_standard": [f.get("form_id") for f in pforms if f.get("form_id") not in got],
            "standard_forms_not_used": [{"form": s["form_oid"], "title": s["title"], "domain": s["domain"],
                                         "source": s["source"]} for s in sforms if s["form_oid"] not in used],
            "warnings": warnings, "proposals": [], "rejected_ids": list(prev.get("rejected_ids") or []),
            "additions": {"status": "not_run"}}
        for m in matched:
            _log(f"{m['protocol_form']} -> {m['standard_form']} ({m['source']}; {m['basis']}; similarity {m['score']})"
                 + (f"; {m['note']}" if m["note"] else ""))
        return out
    except Exception as e:
        _log(f"matching failed, spec unchanged: {type(e).__name__}: {e}")
        return spec


def _dangling_refs(spec, matched_ids):
    out = []
    by_id = {f.get("form_id"): f for f in spec.get("forms") or [] if isinstance(f, dict)}
    for f in by_id.values():
        for dep in f.get("cross_form_dependencies") or []:
            if isinstance(dep, dict) and dep.get("source_form") in matched_ids:
                names = {r.get("name") for r in by_id[dep["source_form"]].get("survey") or [] if isinstance(r, dict)}
                if dep.get("source_field") not in names:
                    out.append(f"{f.get('form_id')} reads {dep.get('source_form')}.{dep.get('source_field')}, which the "
                               f"customer standard form does not have")
    return out


def matched_forms(spec):
    return [f for f in (spec or {}).get("forms") or [] if isinstance(f, dict) and f.get("customer_standard")]


def integrity(spec):
    """Matched forms whose standard content no longer equals what the customer provided (should be none unless a
    DM approved a proposal). [{"form_id", "reason"}]."""
    out = []
    for f in matched_forms(spec):
        cs = f["customer_standard"]
        rows = [r for r in f.get("survey") or [] if isinstance(r, dict) and r.get("provenance") != ADDED
                and not r.get("standard_proposal")]
        added_lists = {a.get("list_name") for a in cs.get("added_fields") or [] if a.get("list_name")}
        choices = [c for c in f.get("choices") or [] if isinstance(c, dict) and c.get("list_name") not in added_lists
                   and c.get("source") == STATUS]
        if content_sha(rows, choices, f.get("settings")) != cs.get("content_sha"):
            out.append({"form_id": f.get("form_id"),
                        "reason": "changed by approved proposals" if cs.get("approved") else "content differs from the standard"})
    return out


# ── Rule 3: add a field only when the protocol clearly specifies it ──────────────

def _norm_text(t):
    t = str(t or "").replace("­", "").replace("’", "'").replace("‘", "'")
    t = t.replace("“", '"').replace("”", '"').replace("–", "-").replace("—", "-")
    t = re.sub(r"-\s*\n\s*", "", t)  # words hyphenated across a line break
    return re.sub(r"\s+", " ", t).strip().lower()


def _concept_keys(form_id, domain, rows):
    """{row name: (concept, qualifier)} by the pipeline's deterministic tagging (aliases are applied later)."""
    std = _cdash()[0]
    out = {}
    if std is None:
        return out
    try:
        import cdisc_concepts
        tmp = {"forms": [{"form_id": form_id, "cdash_domain": domain,
                          "survey": [{"name": r.get("name"), "type": r.get("type")} for r in rows if is_data_row(r)]}]}
        cdisc_concepts.tag_deterministic(tmp, std)
        for r in tmp["forms"][0]["survey"]:
            if r.get("concept"):
                out[r["name"]] = (r["concept"], r.get("concept_qualifier"))
    except Exception as e:
        _log(f"concept tagging for the field comparison failed ({e}); names only")
    return out


def _name_key(name):
    """Field name with numbering normalised (IE01 == IE001) for comparing the same data point across namings."""
    return re.sub(r"0+(\d)", r"\1", str(name or "").upper())


def _covered_by(name, label, s_rows):
    """Name of the standard field that already collects this data point, judged deterministically from the names
    (same name, same numbered item, one name is a part of the other: SYSBP / SYSBP_VSORRES, AEREL / AEREL1) or
    from near-identical labels. None when nothing obvious covers it (the AI check then decides)."""
    key = _name_key(name)
    lab = _norm_title(label)
    for r in s_rows:
        sname = str(r.get("name") or "")
        skey = _name_key(sname)
        if key == skey:
            return sname
        if len(key) >= 3 and len(skey) >= 3:
            toks_s, toks_c = skey.split("_"), key.split("_")
            if key in toks_s or skey in toks_c:
                return sname
            short, long_ = sorted((key, skey), key=len)
            if long_.startswith(short) and len(short) >= 4 and long_[len(short):].strip("_").isdigit():
                return sname
        slab = _norm_title(r.get("label"))
        if lab and slab and difflib.SequenceMatcher(None, lab, slab).ratio() >= 0.85:
            return sname
    return None


def add_candidates(spec, covered=None):
    """Fields of the analysed protocol form whose CDASH concept (or, when untagged, name and label) the matched
    standard form does not have. [{"form_id", "field", "label", "type", "concept", "row", "choices"}]."""
    out = []
    originals = (spec or {}).get("standards_originals") or {}
    for f in matched_forms(spec):
        orig = originals.get(f.get("form_id"))
        if not isinstance(orig, dict):
            continue
        dom = f.get("cdash_domain")
        s_rows = [r for r in f.get("survey") or [] if is_data_row(r)]
        p_rows = [r for r in orig.get("survey") or [] if is_data_row(r)
                  and str(r.get("type") or "").strip().lower() != "calculate"
                  and not str(r.get("name")).upper().endswith(("_CF", "_SF"))]
        s_con, p_con = _concept_keys(f.get("form_id"), dom, s_rows), _concept_keys(f.get("form_id"), dom, p_rows)
        have_con = set(s_con.values()) | {(c, None) for c, _q in s_con.values()}
        by_con = {}
        for n, c in s_con.items():
            by_con.setdefault(c, n)
            by_con.setdefault((c[0], None), n)
        for r in p_rows:
            name = str(r.get("name"))
            con = p_con.get(name)
            same = by_con.get(con) if con and (con in have_con or not con[1]) else None
            same = same or _covered_by(name, r.get("label"), s_rows)
            if same:
                if covered is not None:
                    covered.append({"form_id": f.get("form_id"), "field": name, "covered_by": same})
                continue
            t = str(r.get("type") or "")
            ln = t.split(" ", 1)[1].strip() if t.startswith("select") and " " in t else ""
            out.append({"form_id": f.get("form_id"), "field": name, "label": str(r.get("label") or ""), "type": t,
                        "concept": con[0] if con else "", "row": r,
                        "standard_fields": [str(x.get("name")) for x in s_rows],
                        "choices": [c for c in orig.get("choices") or [] if isinstance(c, dict)
                                    and c.get("list_name") == ln] if ln else []})
    return out


ADD_PROMPT = """You compare data points of a clinical trial eCRF form with a customer's standard form and with the PROTOCOL.

A customer standard eCRF form is being used exactly as the customer provided it. CANDIDATES lists data points that
an earlier analysis attached to that form under its own field names. For EVERY candidate return one entry:

1. covered_by: look through STANDARD FORM FIELDS (name | type | label | choices) for the field that collects the
   same data point: the same measurement, date, criterion or answer, whatever it is named. Names differ between
   the two (a result may be SYSBP in one and SYSBP_VSORRES in the other, a criterion IE01 and IE001, one
   relationship question may be split into several). If the standard form collects it, set covered_by to that
   standard field's name and set quote to null. When in doubt, it is covered.
2. Only when NO standard field collects it: does the PROTOCOL TEXT clearly and directly say that THIS data point
   must be collected or recorded? A general statement (e.g. "adverse events will be recorded") does NOT specify
   an individual field, and a sentence that merely mentions the topic (how many doses are given, that a visit
   takes place) does not either: the quote must itself name this data point as something to collect, record,
   measure, document or report. If it does, set covered_by to null and copy the VERBATIM sentence or phrase from
   the PROTOCOL TEXT into quote. If it does not, set both to null.

Answer ONLY with JSON, no prose, no code fences:
{"additions": [{"form_id": "...", "field": "...", "covered_by": "<standard field name>" | null,
                "quote": "<verbatim protocol text>" | null}]}
The quote is checked against the protocol character by character: do not paraphrase, shorten with "...", or
correct it. A field is added only when covered_by is null and the quote is found in the protocol; "nothing to add"
is the expected outcome for most candidates.
"""


def build_add_request(spec, protocol_text, max_chars=600_000):
    """(prompt, extra_text) for the candidates, or None when there is nothing to check."""
    cands = add_candidates(spec)
    if not cands or not str(protocol_text or "").strip():
        return None
    lines = []
    for f in matched_forms(spec):
        mine = [c for c in cands if c["form_id"] == f.get("form_id")]
        if not mine:
            continue
        lines.append(f"\nFORM {f.get('form_id')} | {f.get('form_title')} | CDASH domain {f.get('cdash_domain') or '-'}")
        lists = {}
        for c in f.get("choices") or []:
            if isinstance(c, dict):
                lists.setdefault(c.get("list_name"), []).append(str(c.get("label") or c.get("name")))
        lines.append("  STANDARD FORM FIELDS:")
        std_lines = []
        for r in f.get("survey") or []:
            if not is_data_row(r):
                continue
            t = str(r.get("type") or "")
            ln = t.split(" ", 1)[1].strip() if t.startswith("select") and " " in t else ""
            std_lines.append(f"    {r.get('name')} | {t.split(' ')[0]} | {str(r.get('label') or '')[:90]}"
                             + (f" | {'; '.join(lists.get(ln, [])[:6])}" if ln in lists else ""))
        lines.append("\n".join(std_lines)[:60000])
        lines.append("  CANDIDATES:")
        for c in mine:
            lines.append(f"    {c['field']} | {c['type'].split(' ')[0]} | {c['label'][:140]}"
                         + (f" | CDASH {c['concept']}" if c["concept"] else ""))
    extra = "CANDIDATES BY FORM:" + "\n".join(lines) + "\n\nPROTOCOL TEXT:\n" + str(protocol_text)[:max_chars]
    return ADD_PROMPT, extra


def _parse_json(text, key):
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", str(text or "").strip())
    try:
        data = json.loads(t)
    except Exception:
        m = re.search(r"\{.*\}", t, re.S)
        data = json.loads(m.group(0)) if m else {}
    v = data.get(key) if isinstance(data, dict) else None
    return v if isinstance(v, list) else []


_REF = re.compile(r"\$\{([^}]+)\}")


def apply_additions(spec, response_text, protocol_text):
    """Validate the AI's answer and add the fields that pass. A field is added only when its quote is found in the
    protocol text (whitespace-normalised). Returns {"candidates", "added", "rejected"}; mutates spec in place."""
    pre = []
    cands = {(c["form_id"], c["field"]): c for c in add_candidates(spec, pre)}
    result = {"status": "done", "candidates": len(cands), "added": [], "rejected": [], "covered": pre,
              "not_specified": []}
    try:
        answers = _parse_json(response_text, "additions")
    except Exception:
        result["rejected"].append({"reason": "unparseable_response"})
        answers = []
    proto = _norm_text(protocol_text)
    forms = {f.get("form_id"): f for f in matched_forms(spec)}
    done = set()
    for a in answers:
        if not isinstance(a, dict):
            result["rejected"].append({"reason": "malformed"})
            continue
        key = (str(a.get("form_id")), str(a.get("field")))
        rej = {"form_id": key[0], "field": key[1]}
        c = cands.get(key)
        quote = str(a.get("quote") or "").strip()
        nq = _norm_text(quote)
        cov = str(a.get("covered_by") or "").strip()
        if c is None:
            result["rejected"].append({**rej, "reason": "not_a_candidate"})
        elif key in done:
            result["rejected"].append({**rej, "reason": "duplicate"})
        elif cov and cov.lower() not in ("null", "none"):
            if cov in c["standard_fields"]:
                result["covered"].append({**rej, "covered_by": cov})
            else:  # says "covered" but names no real standard field: not added either way
                result["rejected"].append({**rej, "reason": "covered_by_unknown_field", "covered_by": cov[:60]})
        elif not quote or quote.lower() in ("null", "none"):
            result["not_specified"].append(rej)
        elif len(nq) < 12 or len(nq.split()) < 3:
            result["rejected"].append({**rej, "reason": "quote_too_short"})
        elif nq not in proto:
            result["rejected"].append({**rej, "reason": "quote_not_in_protocol", "quote": quote[:200]})
        else:
            done.add(key)
            result["added"].append(_add_field(forms[key[0]], c, quote))
    for m in state(spec).get("matched") or []:
        m["fields_added"] = [x for x in result["added"] if x["form_id"] == m["form_id"]]
    state(spec)["additions"] = {k: v for k, v in result.items()}
    return result


def _add_field(form, cand, quote):
    """Append one protocol-specified field to a standard form. Never changes or removes a standard row."""
    row = {k: copy.deepcopy(v) for k, v in cand["row"].items() if k not in META_KEYS or k in ("concept", "concept_qualifier", "concept_source")}
    names = {r.get("name") for r in form.get("survey") or [] if isinstance(r, dict)}
    dropped = []
    for k in ("relevant", "constraint", "calculation"):  # protocol logic is kept only when everything it reads exists
        refs = _REF.findall(str(row.get(k) or ""))
        if refs and any(x not in names and x != row.get("name") for x in refs):
            dropped.append(k)
            row.pop(k, None)
            if k == "constraint":
                row.pop("constraint_message", None)
    concept = cand.get("concept")
    if concept and not str(row.get("label") or "").strip():
        _std, fields, by_var, _d = _cdash()
        rec = next(iter(by_var.get(concept) or []), None)
        row["label"] = (rec or {}).get("prompt") or (rec or {}).get("question") or row.get("name")
    groups = [str(r.get("bind__oc_itemgroup")) for r in form.get("survey") or [] if is_data_row(r)
              and r.get("bind__oc_itemgroup")]
    if groups:
        row["bind__oc_itemgroup"] = groups[-1]
    list_name = ""
    t = str(row.get("type") or "")
    if cand.get("choices") and t.startswith("select") and " " in t:
        kind, ln = t.split(" ", 1)
        ln = ln.strip()
        mine = [(str(c.get("name")), str(c.get("label"))) for c in cand["choices"]]
        theirs = [(str(c.get("name")), str(c.get("label"))) for c in form.get("choices") or [] if c.get("list_name") == ln]
        if theirs and theirs != mine:
            taken = {c.get("list_name") for c in form.get("choices") or []}
            ln = _list_name(ln + "_PROT", taken)
            row["type"] = f"{kind} {ln}"
        if not theirs or theirs != mine:
            list_name = ln
            form.setdefault("choices", []).extend({**{k: v for k, v in c.items() if k != "source"},
                                                   "list_name": ln, "source": "PROTOCOL"} for c in cand["choices"])
    row.update({"library_source": "PROTOCOL_SPECIFIC", "completion_status": "COMPLETE", "provenance": ADDED,
                "protocol_quote": quote, "flag_reason": ""})
    if dropped:
        row.update({"completion_status": "FLAGGED",
                    "flag_reason": f"Protocol logic ({', '.join(dropped)}) reads fields the customer standard form does "
                                   f"not have; it was not carried over."})
    form.setdefault("survey", []).append(row)
    rec = {"form_id": form.get("form_id"), "field": row.get("name"), "label": str(row.get("label") or ""),
           "concept": concept or "", "quote": quote, "list_name": list_name, "logic_dropped": dropped}
    cs = form["customer_standard"]
    cs.setdefault("added_fields", []).append(rec)
    lm = form.setdefault("library_match", {})
    lm["fields_extended_from_protocol"] = len(cs["added_fields"])
    lm["added_from_protocol"] = [{"field": a["field"], "quote": a["quote"]} for a in cs["added_fields"]]
    return rec


# ── Proposals on standard forms (rules 7 and 8) ──────────────────────────────────

def proposals(spec):
    return state(spec).get("proposals") or []


def set_proposals(spec, kind, items):
    """Replace all proposals of one kind ("engine" | "ai"). Rejected ones are never listed again."""
    if not isinstance(spec, dict) or not state(spec):
        return
    st = spec["study_meta"]["standards_match"]
    rejected = set(st.get("rejected_ids") or [])
    st["proposals"] = [p for p in st.get("proposals") or [] if p.get("kind") != kind] + \
                      [p for p in items if p.get("id") not in rejected]


def proposal_id(kind, *parts):
    h = hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()[:8].upper()
    return f"{'STD' if kind == 'engine' else 'AIS'}-{h}"


def _row(form, name):
    return next((r for r in form.get("survey") or [] if isinstance(r, dict) and r.get("name") == name), None)


def _join_message(cur, msg):
    cur, msg = str(cur or "").strip(), str(msg or "").strip()
    if not msg or msg in cur:
        return cur
    return f"{cur}{' ' if cur.endswith(('.', '!', '?')) else '. '}{msg}" if cur else msg


def _plan(form, prop):
    """Validate a proposal against the form as it is now. Returns (steps, None) or (None, reason)."""
    steps = []
    for op in prop.get("ops") or []:
        kind = op.get("op")
        if kind == "row":
            row = _row(form, op.get("field"))
            if row is None:
                return None, f"field {op.get('field')} is not on the form"
            for k, new in (op.get("set") or {}).items():
                cur, before = row.get(k), (op.get("before") or {}).get(k)
                if k == "constraint_message":
                    continue  # handled with the constraint
                if k == "constraint":
                    clause = op.get("clause") or new
                    if clause and clause in str(cur or ""):
                        continue
                    value = new if (cur or None) == (before or None) else f"({cur}) and ({clause})"
                    steps.append(("constraint", row, value, op.get("message") or (op.get("set") or {}).get("constraint_message"),
                                  (op.get("set") or {}).get("constraint_message") if (cur or None) == (before or None) else None))
                elif (cur or None) == (new or None):
                    continue
                elif (cur or None) in ((before or None), None):
                    steps.append(("set", row, k, new))
                else:
                    return None, f"{op.get('field')}.{k} was changed since the proposal was made"
        elif kind == "insert":
            if _row(form, (op.get("row") or {}).get("name")) is None:
                steps.append(("insert", op.get("after"), op.get("row")))
        elif kind == "choice":
            c = op.get("choice") or {}
            if not any(x.get("list_name") == c.get("list_name") and str(x.get("name")) == str(c.get("name"))
                       for x in form.get("choices") or []):
                steps.append(("choice", c))
        elif kind == "setting":
            steps.append(("setting", op.get("key"), op.get("value")))
        elif kind == "dependency":
            if op.get("value") not in (form.get("cross_form_dependencies") or []):
                steps.append(("dependency", op.get("value")))
        else:
            return None, f"unknown operation {kind!r}"
    return steps, None


def apply_proposal(spec, prop):
    """Apply one approved proposal to its customer standard form, deterministically. Validates everything first;
    nothing is changed unless the whole proposal applies. Returns (status, note)."""
    form = next((f for f in spec.get("forms") or [] if isinstance(f, dict)
                 and norm_id(f.get("form_id")) == norm_id(prop.get("target_form"))), None)
    if form is None:
        return "not_found", "form not in the study"
    steps, why = _plan(form, prop)
    if why:
        return "failed", why
    if not steps:
        status, note = "already", "already in the build"
    else:
        for st in steps:
            if st[0] == "constraint":
                _k, row, value, msg, full_msg = st
                row["constraint"] = value
                row["constraint_message"] = full_msg if full_msg else _join_message(row.get("constraint_message"), msg)
                if not row["constraint_message"]:
                    row.pop("constraint_message")
            elif st[0] == "set":
                _k, row, key, new = st
                if new in (None, ""):
                    row.pop(key, None)
                else:
                    row[key] = new
            elif st[0] == "insert":
                survey = form.setdefault("survey", [])
                idx = next((i for i, r in enumerate(survey) if isinstance(r, dict) and r.get("name") == st[1]), len(survey) - 1)
                survey.insert(idx + 1, {**copy.deepcopy(st[2]), "standard_proposal": prop.get("id")})
            elif st[0] == "choice":
                form.setdefault("choices", []).append({**copy.deepcopy(st[1]), "source": "APPROVED_PROPOSAL"})
            elif st[0] == "setting":
                settings = form.setdefault("settings", {})
                if st[1] == "crossform_references":
                    cur = [e for e in str(settings.get(st[1]) or "").split(",") if e]
                    settings[st[1]] = ",".join(cur + [e for e in str(st[2] or "").split(",") if e and e not in cur])
                else:
                    settings[st[1]] = st[2]
            elif st[0] == "dependency":
                form.setdefault("cross_form_dependencies", []).append(copy.deepcopy(st[1]))
        status, note = "applied", "added to the build"
    row = _row(form, prop.get("target_field"))
    cid = prop.get("check_id") or prop.get("id")
    if row is not None and prop.get("check_type") == "Constraint":
        if cid not in (row.get("edit_checks") or []):
            row.setdefault("edit_checks", []).append(cid)
        row.setdefault("edit_check_exprs", {})[cid] = prop.get("logic")
        if prop.get("message"):
            row.setdefault("edit_check_msgs", {})[cid] = prop.get("message")
        if not any(isinstance(d, dict) and d.get("id") == cid for d in row.get("edit_check_details") or []):
            row.setdefault("edit_check_details", []).append(
                {"id": cid, "source": prop.get("source"), "category": prop.get("category") or "",
                 "rationale": prop.get("rationale") or "Approved by a data manager on a customer standard form.",
                 "protocol_reference": prop.get("protocol_reference") or ""})
    elif row is not None and prop.get("check_type") in ("Conditional Display", "Required"):
        key = "relevant" if prop.get("check_type") == "Conditional Display" else "required"
        row.setdefault("edit_check_set_by", {})[key] = prop.get("convention_id") or prop.get("id")
        row.setdefault("edit_check_source", {})[key] = prop.get("source")
    cs = form.setdefault("customer_standard", {})
    if not any(a.get("id") == prop.get("id") for a in cs.get("approved") or []):
        cs.setdefault("approved", []).append({"id": prop.get("id"), "convention_id": prop.get("convention_id"),
                                              "field": prop.get("target_field"), "source": prop.get("source")})
    st = state(spec)
    if st:
        st["proposals"] = [p for p in st.get("proposals") or [] if p.get("id") != prop.get("id")]
    return status, note


def reject_proposal(spec, pid):
    st = state(spec)
    if not st:
        return False
    before = len(st.get("proposals") or [])
    st["proposals"] = [p for p in st.get("proposals") or [] if p.get("id") != pid]
    if pid and pid not in st.setdefault("rejected_ids", []):
        st["rejected_ids"].append(pid)
    return len(st["proposals"]) < before


def is_approved(form, convention_id, field):
    return any(a.get("convention_id") == convention_id and a.get("field") == field
               for a in ((form or {}).get("customer_standard") or {}).get("approved") or [])


# ── Reporting ────────────────────────────────────────────────────────────────────

def summary_lines(spec):
    """Lines for the monday log: every match decision, forms without a standard, standard forms not used."""
    st = state(spec)
    if not st:
        return []
    out = [f"Customer standard forms: {len(st.get('matched') or [])} protocol form(s) use a customer standard form "
           f"exactly as provided."]
    for m in st.get("matched") or []:
        added = m.get("fields_added") or []
        out.append(f"  {m['protocol_form']} -> {m['standard_title']} [{m['standard_form']}] ({m['source']}; {m['basis']}; "
                   f"{m['fields_kept']} fields kept"
                   + (f", {len(added)} added from protocol: {', '.join(a['field'] for a in added)}" if added else "")
                   + ("; no logic in the standard" if not m.get("has_logic") else "") + ")"
                   + (f" NOTE: {m['note']}" if m.get("note") else ""))
    if st.get("protocol_forms_without_standard"):
        out.append("  No customer standard for: " + ", ".join(str(x) for x in st["protocol_forms_without_standard"]))
    if st.get("standard_forms_not_used"):
        out.append("  Standard forms the protocol does not need: "
                   + ", ".join(f"{s['form']} ({s['domain'] or 'no domain'})" for s in st["standard_forms_not_used"]))
    for w in st.get("warnings") or []:
        out.append(f"  WARNING: {w}")
    return out


def provenance_by_domain(spec):
    """{CDASH domain or form id (upper): (customer form name, source)} for the Protocol Summary."""
    out = {}
    for f in matched_forms(spec):
        cs = f["customer_standard"]
        for k in (str(f.get("cdash_domain") or "").upper(), norm_id(f.get("form_id")), _norm_title(f.get("form_title"))):
            if k:
                out.setdefault(k, (cs.get("form_name"), cs.get("source")))
    return out

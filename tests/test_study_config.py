"""Study Configuration (study_config.py): events with calendar, forms at events, every value with a source.
Synthetic spec only."""
import copy
import io
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "skills", "protocol-analysis", "scripts"))
import study_config as sc

PROTOCOL = "Visits: Screening (Day -14 to Day 0), Baseline (Day 0), Day 14 ± 3 after Baseline, and Week 8-10."


def _form(fid, title, visits, **kw):
    return {"form_id": fid, "form_title": title, "visits_assigned": visits, "cdash_domain": kw.pop("domain", None),
            "survey": kw.pop("survey", [{"type": "date", "name": fid + "DAT", "label": "Date"}]), "choices": [],
            "settings": {"form_id": fid, "form_title": title, "version": "1"}, **kw}


def spec():
    return copy.deepcopy({
        "study_meta": {"protocol_number": "SYN-002"},
        "timepoint_csv": {"rows": [
            {"event": "SE_SCREENING", "timepoint": "Screening (Day -14 to Day 0)"},
            {"event": "SE_BASELINE", "timepoint": "Baseline (Day 0)"},
            {"event": "SE_DAY14", "timepoint": "Day 14 ± 3"},
            {"event": "SE_WEEK8", "timepoint": "Week 8-10"},
            {"event": "SE_FU", "timepoint": "Follow-up call"},
            {"event": "SE_UNSCHEDULED", "timepoint": "Unscheduled Visit"},
            {"event": "SE_COMMON", "timepoint": "Common"},
            {"event": "SE_SCREENING", "timepoint": "duplicate row"}]},
        "schedule_of_events": {"form_placements": [
            {"target_visit_oid": "SE_DAY14", "form_id": "VS", "required": False, "notes": "if clinically indicated"}]},
        "forms": [_form("DM", "Demographics", ["SE_SCREENING"]),
                  _form("VS", "Vital Signs", ["SE_SCREENING", "SE_DAY14", "SE_DAY14", "SE_NOT_AN_EVENT"]),
                  _form("AE", "Adverse Events", ["SE_COMMON"], has_repeating_group=True)],
    })


@pytest.fixture(autouse=True)
def _env(monkeypatch):
    monkeypatch.delenv("STUDY_CONFIG", raising=False)


def _ev(cfg, oid):
    return next(e for e in cfg["events"] if e["oid"] == oid)


def test_timing_is_parsed_from_visit_labels():
    assert sc.parse_timing("Day 14 ± 3 after Baseline") == {"offset_days": 14, "window_days": 3, "phrase": "Day 14 ± 3"}
    assert sc.parse_timing("Day 14 +/- 3 days")["window_days"] == 3
    assert sc.parse_timing("Week 4 ± 1 week") == {"offset_days": 28, "window_days": 7, "phrase": "Week 4 ± 1 week"}
    assert sc.parse_timing("Week 2-3") == {"offset_days": 14, "window_days": 7, "phrase": "Week 2-3"}
    assert sc.parse_timing("Month 3")["offset_days"] == 90
    assert sc.parse_timing("Screening (Day -14 to Day 0)") == {"offset_days": -14, "window_days": 14, "phrase": "Day -14 to Day 0"}
    assert sc.parse_timing("Unscheduled Visit") is None


def test_events_type_repeating_and_calendar_with_sources():
    cfg = sc.build(spec(), PROTOCOL)
    assert [e["oid"] for e in cfg["events"]] == ["SE_SCREENING", "SE_BASELINE", "SE_DAY14", "SE_WEEK8", "SE_FU",
                                                 "SE_UNSCHEDULED", "SE_COMMON"]
    assert _ev(cfg, "SE_COMMON")["type"] == {"value": "Common", "source": "pipeline default"}
    assert _ev(cfg, "SE_UNSCHEDULED")["repeating"]["value"] is True and _ev(cfg, "SE_DAY14")["type"]["value"] == "Visit-Based"
    scr = _ev(cfg, "SE_SCREENING")["calendar"]
    assert scr["scheduler"]["trigger"]["value"] == "PARTICIPANT_CREATED" and scr["auto_close_after_days"]["value"] == 14
    d14 = _ev(cfg, "SE_DAY14")["calendar"]
    assert d14["scheduler"]["relative_event"]["value"] == "SE_BASELINE" and d14["scheduler"]["offset_days"]["value"] == 14
    # "Day 14 ± 3" is in the protocol word for word: protocol-sourced, with the quote
    assert d14["auto_close_after_days"] == {"value": 3, "source": "protocol", "quote": "Day 14 ± 3"}
    w8 = _ev(cfg, "SE_WEEK8")["calendar"]
    assert (w8["scheduler"]["offset_days"]["value"], w8["auto_close_after_days"]["value"]) == (56, 14)
    fu = _ev(cfg, "SE_FU")["calendar"]
    assert fu["scheduler"]["trigger"]["value"].startswith("none") and fu["auto_close_after_days"]["value"] is None
    assert _ev(cfg, "SE_COMMON")["calendar"]["auto_close_after_days"]["value"] is None


def test_a_timing_not_found_in_the_protocol_is_ai_proposed():
    cfg = sc.build(spec(), "A protocol that never states the visit days.")
    d14 = _ev(cfg, "SE_DAY14")["calendar"]
    assert d14["scheduler"]["offset_days"] == {"value": 14, "source": "AI-proposed"}


def test_a_timing_is_not_verified_by_a_longer_number_in_the_protocol():
    s = spec()
    s["timepoint_csv"]["rows"] = [{"event": "SE_BASELINE", "timepoint": "Baseline (Day 0)"},
                                  {"event": "SE_W1", "timepoint": "Week 1-2"}, {"event": "SE_D1", "timepoint": "Day 1"}]
    cfg = sc.build(s, "Visits at Week 12-14 and on Day 14.")
    assert _ev(cfg, "SE_W1")["calendar"]["scheduler"]["offset_days"]["source"] == "AI-proposed"
    assert _ev(cfg, "SE_D1")["calendar"]["scheduler"]["offset_days"]["source"] == "AI-proposed"
    cfg = sc.build(s, "Visits at Week 1–2 and on Day 1.")
    assert _ev(cfg, "SE_W1")["calendar"]["scheduler"]["offset_days"]["source"] == "protocol"


def test_the_scheduling_block_wins_over_the_label():
    s = spec()
    s["scheduling"] = [{"event_oid": "SE_SCREENING", "anchor_event_oid": None, "offset_target_days": 0},
                       {"event_oid": "SE_DAY14", "anchor_event_oid": "SE_SCREENING", "offset_target_days": 20,
                        "window_upper_days": 5}]
    d14 = _ev(sc.build(s, PROTOCOL), "SE_DAY14")["calendar"]
    assert d14["scheduler"]["relative_event"]["value"] == "SE_SCREENING" and d14["scheduler"]["offset_days"]["value"] == 20
    assert d14["auto_close_after_days"]["value"] == 5


def test_forms_at_events_follow_the_board_rules():
    cfg = sc.build(spec(), PROTOCOL)
    cards = {(f["event_oid"], f["form_id"]): f for f in cfg["forms"]}
    assert set(cards) == {("SE_SCREENING", "DM"), ("SE_SCREENING", "VS"), ("SE_DAY14", "VS"), ("SE_COMMON", "AE")}
    assert cards[("SE_SCREENING", "DM")]["required"]["value"] is True
    ae = cards[("SE_COMMON", "AE")]
    assert ae["required"]["value"] is False and ae["allow_add"]["value"] is True and ae["hidden"]["value"] is False
    assert cards[("SE_SCREENING", "DM")]["allow_add"]["value"] is False
    # the analysis marked VS optional at Day 14: the board rule is kept, the difference is noted for review
    assert "marks this form optional" in cards[("SE_DAY14", "VS")]["required"]["note"]
    assert all(v["source"] in cfg["sources"] for f in cfg["forms"] for v in (f["required"], f["hidden"], f["allow_add"]))


def test_board_builder_and_configuration_agree():
    src = open(os.path.join(ROOT, "pipeline.py")).read()
    assert '"UNSCH" in event_oid.upper() or "COMMON" in event_oid.upper()' in src  # the rule event_kind mirrors
    assert sc.event_kind("SE_UNSCH_1") == ("Common", True) and sc.event_kind("SE_WEEK_2") == ("Visit-Based", False)
    assert sc.card_required("SE_COMMON") is False and sc.card_required("SE_WEEK_2") is True


def test_dm_values_survive_a_rebuild_and_kill_switch(monkeypatch):
    s = spec()
    cfg = sc.apply(s, PROTOCOL)
    assert s["study_configuration"] is cfg
    next(f for f in cfg["forms"] if f["form_id"] == "DM")["hidden"] = {"value": True, "source": "DM"}
    _ev(cfg, "SE_DAY14")["calendar"]["auto_close_after_days"] = {"value": 10, "source": "DM"}
    again = sc.apply(s, PROTOCOL)
    assert next(f for f in again["forms"] if f["form_id"] == "DM")["hidden"] == {"value": True, "source": "DM"}
    assert _ev(again, "SE_DAY14")["calendar"]["auto_close_after_days"]["value"] == 10
    assert _ev(again, "SE_WEEK8")["calendar"]["auto_close_after_days"]["value"] == 14
    monkeypatch.setenv("STUDY_CONFIG", "0")
    s2 = spec()
    assert sc.apply(s2, PROTOCOL) is None and "study_configuration" not in s2


def test_xlsx_sheet_and_pdf_section():
    import openpyxl
    from generate_study_spec_xlsx import build_study_config_sheet
    from generate_study_spec_pdf import build_study_config_page, make_styles
    s = spec()
    sc.apply(s, PROTOCOL)
    wb = openpyxl.Workbook()
    build_study_config_sheet(wb, s)
    ws = wb["STUDY_CONFIG"]
    cells = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    flat = [str(c) for r in cells for c in r if c is not None]
    assert any(c.startswith("EVENTS") for c in flat) and any(c.startswith("FORMS AT EVENTS") for c in flat)
    row = next(r for r in cells if r[0] == "SE_DAY14" and r[2] == "Visit-Based")
    assert row[5] == "SE_BASELINE" and row[6] == "14" and row[7] == "3" and 'protocol: "Day 14 ± 3"' in row[8]
    assert build_study_config_page(s, make_styles())
    # no configuration: no sheet, no section
    wb2 = openpyxl.Workbook()
    build_study_config_sheet(wb2, spec())
    assert "STUDY_CONFIG" not in wb2.sheetnames and build_study_config_page(spec(), make_styles()) == []


def test_summary_line():
    line = sc.summary_line(sc.build(spec(), PROTOCOL))
    assert line.startswith("Study Configuration: 7 event(s) (2 scheduled relative to another event, 3 with auto-close), 4 form(s)")

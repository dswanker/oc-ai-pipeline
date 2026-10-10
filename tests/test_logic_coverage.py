"""The logic-coverage skill (skills/logic-coverage): the catalog, build-or-propose, cross-form checks, derived helper
items, the report. All forms are synthetic; CDASH concepts are set on the rows as cdisc_concepts.py would."""
import contextlib, copy, io, json, os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
for p in (ROOT, HERE, os.path.join(ROOT, "skills", "logic-coverage", "scripts"),
          os.path.join(ROOT, "skills", "dvs-specification", "scripts")):
    if p not in sys.path:
        sys.path.insert(0, p)
import openpyxl

import logic_coverage as lc
import standalone
import standards_match as sm
import dvs_edits
from standards import fixtures as fx

NY = [{"list_name": "NY", "name": "Y", "label": "Yes"}, {"list_name": "NY", "name": "N", "label": "No"}]


def row(type_, name, label=None, concept=None, qual=None, group="G", **kw):
    r = {"type": type_, "name": name, "label": label if label is not None else f"Question {name}",
         "bind__oc_itemgroup": group, **kw}
    if concept:
        r["concept"] = concept
    if qual:
        r["concept_qualifier"] = qual
    return r


def form(fid, rows, choices=None, standard=None, domain=None, visits=("SE_ONE",), title=None):
    f = {"form_id": fid, "form_title": title or f"Form {fid}", "cdash_domain": domain, "survey": rows,
         "choices": NY + (choices or []), "settings": {"form_id": fid, "form_title": title or f"Form {fid}"},
         "visits_assigned": list(visits)}
    if standard:                       # "odm": structure only; "xls": an XLSForm that carries its own logic
        f["customer_standard"] = {"source": "uploaded ODM" if standard == "odm" else "uploaded XLSForm",
                                  "has_logic": standard == "xls", "verbatim": standard == "xls", "approved": [],
                                  "added_fields": []}
        for r in rows:
            r.setdefault("library_source", "CUSTOM")
    return f


def consent():
    return form("ICF", [row("date", "ICFDAT", concept="RFICDAT", group="ICF")])


def spec(*forms):
    return {"study_meta": {"protocol_number": "T"}, "forms": list(forms), "review_flags": {}}


def run(s, mode="apply"):
    with contextlib.redirect_stdout(io.StringIO()):
        return lc.run(s, mode)


def r_(s, fid, name):
    return next(r for f in s["forms"] if f["form_id"] == fid for r in f["survey"] if r.get("name") == name)


def total(res, cat):
    return res["totals"][cat]


# ── field level ───────────────────────────────────────────────────────────────

def test_dates_cannot_be_in_the_future_unless_the_field_asks_for_a_future_date():
    s = spec(form("AA", [row("date", "D1"), row("datetime", "D2"),
                         row("date", "PLANNED", constraint=". >= today()"),
                         row("date", "HAS", constraint=". <= today()"),
                         row("date", "CALC", calculation="today()", readonly="yes")]))
    res = run(s)
    assert r_(s, "AA", "D1")["constraint"] == ". <= today()"
    assert r_(s, "AA", "D1")["constraint_message"] == "Date cannot be in the future."
    assert r_(s, "AA", "D2")["constraint"] == ". <= now()"
    assert r_(s, "AA", "PLANNED")["constraint"] == ". >= today()"            # not applicable
    assert r_(s, "AA", "HAS")["constraint"] == ". <= today()"                # covered: not doubled
    assert "constraint" not in r_(s, "AA", "CALC")
    assert total(res, "DATE_FUTURE") == {"applicable": 3, "covered": 1, "proposed_before": 0, "added": 2,
                                         "proposed": 0, "missing": 0}


def test_partial_date_parts_get_ranges():
    s = spec(form("AA", [row("integer", "ST_YEAR"), row("integer", "ST_MON"), row("integer", "ST_DAY"),
                         row("integer", "LONE_YEAR")]))
    run(s)
    assert r_(s, "AA", "ST_YEAR")["constraint"] == ". >= 1900 and . <= number(format-date(today(), '%Y'))"
    assert r_(s, "AA", "ST_MON")["constraint"] == ". >= 1 and . <= 12"
    assert r_(s, "AA", "ST_DAY")["constraint"] == ". >= 1 and . <= 31"
    assert r_(s, "AA", "LONE_YEAR")["constraint"] == ". >= 0"                # no sibling part: an ordinary count


def test_numeric_ranges_by_concept_and_unit():
    units = [{"list_name": "TU", "name": "C", "label": "C"}, {"list_name": "TU", "name": "F", "label": "F"},
             {"list_name": "WU", "name": "kg", "label": "kg"}]
    s = spec(form("VS", [row("decimal", "SYS", concept="VSORRES", qual="SYSBP"),
                         row("decimal", "TEMP", concept="VSORRES", qual="TEMP"),
                         row("select_one TU", "TEMPU", concept="VSORRESU", qual="TEMP"),
                         row("decimal", "WT", concept="VSORRES", qual="WEIGHT"),
                         row("select_one WU", "WTU", concept="VSORRESU", qual="WEIGHT"),
                         row("integer", "AGE", concept="AGE"), row("integer", "COUNT"),
                         row("decimal", "LAB", concept="LBORRES", qual="WBC"), row("decimal", "MYSTERY"),
                         row("decimal", "DONE", constraint=". > 0 and . <= 42")], units, domain="VS"))
    res = run(s)
    assert r_(s, "VS", "SYS")["constraint"] == ". >= 60 and . <= 250"
    assert r_(s, "VS", "TEMP")["constraint"] == \
        "(${TEMPU} = 'F' and . >= 93 and . <= 108) or (${TEMPU} != 'F' and . >= 34 and . <= 42)"
    assert r_(s, "VS", "WT")["constraint"] == ". >= 20 and . <= 300"         # the form offers kg only
    assert r_(s, "VS", "AGE")["constraint"] == ". >= 0 and . <= 120"
    assert r_(s, "VS", "COUNT")["constraint"] == ". >= 0" and r_(s, "VS", "LAB")["constraint"] == ". >= 0"
    assert "constraint" not in r_(s, "VS", "MYSTERY")                        # meaning unknown: no invented range
    assert r_(s, "VS", "DONE")["constraint"] == ". > 0 and . <= 42"
    assert [n["field"] for n in res["not_generated"]] == ["MYSTERY"]
    assert total(res, "NUM_RANGE")["missing"] == 1 and total(res, "NUM_RANGE")["covered"] == 1


def test_other_specify_is_shown_only_with_other():
    ch = [{"list_name": "RT", "name": "ORAL", "label": "Oral"}, {"list_name": "RT", "name": "c_99", "label": "Other",
                                                                  "cdisc_submission_value": "OTHER"},
          {"list_name": "MS", "name": "A", "label": "A"}, {"list_name": "MS", "name": "OTHER", "label": "Other"}]
    s = spec(form("CM", [row("select_one RT", "ROUTE"), row("text", "ROUTEO"),
                         row("select_multiple MS", "MULTI"), row("text", "MULTI_SP"),
                         row("select_one NY", "NOOTHER"), row("text", "NOOTHERX")], ch))
    run(s)
    assert r_(s, "CM", "ROUTEO")["relevant"] == "${ROUTE} = 'c_99'"          # the list's own code for OTHER
    assert r_(s, "CM", "MULTI_SP")["relevant"] == "selected(${MULTI}, 'OTHER')"
    assert "relevant" not in r_(s, "CM", "NOOTHERX")


def test_gate_hides_only_items_of_known_meaning_and_reason_shows_when_not_done():
    s = spec(form("LB", [row("select_one NY", "PERF", concept="LBPERF"),
                         row("text", "WHYNOT", concept="LBREASND"),
                         row("date", "LBDAT", concept="LBDAT"), row("decimal", "RES", concept="LBORRES"),
                         row("text", "UNKNOWN"),                                   # no concept: never hidden
                         row("text", "ELSEWHERE", concept="LBNAM", group="OTHER"),  # another item group
                         row("text", "OWN", concept="LBSPEC", relevant="${RES} > 1")], domain="LB"))
    res = run(s)
    assert r_(s, "LB", "LBDAT")["relevant"] == "${PERF} = 'Y'" and r_(s, "LB", "RES")["relevant"] == "${PERF} = 'Y'"
    assert "relevant" not in r_(s, "LB", "UNKNOWN") and "relevant" not in r_(s, "LB", "ELSEWHERE")
    assert r_(s, "LB", "OWN")["relevant"] == "${RES} > 1"                    # the author's show-when stays
    why = r_(s, "LB", "WHYNOT")
    assert why["relevant"] == "${PERF} = 'N'" and why["required"] == "yes"
    assert why["edit_check_source"]["relevant"] == "CDISC CORE (CORE-000440)"
    assert total(res, "GATE")["added"] == 2 and total(res, "GATE")["covered"] == 1
    assert total(res, "NOT_DONE_REASON")["added"] == 1


def test_dose_gate_names_its_core_rule():
    s = spec(form("EX", [row("select_one NY", "GIVEN", concept="EXOCCUR"), row("decimal", "DOSE", concept="EXDOSE")],
                  domain="EX"))
    res = run(s)
    assert r_(s, "EX", "DOSE")["relevant"] == "${GIVEN} = 'Y'"
    assert r_(s, "EX", "DOSE")["edit_check_source"]["relevant"] == "CDISC CORE (CORE-000004/137)"
    assert res["core_rules"] == ["CORE-000004/137"]


# ── within a form ─────────────────────────────────────────────────────────────

def _ae():
    out = [{"list_name": "OUT", "name": "REC", "label": "Recovered"},
           {"list_name": "OUT", "name": "c_5", "label": "Fatal", "cdisc_submission_value": "FATAL"}]
    return form("AE", [row("date", "START", concept="AESTDAT"), row("select_one NY", "ONGO", concept="AEONGO"),
                       row("date", "END", concept="AEENDAT"),
                       row("select_one NY", "SER", concept="AESER"), row("select_one NY", "DTH", concept="AESDTH"),
                       row("select_one NY", "HOSP", concept="AESHOSP"), row("select_one OUT", "OUTC", concept="AEOUT"),
                       row("select_one NY", "SIG", concept="PECLSIG"), row("text", "SIGTXT", concept="PEDESC")],
                out, domain="AE", visits=("SE_COMMON",))


def test_within_form_consistency_checks():
    s = spec(_ae())
    res = run(s)
    end = r_(s, "AE", "END")
    assert ". = '' or ${START} = '' or . >= ${START}" in end["constraint"] and ". <= today()" in end["constraint"]
    assert end["relevant"] == "${ONGO} = 'N'" and end["required"] == "yes"
    assert r_(s, "AE", "SER")["constraint"] == ". != 'Y' or (${DTH} = 'Y' or ${HOSP} = 'Y')"
    assert r_(s, "AE", "OUTC")["constraint"] == ". != 'c_5' or ${DTH} = 'Y'"
    assert r_(s, "AE", "SIGTXT")["relevant"] == "${SIG} = 'Y'" and r_(s, "AE", "SIGTXT")["required"] == "yes"
    for cat in ("START_END", "ONGOING_END", "SERIOUS_CRITERIA", "FATAL_DEATH", "CLSIG_DESC"):
        assert total(res, cat)["added"] == 1, cat
    assert "CORE-000022" in res["core_rules"]
    details = {d["id"]: d for d in r_(s, "AE", "SER")["edit_check_details"]}
    assert details["LC.CORE-000022"]["source"] == "CDISC CORE (CORE-000022)"


def test_existing_logic_is_never_duplicated():
    ae = _ae()
    r = {x["name"]: x for x in ae["survey"]}
    r["END"]["constraint"] = ". >= ${START}"
    r["END"]["relevant"] = "${ONGO} = 'N'"
    r["SER"]["constraint"] = ". = 'N' or ${DTH} = 'Y' or ${HOSP} = 'Y'"
    s = spec(ae)
    res = run(s)
    assert r_(s, "AE", "END")["constraint"] == "(. >= ${START}) and (. <= today())"   # only the missing rule added
    assert r_(s, "AE", "SER")["constraint"] == ". = 'N' or ${DTH} = 'Y' or ${HOSP} = 'Y'"
    for cat in ("START_END", "ONGOING_END", "SERIOUS_CRITERIA"):
        assert total(res, cat) == {"applicable": 1, "covered": 1, "proposed_before": 0, "added": 0, "proposed": 0,
                                   "missing": 0}, cat


def test_a_second_run_changes_nothing_and_reports_the_same():
    s = spec(consent(), _ae())
    first = copy.deepcopy(run(s))
    forms = copy.deepcopy(s["forms"])
    second = run(s)
    assert s["forms"] == forms and second["rows"] == first["rows"] and second["counts"] == first["counts"]


# ── cross form ────────────────────────────────────────────────────────────────

def test_cross_form_checks_fetch_from_the_right_form_and_item():
    ex = form("EX", [row("date", "DOSEDT", concept="EXSTDAT", group="EX")], domain="EX", visits=("SE_DAY1", "SE_DAY8"))
    ds = form("DS", [row("date", "DSDT", concept="DSSTDAT"), row("date", "DIED", concept="DTHDAT")], domain="DS")
    dd = form("DD", [row("date", "DEATHDT", concept="DTHDAT", group="DD")], domain="DD")
    mh = form("MH", [row("date", "MHST", concept="MHSTDAT")], domain="MH")
    mi = form("MI", [row("date", "BIOPSY", concept="MIDAT")], domain="MI")
    s = spec(consent(), ex, ds, dd, mh, mi, _ae())
    res = run(s)
    # consent floor on on-study dates, with the hidden fetch of ICF.ICFDAT
    dsdt = r_(s, "DS", "DSDT")
    assert ". = '' or ${ICFDAT_CF} = '' or . >= ${ICFDAT_CF}" in dsdt["constraint"]
    cf = r_(s, "DS", "ICFDAT_CF")
    assert cf["type"] == "calculate" and "@FormOID='F_ICF'" in cf["calculation"] and "ItemName='ICFDAT'" in cf["calculation"]
    assert "ICFDAT_CF" in r_(s, "EX", "DOSEDT")["constraint"] and "ICFDAT_CF" in r_(s, "AE", "START")["constraint"]
    # death date equal across forms (CORE-000034), both directions
    assert ". = '' or ${DTHDAT_CF} = '' or . = ${DTHDAT_CF}" in r_(s, "DS", "DIED")["constraint"]
    assert "@FormOID='F_DD'" in r_(s, "DS", "DTHDAT_CF")["calculation"] and "ItemName='DEATHDT'" in r_(s, "DS", "DTHDAT_CF")["calculation"]
    assert "@FormOID='F_DS'" in r_(s, "DD", "DTHDAT_CF")["calculation"]
    # review checks are proposals, each with what to confirm; the forms are untouched by them
    props = {(p["target_form"], p["category"]): p for p in res["proposals"]}
    assert set(props) == {("MH", "HISTORY_BEFORE_CONSENT"), ("MI", "CONSENT_FLOOR"), ("AE", "AE_FIRST_DOSE")}
    assert "constraint" not in r_(s, "MH", "MHST") or "ICFDAT_CF" not in r_(s, "MH", "MHST")["constraint"]
    assert "FIRSTDOSE_CF" not in json.dumps(next(f for f in s["forms"] if f["form_id"] == "AE")["survey"])
    first_dose = props[("AE", "AE_FIRST_DOSE")]
    assert first_dose["cross_form"] == "EX.DOSEDT" and "pre-treatment" in first_dose["message"]
    helper = next(o["row"] for o in first_dose["ops"] if o["op"] == "insert")
    assert "[@StudyEventOID='SE_DAY1']" in helper["calculation"]                # the first dosing visit
    assert "confirm against the protocol" in props[("MI", "CONSENT_FLOOR")]["note"]
    listed = {(x["form"], x["field"], x.get("status", "added")) for x in res["cross_form"]}
    assert ("DS", "DSDT", "added") in listed and ("AE", "START", "proposed") in listed
    assert "CORE-000034" in res["core_rules"]
    # dosing window: applies, is not generated, and says why
    assert total(res, "DOSING_WINDOW") == {"applicable": 1, "covered": 0, "proposed_before": 0, "added": 0,
                                           "proposed": 0, "missing": 1}
    assert any(n["category"] == "DOSING_WINDOW" and "no visit windows" in n["reason"] for n in res["not_generated"])


def test_eligibility_and_pregnancy_checks_are_review_proposals():
    ie = form("IE", [row("select_one NY", "ELIG", concept="IEYN", group="IE")], domain="IE")
    ex = form("EX", [row("date", "DOSEDT", concept="EXSTDAT", group="EX")], domain="EX")
    rp = form("RP", [row("date", "LMPDT", concept="RPSTDAT")], domain="RP")
    res = run(spec(ie, ex, rp))
    cats = {(p["target_form"], p["category"]) for p in res["proposals"]}
    assert ("EX", "ELIGIBILITY_DOSING") in cats and ("RP", "PREGNANCY_DOSING") in cats
    elig = next(p for p in res["proposals"] if p["category"] == "ELIGIBILITY_DOSING")
    assert elig["logic"] == ". = '' or ${IEYN_CF} = '' or ${IEYN_CF} = 'Y'" and elig["cross_form"] == "IE.ELIG"


def test_required_from_cdash_is_a_proposal():
    s = spec(form("AA", [row("text", "TERM", concept="AETERM", cdash={"core": "HR", "variable": "AETERM"}),
                         row("text", "OPT", concept="AEX", cdash={"core": "O", "variable": "AEX"}),
                         row("text", "ISREQ", concept="AEY", cdash={"core": "HR"}, required="yes")]))
    res = run(s)
    assert "required" not in r_(s, "AA", "TERM")
    assert [(p["target_field"], p["check_type"]) for p in res["proposals"]] == [("TERM", "Required")]
    assert total(res, "REQUIRED")["covered"] == 1


# ── build or propose ──────────────────────────────────────────────────────────

def _std_ae(kind):
    f = _ae()
    return form("AE", f["survey"], f["choices"], standard=kind, domain="AE")


def test_a_logic_free_standard_gets_checks_built_with_the_reason():
    s = spec(_std_ae("odm"))
    res = run(s)
    assert res["counts"]["applied"] >= 6 and not res["proposals"]
    assert r_(s, "AE", "END")["relevant"] == "${ONGO} = 'N'"
    approved = s["forms"][0]["customer_standard"]["approved"]
    assert approved and all(a["auto"].startswith("Applied: the customer standard form carries no logic") for a in approved)


def test_a_standard_with_its_own_logic_gets_proposals_only():
    f = _std_ae("xls")
    pristine = copy.deepcopy(f["survey"])
    s = spec(f)
    res = run(s)
    assert s["forms"][0]["survey"] == pristine and res["counts"]["applied"] == 0
    assert len(res["proposals"]) >= 6
    assert all(p["kind"] == "coverage" and p["id"].startswith("LCP-") and p["ops"] and p["source"] for p in res["proposals"])
    assert all("carries its own logic" in p["note"] for p in res["proposals"])
    assert sm.proposals(s) == res["proposals"]                                # a data manager sees them in the DVS


def test_kill_switch_makes_a_logic_free_standard_proposals_only(monkeypatch):
    monkeypatch.setenv("STANDARD_LOGIC_FREE_APPLY", "0")
    res = run(spec(_std_ae("odm")))
    assert res["counts"]["applied"] == 0 and res["proposals"]


def test_report_mode_changes_no_form():
    s = spec(consent(), _ae(), _std_ae("odm") | {"form_id": "AE2"})
    before = copy.deepcopy(s["forms"])
    res = run(s, "report")
    assert s["forms"] == before and res["counts"]["applied"] == 0 and res["mode"] == "report"
    assert len(res["proposals"]) >= 10


def test_approve_builds_a_proposal_and_reject_retires_it():
    s = spec(_std_ae("xls"))
    res = run(s)
    end_order = next(p for p in res["proposals"] if p["category"] == "START_END")
    gate = next(p for p in res["proposals"] if p["category"] == "ONGOING_END")
    assert sm.apply_proposal(s, end_order)[0] == "applied"
    assert "${START}" in r_(s, "AE", "END")["constraint"]
    assert sm.reject_proposal(s, gate["id"])
    again = run(s)
    ids = {p["id"] for p in again["proposals"]}
    assert end_order["id"] not in ids and gate["id"] not in ids
    assert total(again, "START_END")["added"] == 1                            # the approved check counts as added
    assert gate["id"] in again["rejected_ids"]


def test_approving_on_a_pipeline_form_does_not_turn_it_into_a_customer_standard():
    s = spec(consent(), form("EX", [row("date", "DOSEDT", concept="EXSTDAT", group="EX")], domain="EX"), _ae())
    res = run(s)
    p = next(x for x in res["proposals"] if x["category"] == "AE_FIRST_DOSE")
    assert sm.apply_proposal(s, p)[0] == "applied"
    ae = next(f for f in s["forms"] if f["form_id"] == "AE")
    assert "customer_standard" not in ae and "FIRSTDOSE_CF" in r_(s, "AE", "START")["constraint"]


# ── derived helper items ──────────────────────────────────────────────────────

def test_derived_helper_items_are_rebuilt_or_made_read_only_and_reported():
    donor = form("DOV", [row("calculate", "EVT_CF", label="", calculation="instance('clinicaldata')/ODM/@X",
                             bind__oc_external="clinicaldata"),
                         row("calculate", "TPT", label="", calculation="pulldata('t','timepoint','event',${EVT_CF})")])
    odm = form("VS", [row("text", "TPT", label="TPT"),                    # calculated elsewhere in the study
                      row("text", "X_CALC", label=""),                    # nothing to derive it from
                      row("text", "REQD", label="REQD", required="yes"),  # a user must answer it: not a helper
                      row("text", "REAL", label="A real question"),
                      row("decimal", "USES", calculation="${X_CALC} + 1", readonly="yes")], standard="odm")
    s = spec(donor, odm)
    res = run(s)
    tpt = r_(s, "VS", "TPT")
    assert tpt["calculation"] == "pulldata('t','timepoint','event',${EVT_CF})" and tpt["readonly"] == "yes"
    assert r_(s, "VS", "EVT_CF")["calculation"] == "instance('clinicaldata')/ODM/@X"   # its helper came along
    x = r_(s, "VS", "X_CALC")
    assert x["readonly"] == "yes" and "calculation" not in x and x["completion_status"] == "FLAGGED"
    assert "readonly" not in r_(s, "VS", "REQD") and "readonly" not in r_(s, "VS", "REAL")
    got = {h["item"]: h for h in res["helpers"]}
    assert set(got) == {"TPT", "X_CALC"}
    assert got["TPT"]["action"] == "calculation rebuilt" and got["X_CALC"]["action"] == "made read-only"
    assert "referenced by other logic" in got["X_CALC"]["detail"]
    assert total(res, "DERIVED_ITEM")["added"] == 2
    assert {h["item"] for h in run(s)["helpers"]} == {"TPT", "X_CALC"}        # still reported on a later run


def test_helper_switch_reports_them_and_changes_none(monkeypatch):
    monkeypatch.setenv("LOGIC_COVERAGE_HELPERS", "0")
    s = spec(form("VS", [row("text", "X_CALC", label="X_CALC"), row("date", "D1")], standard="odm"))
    res = run(s)
    assert "readonly" not in r_(s, "VS", "X_CALC") and res["helpers"][0]["action"].startswith("reported")
    assert r_(s, "VS", "D1")["constraint"] == ". <= today()"          # the checks are still built


def test_helper_items_are_only_reported_in_report_mode_and_never_on_forms_with_logic():
    odm = form("VS", [row("text", "X_CALC", label="X_CALC")], standard="odm")
    xls = form("PE", [row("text", "Y_CALC", label="Y_CALC")], standard="xls")
    s = spec(odm, xls)
    res = run(s, "report")
    assert "readonly" not in r_(s, "VS", "X_CALC")
    assert [(h["item"], h["action"].split(" (")[0]) for h in res["helpers"]] == [("X_CALC", "reported")]


# ── report, DVS, standalone ───────────────────────────────────────────────────

def test_report_has_a_row_per_form_and_category_and_a_specification_section():
    s = spec(consent(), _ae())
    res = run(s)
    assert len(res["rows"]) == 2 * len(lc.CATEGORIES)
    sheet = lc.sheet_rows(s)
    icf = next(r for r in sheet if r[0] == "ICF" and r[2] == "End on or after start")
    assert icf[3] == 0 and icf[9] == "Not applicable"
    ae = next(r for r in sheet if r[0] == "AE" and r[2] == "End on or after start")
    assert ae[3:9] == [1, 0, 0, 1, 0, 0] and ae[9] == "1 added" and ae[10] == "END"
    title, _note, headers, rows, weights = lc.section(s)
    assert title == "LOGIC COVERAGE" and len(headers) == len(weights) == 8
    assert all(r[2] > 0 for r in rows)
    assert lc.summary_lines(s)[0].startswith("Logic coverage (apply): ")
    assert lc.section({"study_meta": {}}) is None


def _dvs(s):
    from extract_dvs_from_forms import extract_dvs_data
    from generate_dvs import build_dvs
    fj = standalone.original_forms_json(s)
    path = os.path.join(tempfile.mkdtemp(), "dvs.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        data = extract_dvs_data(s, fj)
        build_dvs(data, path)
    return data, path


def test_dvs_lists_added_checks_as_draft_proposals_as_proposed_and_has_the_coverage_sheet():
    s = spec(consent(), _std_ae("odm"), _std_ae("xls") | {"form_id": "AEX"})
    run(s)
    data, path = _dvs(s)
    rows = data["dvs_oc4"]
    added = [r for r in rows if r["Target Form OID"] == "AE" and r["Target Item Name"] == "SER" and r["Check Type"] == "Constraint"]
    assert added and added[0]["Status"] == "Draft" and added[0]["Check Source"] == "CDISC CORE (CORE-000022)"
    assert "carries no logic of its own" in added[0]["Notes"]
    floor = next(r for r in rows if r["Target Form OID"] == "AE" and r["Target Item Name"] == "START" and "ICFDAT_CF" in r["Expression / Calculation"])
    assert floor["Check Source"] == "Logic Coverage" and floor["Source Form OID(s)"] in ("ICF", "F_ICF")
    proposed = [r for r in rows if r.get("Status") == "Proposed" and r["Target Form OID"] == "AEX"]
    assert proposed and all("(proposed, logic coverage)" in r["Check Name"] for r in proposed)
    assert all(json.loads(r["Machine Data"])["kind"] == "standard_proposal" for r in proposed)
    wb = openpyxl.load_workbook(path)
    ws = wb["LOGIC_COVERAGE"]
    head = [c.value for c in ws[2]]
    assert head == lc.SHEET_HEADERS
    assert ws.max_row >= 2 + 3 * len(lc.CATEGORIES)
    assert any(c.get("Item_Name") == "END" for c in data["uat_cases"])        # UAT cases from the added logic


def test_dvs_approve_action_builds_a_coverage_proposal():
    s = spec(_std_ae("xls"))
    run(s)
    data, _path = _dvs(s)
    target = next(r for r in data["dvs_oc4"] if r.get("Status") == "Proposed" and "${START}" in str(r["Expression / Calculation"]))
    res = dvs_edits.apply_actions(s, [{**{k: "" if v is None else str(v) for k, v in target.items()},
                                        "Action": "approve"}])
    assert res["applied"] == 1 and "${START}" in r_(s, "AE", "END")["constraint"]


def test_standalone_on_an_odm_builds_logic_and_on_an_xlsform_with_logic_proposes():
    out = tempfile.mkdtemp()
    res = standalone.run_existing([("standard.xml", fx.ODM)], "apply", "T-1", out)
    s = res["spec"]
    assert res["build_log"]["forms_built"] and not res["build_log"]["build_errors"]
    assert res["summary"]["counts"]["applied"] > 0 and os.path.exists(res["paths"]["dvs"])
    built = openpyxl.load_workbook(os.path.join(res["paths"]["forms"], "AE.xlsx"))["survey"]
    head = [c.value for c in built[1]]
    cons = [r[head.index("constraint")].value for r in built.iter_rows(min_row=2) if r[head.index("name")].value == "AESTDAT"]
    assert cons and "today()" in cons[0]
    assert "# Logic coverage report" in open(res["paths"]["report"]).read()
    assert all(f["customer_standard"]["source"] == "existing build" for f in s["forms"])

    res2 = standalone.run_existing([("AEGEN.xlsx", fx.xlsform_bytes())], "apply", "T-2", tempfile.mkdtemp())
    assert res2["summary"]["counts"]["applied"] == 0 and res2["summary"]["counts"]["proposed"] > 0

    res3 = standalone.run_existing([("standard.xml", fx.ODM)], "report", "T-3", tempfile.mkdtemp())
    assert res3["summary"]["counts"]["applied"] == 0 and "forms" not in res3["paths"]
    assert res3["dvs"]["uat_cases"]


def test_odm_events_place_forms_on_their_visits():
    data = fx.ODM.replace(b'<MetaDataVersion OID="v1" Name="v1">',
                          b'<MetaDataVersion OID="v1" Name="v1"><StudyEventDef OID="SE_A" Name="A" Repeating="No" '
                          b'Type="Scheduled"><FormRef FormOID="F_AE" Mandatory="No"/></StudyEventDef>')
    assert standalone.odm_form_events(data) == {"F_AE": ["SE_A"]}
    s = standalone.study_from_files([("s.xml", data)], "T")
    assert next(f for f in s["forms"] if f["form_id"] == "AE")["visits_assigned"] == ["SE_A"]


def test_disabled_by_its_switch(monkeypatch):
    assert lc.enabled()
    monkeypatch.setenv("LOGIC_COVERAGE", "0")
    assert not lc.enabled()


# ── in the pipeline ───────────────────────────────────────────────────────────

def test_the_pipeline_runs_the_audit_after_the_logic_is_assembled_and_before_the_chains():
    src = open(os.path.join(ROOT, "pipeline.py")).read()
    step = src.index("async def _logic_coverage_step(")
    body = src[step:src.index("\nasync def ", step + 10)]
    assert "_lc.enabled()" in body and "except Exception" in body and "build continues" in body
    call = src.index("await _logic_coverage_step(item_id, struct_json)")
    assert src.index("await _propose_ai_edit_checks(item_id, struct_json)", call - 2000) < call
    assert call < src.index("# ── Chain A: Study Spec files", call)


def test_study_specification_generators_render_the_section():
    for name in ("generate_study_spec_xlsx.py", "generate_study_spec_pdf.py"):
        text = open(os.path.join(ROOT, "skills", "protocol-analysis", "scripts", name)).read()
        assert "logic_coverage.section(data)" in text

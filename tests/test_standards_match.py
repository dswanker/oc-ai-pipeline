"""Customer standard form matching (standards_match.py; docs/OC_STANDARD_FORM_MATCHING_PLAN.md, agreed rules).
All fixtures are synthetic (tests/standards/fixtures.py)."""
import contextlib, copy, io, json, os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "..", "skills", "dvs-specification", "scripts"))
sys.path.insert(0, os.path.join(HERE, "..", "skills", "edc-builder", "scripts"))
sys.path.insert(0, os.path.join(HERE, "..", "skills", "protocol-analysis", "scripts"))
import openpyxl
import pytest

import standards_match as sm
import dvs_edits
from conventions_engine import apply_conventions
from standards import fixtures as fx


def _sources(xls=True, odm=True, referenced=None):
    files = ([("AEGEN.xlsx", fx.xlsform_bytes())] if xls else []) + ([("standard.xml", fx.ODM)] if odm else [])
    return sm.load_sources(files, referenced)


def _matched(**kw):
    with contextlib.redirect_stdout(io.StringIO()):
        return sm.apply(fx.spec(), _sources(**kw))


def _form(spec, fid):
    return next(f for f in spec["forms"] if f["form_id"] == fid)


def _engine(spec):
    with contextlib.redirect_stdout(io.StringIO()):
        apply_conventions(spec, study_id="T", customer_subdomain="", client_name="")
    return spec


def _standard_rows():
    return sm.parse_xlsform(fx.xlsform_bytes(), "AEGEN.xlsx")["survey"]


# ── detection: every file is used or named ───────────────────────────────────────

def test_detects_single_xlsx_docx_zip_and_odm_by_content():
    assert sm.detect_kind(fx.xlsform_bytes()) == "XLSFORM"
    assert sm.detect_kind(fx.xlsform_bytes(with_survey=False)) == "UNSTRUCTURED"  # a workbook that is not an XLSForm
    assert sm.detect_kind(fx.docx_bytes()) == "UNSTRUCTURED"                      # .docx is a zip container too
    assert sm.detect_kind(fx.zip_bytes({"forms/AEGEN.xlsx": fx.xlsform_bytes()})) == "XLSFORM_ZIP"
    assert sm.detect_kind(fx.ODM) == "ODM_XML"
    assert sm.detect_kind(b"just some notes") == "UNSTRUCTURED"
    assert sm.detect_kind(b"") == "UNSTRUCTURED"


def test_every_file_is_used_or_named_as_not_usable():
    src = sm.load_sources([("AEGEN.xlsx", fx.xlsform_bytes()), ("notes.docx", fx.docx_bytes()),
                           ("lists.xlsx", fx.xlsform_bytes(with_survey=False)), ("broken.zip", b"PK\x03\x04nope"),
                           ("pack.zip", fx.zip_bytes({"VS.xlsx": fx.xlsform_bytes(fx.VS_SURVEY, [], fx.VS_SETTINGS),
                                                      "readme.txt": "hello", "__MACOSX/._VS.xlsx": "x"}))])
    by = {r["file"]: r for r in src["files"]}
    assert by["AEGEN.xlsx"]["usable"] and by["pack.zip"]["usable"] and by["pack.zip"]["forms"] == 1
    assert "readme.txt" in by["pack.zip"]["note"]
    for name in ("notes.docx", "lists.xlsx", "broken.zip"):
        assert not by[name]["usable"] and by[name]["note"]
    log = " | ".join(sm.files_log(src))
    assert all(n in log for n in by) and "NOT usable" in log
    assert {f["form_oid"] for f in src["forms"]} == {"AEGEN", "VS"}


# ── parsers ──────────────────────────────────────────────────────────────────────

def test_xlsform_parser_keeps_logic_columns_and_settings():
    f = sm.parse_xlsform(fx.xlsform_bytes(), "AEGEN.xlsx")
    rows = {r.get("name"): r for r in f["survey"] if r.get("name")}
    assert f["form_oid"] == "AEGEN" and f["title"] == "AE General" and f["has_logic"] and f["format"] == "xlsform"
    assert rows["AESTDAT"]["constraint"] == ". <= today()"
    assert rows["AESTDAT"]["constraint_message"] == "Start date cannot be in the future!"
    assert rows["AEENDAT"]["relevant"] == "${AEONGO} = 'N'"
    assert rows["AETERM"]["bind__oc_itemgroup"] == "AEG" and rows["AETERM"]["oc:custom_col"] == "keep-me"
    assert rows["EVT_CF"]["bind__oc_external"] == "clinicaldata"
    assert f["extra_cols"] == ["oc:custom_col"]
    assert f["settings"]["version"] == "7" and len(f["choices"]) == 5
    assert [r["type"] for r in f["survey"]][1] == "begin group" and f["survey"][-1] == {"type": "end group"}


def test_odm_parser_structure_types_lists_and_no_logic():
    forms = {f["form_oid"]: f for f in sm.parse_odm(fx.ODM, "standard.xml")}
    assert set(forms) == {"AE", "MEDH"}
    ae = forms["AE"]
    assert ae["has_logic"] is False and ae["format"] == "odm" and ae["title"] == "Adverse Events"
    assert [r["name"] for r in ae["survey"]] == ["AETERM", "AESTDAT", "AEENDAT", "AESEV", "AEACN", "AEWT", "AENOQ", "AEDESC", "AELEFT"]  # OrderNumber
    rows = {r["name"]: r for r in ae["survey"]}
    assert rows["AESTDAT"]["type"] == "date" and rows["AEWT"]["type"] == "decimal"
    assert rows["AESEV"]["type"] == "select_one AESEV" and rows["AEACN"]["type"] == "select_multiple AEACN_2" \
        or rows["AEACN"]["type"] == "select_multiple AEACN"
    assert rows["AETERM"]["required"] == "yes" and "required" not in rows["AEENDAT"]
    assert rows["AETERM"]["label"] == "Event term" and rows["AETERM"]["bind__oc_itemgroup"] == "AE"
    assert {(c["list_name"], c["name"], c["label"]) for c in ae["choices"]} >= {("AESEV", "1", "Mild"), ("AESEV", "2", "Moderate")}
    assert any(c["name"] == "DRUG" and c["label"] == "Drug withdrawn" for c in ae["choices"])
    assert forms["MEDH"]["repeating"] is True


def test_domain_from_form_oid_and_from_field_concepts():
    by = {f["form_oid"]: f for f in _sources()["forms"]}
    assert by["AEGEN"]["domain"] == "AE" and "field concepts" in by["AEGEN"]["domain_basis"]   # AEGEN matches AE
    assert by["MEDH"]["domain"] == "MH"
    assert sm.derive_domain("XYZ", "Something else", [{"type": "text", "name": "FOO"}])[0] is None
    assert sm.derive_domain("VS", "Vitals", [])[:2] == ("VS", "form OID")


def test_xlsform_wins_over_odm_and_referenced_sits_between():
    ref = [{"label": "referenced study REF1", "forms": [("AE.xlsx", fx.xlsform_bytes(settings={**fx.AE_SETTINGS, "form_id": "AE"}))]}]
    src = sm.load_sources([("standard.xml", fx.ODM)], ref)
    ae = next(f for f in src["forms"] if f["form_oid"] == "AE")
    assert ae["source"] == "referenced study REF1"            # the same form from the ODM is superseded
    with contextlib.redirect_stdout(io.StringIO()):
        out = sm.apply(fx.spec(), sm.load_sources([("AEGEN.xlsx", fx.xlsform_bytes()), ("standard.xml", fx.ODM)], ref))
    m = sm.state(out)["matched"][0]
    assert (m["standard_form"], m["source"]) == ("AEGEN", "uploaded XLSForm")   # uploaded XLSForm has priority


# ── match + exact splice ─────────────────────────────────────────────────────────

def test_matched_form_is_the_standard_exactly_and_keeps_scheduling():
    before = fx.spec()
    out = _matched()
    assert before == fx.spec()
    form = _form(out, "AEGEN")
    std = sm.parse_xlsform(fx.xlsform_bytes(), "AEGEN.xlsx")
    assert [sm.core_row(r) for r in form["survey"]] == std["survey"]            # every row, every column, in order
    assert [{k: v for k, v in c.items() if k != "source"} for c in form["choices"]] == std["choices"]
    assert form["settings"] == std["settings"] and form["extra_cols"] == ["oc:custom_col"]
    assert form["visits_assigned"] == ["SE_COMMON"] and form["reuse_count"] == 1    # protocol scheduling is kept
    assert form["library_match"]["status"] == "CUSTOMER_STANDARD"
    assert form["library_match"]["customer_form_name"] == form["customer_form_name"] == "AE General"
    assert form["library_match"]["source"] == "uploaded XLSForm"
    assert all(r["library_source"] == "CUSTOM" for r in form["survey"])
    assert out["schedule_of_events"]["form_placements"][0]["form_id"] == "AEGEN"    # references follow the form id
    st = sm.state(out)
    assert st["protocol_forms_without_standard"] == ["VS", "DIARY"]
    assert {s["form"] for s in st["standard_forms_not_used"]} == {"AE", "MEDH"}
    assert sm.integrity(out) == []


def test_ambiguous_domain_is_decided_by_name_and_logged():
    two = sm.load_sources([("AEGEN.xlsx", fx.xlsform_bytes()),
                           ("AESER.xlsx", fx.xlsform_bytes(settings={**fx.AE_SETTINGS, "form_id": "AESAE",
                                                                    "form_title": "Serious Adverse Events"}))])
    with contextlib.redirect_stdout(io.StringIO()):
        out = sm.apply(fx.spec(), two)
    m = sm.state(out)["matched"][0]
    assert m["protocol_form"] == "AE" and "share domain AE" in m["note"] and "best overall fit" in m["note"]
    assert len(sm.state(out)["standard_forms_not_used"]) == 1


def test_idempotent_and_rematch_only_when_sources_change():
    src = _sources()
    with contextlib.redirect_stdout(io.StringIO()):
        once = sm.apply(fx.spec(), src)
        assert sm.apply(once, src) is once and not sm.needs_match(once, src)
        other = sm.load_sources([("standard.xml", fx.ODM)])
        again = sm.apply(once, other)
    assert sm.needs_match(once, other)
    form = _form(again, "AE")                                    # restored, then matched to the ODM's AE form
    assert form["customer_standard"]["source"] == "uploaded ODM"
    assert [r["name"] for r in form["survey"]] == ["AETERM", "AESTDAT", "AEENDAT", "AESEV", "AEACN", "AEWT", "AENOQ", "AEDESC", "AELEFT"]
    assert again["schedule_of_events"]["form_placements"][0]["form_id"] == "AE"
    with contextlib.redirect_stdout(io.StringIO()):
        none = sm.apply(again, sm.load_sources([]))
    assert not sm.matched_forms(none) and [f["form_id"] for f in none["forms"]] == ["AE", "VS", "DIARY"]
    assert [r["name"] for r in _form(none, "AE")["survey"]] == ["AETERM", "AESTDAT", "AEENDAT", "AEHOSP", "AECOMMENT"]


def test_no_standards_leaves_the_spec_untouched():
    spec = fx.spec()
    assert sm.apply(spec, sm.load_sources([])) is spec and sm.apply(spec, None) is spec
    assert "standards_match" not in spec["study_meta"]


def test_kill_switch_and_exception_return_the_spec_unchanged(monkeypatch):
    spec, src = fx.spec(), _sources()
    monkeypatch.setenv("STANDARDS_MATCHING", "0")
    assert sm.apply(spec, src) is spec
    monkeypatch.delenv("STANDARDS_MATCHING")
    monkeypatch.setattr(sm, "match_forms", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    with contextlib.redirect_stdout(io.StringIO()):
        assert sm.apply(spec, src) is spec
    assert spec == fx.spec()


# ── rule 3 / 4 / 5: add only with a verified quote, never remove ─────────────────

def test_field_added_only_when_the_quote_is_in_the_protocol():
    out = _matched()
    assert [c["field"] for c in sm.add_candidates(out)] == ["AEHOSP", "AECOMMENT"]     # compared by concept / name
    prompt, extra = sm.build_add_request(out, fx.PROTOCOL_TEXT)
    assert "AEHOSP" in extra and "PROTOCOL TEXT" in extra and "VERBATIM" in prompt
    answer = json.dumps({"additions": [
        {"form_id": "AEGEN", "field": "AEHOSP", "quote": "record whether the event led to   hospitalisation of the participant"},
        {"form_id": "AEGEN", "field": "AECOMMENT", "quote": "a free-text comment must be captured for each event"},
        {"form_id": "AEGEN", "field": "AETERM", "quote": "All adverse events will be recorded from consent."},
        {"form_id": "AEGEN", "field": "AEHOSP", "quote": "All adverse events will be recorded from consent."}]})
    res = sm.apply_additions(out, answer, fx.PROTOCOL_TEXT)
    assert [a["field"] for a in res["added"]] == ["AEHOSP"]
    assert {c["field"] for c in res["covered"]} == {"AETERM", "AESTDAT", "AEENDAT"}   # already on the standard
    assert res["not_specified"] == []
    reasons = {(r.get("field"), r["reason"]) for r in res["rejected"]}
    assert reasons == {("AECOMMENT", "quote_not_in_protocol"), ("AETERM", "not_a_candidate"), ("AEHOSP", "duplicate")}
    form = _form(out, "AEGEN")
    row = form["survey"][-1]
    assert row["name"] == "AEHOSP" and row["provenance"] == "Added from protocol" and "hospitalisation" in row["protocol_quote"]
    assert row["library_source"] == "PROTOCOL_SPECIFIC" and row["bind__oc_itemgroup"] == "AEG"
    assert {c["name"] for c in form["choices"] if c["list_name"] == "YN"} == {"Y", "N"}
    assert form["library_match"]["fields_extended_from_protocol"] == 1
    assert form["library_match"]["added_from_protocol"][0]["field"] == "AEHOSP"
    assert [sm.core_row(r) for r in form["survey"][:-1]] == _standard_rows()            # the standard is untouched
    assert sm.integrity(out) == []
    assert sm.state(out)["matched"][0]["fields_added"][0]["field"] == "AEHOSP"


def test_same_data_point_under_another_name_is_never_added():
    """Compared by concept / data point, not by name: SYSBP vs SYSBP_VSORRES, IE01 vs IE001, AEREL vs AEREL1."""
    rows = [{"type": "decimal", "name": "SYSBP_VSORRES", "label": "Systolic blood pressure:"},
            {"type": "select_one NY", "name": "IE001", "label": "1. Did the participant consent?"},
            {"type": "select_one REL", "name": "AEREL1", "label": "Relationship to drug A"},
            {"type": "date", "name": "VSDAT", "label": "Date of measurements:"}]
    assert sm._covered_by("SYSBP", "Systolic BP (mmHg)", rows) == "SYSBP_VSORRES"
    assert sm._covered_by("IE01", "Signed consent", rows) == "IE001"
    assert sm._covered_by("AEREL", "Relationship to study drug", rows) == "AEREL1"
    assert sm._covered_by("MEASDT", "Date of measurements", rows) == "VSDAT"        # near-identical label
    assert sm._covered_by("TEMP", "Temperature", rows) is None
    # the AI names the standard field that covers a candidate: it is not added, even with a real quote
    out = _matched()
    res = sm.apply_additions(out, json.dumps({"additions": [
        {"form_id": "AEGEN", "field": "AEHOSP", "covered_by": "AESER",
         "quote": "record whether the event led to hospitalisation of the participant"},
        {"form_id": "AEGEN", "field": "AECOMMENT", "covered_by": "NOT_A_FIELD", "quote": None}]}), fx.PROTOCOL_TEXT)
    assert res["added"] == [] and {"form_id": "AEGEN", "field": "AEHOSP", "covered_by": "AESER"} in res["covered"]
    assert res["rejected"][0]["reason"] == "covered_by_unknown_field"
    assert [sm.core_row(r) for r in _form(out, "AEGEN")["survey"]] == _standard_rows()


def test_nothing_changes_when_the_protocol_specifies_nothing():
    out = _matched()
    for answer in ('{"additions": []}', "not json at all", "",
                   '{"additions": [{"form_id": "AEGEN", "field": "AECOMMENT", "covered_by": null, "quote": null}]}'):
        res = sm.apply_additions(out, answer, fx.PROTOCOL_TEXT)
        assert res["added"] == []
    assert [sm.core_row(r) for r in _form(out, "AEGEN")["survey"]] == _standard_rows()
    assert sm.build_add_request(out, "") is None                       # no protocol text: nothing is asked


def test_standard_fields_are_never_removed_by_the_rest_of_the_pipeline():
    out = _matched()
    import cdisc_ct, cdisc_cdash, cdisc_concepts
    std = cdisc_ct.load_standards(None)
    if std is not None:
        cdisc_ct.apply_to_spec(out, std, set())
        cdisc_cdash.apply_to_spec(out, std, set())
        cdisc_concepts.tag_deterministic(out, std)
    _engine(out); _engine(out)
    form = _form(out, "AEGEN")
    assert [sm.core_row(r) for r in form["survey"]] == _standard_rows()      # logic preserved byte for byte
    assert sm.integrity(out) == []


def test_cdisc_ct_never_overrides_a_matched_form():
    import cdisc_ct
    std = cdisc_ct.load_standards(None)
    if std is None:
        pytest.skip("CDISC CT not available")
    out = _matched()
    plain = fx.spec()
    plain["forms"][0]["survey"].append({"type": "select_one SEV", "name": "AESEV", "label": "Severity",
                                        "library_source": "CDASH_DEFAULT"})
    plain["forms"][0]["choices"] += [{"list_name": "SEV", "label": "Mild (custom)", "name": "1", "source": "CDASH"}]
    d_plain = [d for d in cdisc_ct.apply_to_spec(plain, std, set()) if d["field"] == "AESEV"]
    d_std = [d for d in cdisc_ct.apply_to_spec(out, std, set()) if d["form_id"] == "AEGEN"]
    assert d_std and all(d["action"] == "kept" for d in d_std)
    assert d_plain and d_plain[0]["action"] != "kept"                       # the same field is CT-managed elsewhere
    assert {c["label"] for c in _form(out, "AEGEN")["choices"] if c["list_name"] == "SEV"} == \
        {"Mild (custom)", "Moderate (custom)", "Severe (custom)"}


# ── rule 7: rules-engine checks on a standard form are proposals ─────────────────

def _dvs_bytes(spec):
    from extract_dvs_from_forms import extract_dvs_data
    from generate_dvs import build_dvs
    fj = {"forms": {f"{f['form_id']}.xlsx": {"survey": [{("bind::oc:itemgroup" if k == "bind__oc_itemgroup" else k): v
                                                          for k, v in r.items()} for r in f["survey"]],
                                               "choices": f.get("choices") or []} for f in spec["forms"]}}
    path = os.path.join(tempfile.mkdtemp(), "dvs.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        build_dvs(extract_dvs_data(spec, fj), path)
    return open(path, "rb").read()


def _set_actions(xbytes, actions):
    wb = openpyxl.load_workbook(io.BytesIO(xbytes)); ws = wb["DVS_OC4"]
    col = {str(c.value).strip(): i for i, c in enumerate(ws[3], 1) if c.value}
    seen = {}
    for r in range(4, ws.max_row + 1):
        cid = ws.cell(r, col["Check ID"]).value
        if cid in actions:
            ws.cell(r, col["Action"]).value = actions[cid]
            seen[cid] = {h: ws.cell(r, i).value for h, i in col.items()}
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue(), seen


def test_engine_proposes_and_does_not_mutate_a_standard_form():
    out = _engine(_matched())
    props = {p["target_field"]: p for p in sm.proposals(out) if p["kind"] == "engine"}
    assert "AEENDAT" in props and props["AEENDAT"]["check_type"] == "Constraint"
    p = props["AEENDAT"]
    assert p["convention_id"] and p["logic"] and p["ops"] and p["id"].startswith("STD-")
    assert p["source"] in ("CDISC Standard", "Global Rule", "Customer Rule") or p["source"].startswith("CDISC CORE (")
    form = _form(out, "AEGEN")
    assert [sm.core_row(r) for r in form["survey"]] == _standard_rows()
    assert "constraint" not in next(r for r in form["survey"] if r.get("name") == "AEENDAT")
    again = copy.deepcopy(sm.proposals(_engine(out)))                       # the engine runs several times per build
    assert again == sm.proposals(_engine(out)) and len({x["id"] for x in again}) == len(again)
    # the same convention is applied as usual on a form that is not a customer standard
    plain = _engine(fx.spec())
    assert next(r for r in plain["forms"][0]["survey"] if r["name"] == "AEENDAT").get("constraint")


def test_dvs_lists_proposals_and_approve_reject_round_trip():
    out = _engine(_matched())
    props = [p for p in sm.proposals(out) if p["kind"] == "engine"]
    assert len(props) >= 2
    approve, reject = props[0], props[1]
    xb, seen = _set_actions(_dvs_bytes(out), {approve["id"]: "Approve", reject["id"]: "Reject"})
    row = seen[approve["id"]]
    assert row["Status"] == "Proposed" and row["Check Source"] == approve["source"]
    assert row["Expression / Calculation"] == approve["logic"] and row["Target Form OID"] == "AEGEN"
    assert json.loads(row["Machine Data"])["kind"] == "standard_proposal"
    own = [r for r in dvs_edits.parse_actions(_dvs_bytes(out)) or []]
    assert own == []                                                        # nothing has an action yet
    res = dvs_edits.apply_actions(out, dvs_edits.parse_actions(xb))
    assert res["applied"] == 2, res["results"]
    field = next(r for r in _form(out, "AEGEN")["survey"] if r.get("name") == approve["target_field"])
    assert approve["logic"] in str(field.get("constraint") or field.get("relevant") or field.get("required"))
    ids = {p["id"] for p in sm.proposals(out)}
    assert approve["id"] not in ids and reject["id"] not in ids
    assert sm.integrity(out) == [{"form_id": "AEGEN", "reason": "changed by approved proposals"}]
    snapshot = copy.deepcopy(_form(out, "AEGEN")["survey"])
    _engine(out); _engine(out)                                              # approved check stays, rejected stays away
    assert [sm.core_row(r) for r in _form(out, "AEGEN")["survey"]] == [sm.core_row(r) for r in snapshot]
    ids = {p["id"] for p in sm.proposals(out)}
    assert approve["id"] not in ids and reject["id"] not in ids
    again = dvs_edits.apply_actions(out, dvs_edits.parse_actions(xb))       # the same DVS uploaded twice
    assert again["applied"] == 0 and again["failed"] == 0
    # the approved check is now a build check with its source; the standard's own logic reads Customer Standard
    from extract_dvs_from_forms import extract_dvs_data
    fj = {"forms": {"AEGEN.xlsx": {"survey": [{("bind::oc:itemgroup" if k == "bind__oc_itemgroup" else k): v
                                                for k, v in r.items()} for r in _form(out, "AEGEN")["survey"]],
                                   "choices": _form(out, "AEGEN")["choices"]}}}
    with contextlib.redirect_stdout(io.StringIO()):
        rows = extract_dvs_data(out, fj)["dvs_oc4"]
    own = [r for r in rows if r["Target Item Name"] == "AESTDAT" and r["Check Type"] == "Constraint"]
    assert own and own[0]["Check Source"] == "Customer Standard" and own[0]["Status"] != "Proposed"
    if approve["check_type"] == "Constraint":
        done = [r for r in rows if r["Target Item Name"] == approve["target_field"] and r["Status"] != "Proposed"
                and r["Expression / Calculation"] == approve["logic"]]
        assert done and done[0]["Check Source"] == approve["source"]


def test_apply_proposal_is_all_or_nothing():
    out = _engine(_matched())
    p = copy.deepcopy(sm.proposals(out)[0])
    p["ops"].append({"op": "row", "field": "NOPE", "set": {"relevant": "1"}, "before": {}})
    before = copy.deepcopy(_form(out, "AEGEN"))
    assert sm.apply_proposal(out, p)[0] == "failed"
    assert _form(out, "AEGEN") == before


def test_move_to_form_never_takes_a_field_out_of_a_standard_form():
    from conventions_engine import EntityContext, effects
    out = _matched()
    form = _form(out, "AEGEN")
    row = next(r for r in form["survey"] if r.get("name") == "AESEV")
    ctx = EntityContext(kind="field", entity=row, parent=form, spec=out, path="x")
    effects.apply_effect({"move_to_form": "VS"}, ctx, out, "test.move")
    effects.sweep_pending_removals(out)
    assert any(r.get("name") == "AESEV" for r in form["survey"])
    assert not any(r.get("name") == "AESEV" for r in _form(out, "VS")["survey"])
    assert out["review_flags"]["customer_standard_not_applied"]


# ── build + documents ────────────────────────────────────────────────────────────

def test_builder_writes_an_xlsform_standard_verbatim():
    from build_xlsforms import build_single_xlsform
    out = _matched()
    form = _form(out, "AEGEN")
    path = os.path.join(tempfile.mkdtemp(), "AEGEN.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        build_single_xlsform(copy.deepcopy(form), path, {"placeholder_applied": []})
    built = sm.parse_xlsform(open(path, "rb").read(), "AEGEN.xlsx")
    assert built["survey"] == _standard_rows()                      # groups, logic, messages and extra columns
    assert built["choices"] == sm.parse_xlsform(fx.xlsform_bytes(), "AEGEN.xlsx")["choices"]
    assert {k: v for k, v in built["settings"].items() if k != "version"} == \
        {k: v for k, v in fx.AE_SETTINGS.items() if k != "version"}


def test_provenance_in_study_spec_and_protocol_summary():
    from generate_study_spec_pdf import build_edc_pdf
    from generate_study_spec_xlsx import build_edc_xlsx
    from generate_protocol_summary_pdf import _apply_customer_standard_provenance
    out = _matched()
    sm.apply_additions(out, json.dumps({"additions": [{"form_id": "AEGEN", "field": "AEHOSP",
                       "quote": "record whether the event led to hospitalisation of the participant"}]}), fx.PROTOCOL_TEXT)
    tmp = tempfile.mkdtemp()
    with contextlib.redirect_stdout(io.StringIO()):
        build_edc_pdf(copy.deepcopy(out), os.path.join(tmp, "s.pdf"))
        build_edc_xlsx(copy.deepcopy(out), os.path.join(tmp, "s.xlsx"))
    assert os.path.getsize(os.path.join(tmp, "s.pdf")) > 1000
    wb = openpyxl.load_workbook(os.path.join(tmp, "s.xlsx"))
    text = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    assert "CUSTOMER_STANDARD" in text and "AE General" in text and "uploaded XLSForm" in text
    assert "AEHOSP" in text and "hospitalisation of the participant" in text
    details = [{"domain_name": "Adverse Events", "cdash_code": "AE", "source": "CDASH_ESTIMATE", "customer_form_name": None},
               {"domain_name": "Vital Signs", "cdash_code": "VS", "source": "CDASH_ESTIMATE", "customer_form_name": None}]
    got = _apply_customer_standard_provenance(details, out)
    assert got[0]["source"] == "CUSTOMER_STANDARD" and got[0]["customer_form_name"] == "AE General"
    assert got[0]["customer_standard_source"] == "uploaded XLSForm"
    assert got[1]["source"] == "CDASH_ESTIMATE" and details[0]["source"] == "CDASH_ESTIMATE"
    lines = "\n".join(sm.summary_lines(out))
    assert "AE -> AE General [AEGEN]" in lines and "added from protocol: AEHOSP" in lines


def test_catalog_is_compact_and_names_forms_domains_and_fields():
    cat = sm.catalog_text(_sources())
    assert "FORM AEGEN | AE General | CDASH domain: AE" in cat and "AETERM, AESTDAT" in cat
    assert "today()" not in cat and len(cat) < 1500


def test_odm_standard_builds_without_changing_the_study_spec():
    """An ODM-only standard has no settings beyond title and id: the builder fills the OC defaults and the version
    in the file only, the build QA accepts the customer's settings, and the spec keeps the standard exactly."""
    from build_xlsforms import build_single_xlsform
    from build_checklist import run_qa_checks
    out = _matched(xls=False)
    form = _form(out, "AE")
    assert form["customer_standard"]["verbatim"] is False and form["customer_standard"]["has_logic"] is False
    before = copy.deepcopy(form)
    path = os.path.join(tempfile.mkdtemp(), "AE.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        build_single_xlsform(form, path, {"placeholder_applied": []})
    assert form == before and sm.integrity(out) == []
    built = sm.parse_xlsform(open(path, "rb").read(), "AE.xlsx")
    assert built["settings"]["style"] and built["settings"]["namespaces"] and built["settings"]["form_id"] == "AE"
    assert [r["name"] for r in built["survey"] if r.get("name")] == [r["name"] for r in form["survey"]]
    qa = {c: status for c, status, _note in run_qa_checks(form, {})}
    assert qa["settings_complete"] == "PASS"
    plain = fx.spec()["forms"][1]
    plain["settings"] = {"form_title": "Vital Signs", "form_id": "VS"}
    assert {c: st for c, st, _n in run_qa_checks(plain, {})}["settings_complete"] == "FAIL"   # unchanged elsewhere


def test_odm_item_without_question_text_never_gets_an_empty_label():
    """Label fallback: Question, else the ODM item Description, else the OpenClinica item details, else the item
    name with the review flag "label from ODM name, no question text"."""
    ae = next(f for f in sm.parse_odm(fx.ODM, "standard.xml") if f["form_oid"] == "AE")
    rows = {r["name"]: r for r in ae["survey"]}
    assert all(str(r.get("label") or "").strip() for f in sm.parse_odm(fx.ODM, "standard.xml") for r in f["survey"])
    assert rows["AETERM"]["label"] == "Event term" and "flag_reason" not in rows["AETERM"]
    assert rows["AEDESC"]["label"] == "Event description text" and "flag_reason" not in rows["AEDESC"]
    assert rows["AELEFT"]["label"] == "Left item text label" and "flag_reason" not in rows["AELEFT"]
    assert rows["AENOQ"]["label"] == "AENOQ"
    assert rows["AENOQ"]["flag_reason"] == "label from ODM name, no question text"
    assert rows["AENOQ"]["completion_status"] == "FLAGGED"
    out = _matched(xls=False)
    row = next(r for r in _form(out, "AE")["survey"] if r["name"] == "AENOQ")
    assert row["label"] == "AENOQ" and row["completion_status"] == "FLAGGED"
    assert row["flag_reason"] == "label from ODM name, no question text"
    assert out["review_flags"]["customer_standard_label_from_name"] == ["AE.AENOQ: label from ODM name, no question text"]
    assert next(r for r in _form(out, "AE")["survey"] if r["name"] == "AETERM")["completion_status"] == "COMPLETE"
    assert sm.integrity(out) == []
    with contextlib.redirect_stdout(io.StringIO()):
        gone = sm.apply(out, sm.load_sources([]))            # the standard is removed: its flags go with it
    assert "customer_standard_label_from_name" not in gone["review_flags"]

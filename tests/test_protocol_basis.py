"""Protocol basis (protocol_basis.py): the protocol defines the forms, so a form the analysis created that no protocol
text asks for is not built. Synthetic protocols and forms only."""
import copy
import json
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "skills", "protocol-analysis", "scripts"))
import protocol_basis as pb
import protocol_forms as pf
import standards_match as sm
from tests.standards import fixtures as fx

PROTOCOL = ("6.1 Vital signs will be recorded at every visit. 6.2 The date of the most recent widget calibration "
            "will be collected at screening. 6.3 Participants who complete the washout period enter the treatment "
            "period. 9.1 All adverse events will be reported from consent until the last visit.")
Q_VS = "Vital signs will be recorded at every visit."
Q_WID = "The date of the most recent widget calibration will be collected at screening."
Q_MENTION = "Participants who complete the washout period enter the treatment period."


@pytest.fixture(autouse=True)
def env(monkeypatch):
    monkeypatch.setitem(sm._CDASH, "v", (None, {}, {}, set()))
    monkeypatch.delenv("PROTOCOL_BASIS_CHECK", raising=False)


def _vs():
    return fx._f("VS", "Vital Signs", "VS", [("date", "VSDAT", "Date", {}), ("decimal", "SYSBP", "Systolic", {})])


def _wid():
    return fx._f("WID", "Widget Calibration", None, [("date", "WIDDAT", "Date of last calibration", {})])


def _invented(fid="WASH", title="Washout Period", visits=("SE_SCREENING",)):
    f = fx._f(fid, title, None, [("date", f"{fid}DAT", "Washout start date", {}), ("text", f"{fid}TXT", "Comment", {})],
              visits=visits)
    f["library_match"] = {"status": "CUSTOM", "source_type": "protocol"}
    return f


def _spec(*forms):
    s = {"study_meta": {}, "forms": list(forms), "review_flags": {},
         "schedule_of_events": {"form_placements": [{"target_visit_oid": v, "form_id": f["form_id"]}
                                                    for f in forms for v in f.get("visits_assigned") or []]}}
    return s


def _resp(**forms):
    return json.dumps({"forms": [dict({"form_id": k}, **v) for k, v in forms.items()]})


def _req(q, section="6", fields=("VSDAT", "WIDDAT", "WASHDAT"), kind="record"):
    return {"quotes": [{"text": q, "section": section, "kind": kind}], "fields": list(fields), "why": "asked for"}


MENTION = {"quotes": [{"text": Q_MENTION, "section": "6.3", "kind": "course"}], "fields": [], "why": "a study period, no data asked for"}
NONE = {"quotes": [], "fields": []}


def test_a_form_the_protocol_never_asks_for_is_removed_everywhere_and_listed_as_not_built():
    spec = _spec(_vs(), _wid(), _invented())
    spec["study_meta"]["ai_edit_checks"] = {"proposals": [{"target_form": "WASH", "target_field": "WASHDAT", "source_form": "WASH"},
                                                          {"target_form": "VS", "target_field": "VSDAT", "source_form": "VS"}]}
    spec["study_meta"]["standards_match"] = {"protocol_forms_without_standard": ["WID", "WASH"]}
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WID=_req(Q_WID), WASH=MENTION), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "WID"]
    assert all(p["form_id"] != "WASH" for p in spec["schedule_of_events"]["form_placements"])
    assert spec["study_meta"]["ai_edit_checks"]["proposals"] == [{"target_form": "VS", "target_field": "VSDAT", "source_form": "VS"}]
    assert spec["study_meta"]["standards_match"]["protocol_forms_without_standard"] == ["WID"]
    assert spec["study_meta"]["protocol_basis_removed"]["WASH"]["form_title"] == "Washout Period"
    rem = st["removed"]
    assert [r["form_id"] for r in rem] == ["WASH"] and rem[0]["reason"].startswith("no protocol text asks for this data")
    assert rem[0]["mention"] == Q_MENTION and st["events_emptied"] == []
    assert any("WASH (Washout Period): not built, no protocol text asks for this data" in m for m in spec["review_flags"][pb.FLAG])
    title, _note, headers, rows, _w = pb.section(spec)
    assert title == "FORMS NOT BUILT" and headers[:5] == ["Form", "Title", "Content from", "Status", "Reason"]
    assert rows == [["WASH", "Washout Period", "protocol analysis", "Not built", rem[0]["reason"], Q_MENTION]]
    assert any("WASH (Washout Period): not built" in line for line in pb.summary_lines(spec))
    assert not pb.needs_check(spec, PROTOCOL)   # runs once; the record of the removed form stays
    import openpyxl
    from generate_study_spec_xlsx import build_forms_not_built_sheet
    from generate_study_spec_pdf import build_forms_not_built_block, make_styles
    wb = openpyxl.Workbook()
    build_forms_not_built_sheet(wb, spec)
    assert wb["FORMS_NOT_BUILT"].cell(row=3, column=1).value == "WASH"
    assert len(build_forms_not_built_block(spec, make_styles())) == 4
    wb2 = openpyxl.Workbook()
    build_forms_not_built_sheet(wb2, {"study_meta": {}})
    assert "FORMS_NOT_BUILT" not in wb2.sheetnames and build_forms_not_built_block({"study_meta": {}}, make_styles()) == []


def test_a_form_with_a_verified_quote_is_kept_and_the_quote_is_recorded():
    spec = _spec(_vs(), _wid())
    st = pb.apply(spec, _resp(VS=_req(Q_VS, "6.1"), WID=_req(Q_WID, "6.2")), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "WID"] and st["removed"] == []
    rec = {r["form_id"]: r for r in st["forms"]}
    assert rec["WID"]["basis"] == "verified protocol quote (instruction to record)" and rec["WID"]["quote"] == Q_WID and rec["WID"]["section"] == "6.2"
    assert pb.FLAG not in spec["review_flags"] and pb.section(spec) is None


def test_a_quote_that_is_not_in_the_protocol_does_not_support_the_form():
    spec = _spec(_vs(), _invented())
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WASH=_req("The washout start date will be recorded on the case report form.")),
                  PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS"]
    assert "the quoted text is not in the protocol" in st["removed"][0]["reason"]
    assert st["rejected"] == {"quote not found in the protocol": 1}


def test_a_real_quote_that_asks_for_none_of_the_forms_fields_does_not_support_it():
    spec = _spec(_vs(), _invented())
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WASH=_req(Q_MENTION, fields=["NOT_A_FIELD"])), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS"]
    assert "the protocol only mentions it" in st["removed"][0]["reason"] and st["removed"][0]["mention"] == Q_MENTION
    assert st["rejected"] == {"requirement without a field of the form": 1}
    # the kind of passage decides, not the model's own verdict: a real quote that is no instruction supports nothing
    for kind in ("eligibility", "course", "process", "definition", "time", "other", "made up"):
        spec = _spec(_vs(), _invented())
        pb.apply(spec, _resp(VS=_req(Q_VS, kind="schedule"), WASH=_req(Q_MENTION, kind=kind)), PROTOCOL, fresh=True)
        assert [f["form_id"] for f in spec["forms"]] == ["VS"], kind
    assert {r["form_id"]: r for r in st["forms"]}["VS"]["fields"] == ["VSDAT"]


def test_an_assessment_of_the_completeness_check_supports_a_form_without_a_quote():
    spec = _spec(_vs(), _invented())
    spec["study_meta"]["protocol_forms"] = {"status": "done", "assessments": [
        {"assessment": "Washout", "form": "WASH", "section": "6.3", "quote": Q_MENTION}]}
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WASH=NONE), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "WASH"] and st["removed"] == []
    assert "assessment of the completeness check" in {r["form_id"]: r for r in st["forms"]}["WASH"]["basis"]


def test_a_customer_standard_form_without_protocol_text_is_kept_and_flagged():
    oc4 = _invented("SPREV", "Sponsor Review")
    oc4["customer_standard"] = {"form_oid": "SPREV", "source": "uploaded ODM"}
    crf = _invented("LIBF", "Site Checklist")
    crf["library_match"] = {"status": "LIBRARY_MATCH", "source_type": "customer"}
    spec = _spec(_vs(), oc4, crf)
    st = pb.apply(spec, _resp(VS=_req(Q_VS), SPREV=NONE, LIBF=NONE),
                  PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "SPREV", "LIBF"] and st["removed"] == []
    rec = {r["form_id"]: r for r in st["forms"]}
    for fid in ("SPREV", "LIBF"):
        assert rec[fid]["action"] == pb.KEPT_STANDARD and rec[fid]["reason"] == "customer standard form, no protocol text found"
    assert sum("customer standard form, no protocol text found" in m for m in spec["review_flags"][pb.FLAG]) == 2
    assert pb.section(spec) is None


def test_a_form_another_forms_logic_reads_is_kept_and_flagged():
    xp = "instance('clinicaldata')/ODM/ClinicalData/SubjectData/StudyEventData/FormData[@FormOID='F_WASH']/ItemGroupData/ItemData[@OpenClinica:ItemName='WASHDAT']/@Value"
    vs = _vs()
    vs["survey"].append({"type": "calculate", "name": "WASHDAT_CF", "label": "", "calculation": xp})
    spec = _spec(vs, _invented())
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WASH=MENTION), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "WASH"] and st["removed"] == []
    rec = {r["form_id"]: r for r in st["forms"]}["WASH"]
    assert rec["action"] == pb.KEPT_REFERENCED and rec["referenced_by"] == ["VS"]
    assert any("kept because VS reads its fields" in m for m in spec["review_flags"][pb.FLAG])
    # the declared dependency counts as well; a form only another removed form reads goes with it
    vs2 = _vs()
    vs2["cross_form_dependencies"] = [{"source_form": "WASH", "source_field": "WASHDAT"}]
    assert pb.referenced_by(_spec(vs2, _invented()), _invented()) == ["VS"]
    other = _invented("WASHB", "Washout Follow-up")
    other["survey"].append({"type": "calculate", "name": "X_CF", "label": "", "calculation": xp})
    spec = _spec(_vs(), _invented(), other)
    pb.apply(spec, _resp(VS=_req(Q_VS), WASH=MENTION, WASHB=MENTION), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS"]


def test_a_reused_specification_is_never_changed_the_result_is_recorded_and_flagged():
    spec = _spec(_vs(), _invented())
    before = copy.deepcopy(spec["forms"]), copy.deepcopy(spec["schedule_of_events"])
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WASH=MENTION), PROTOCOL, fresh=False)
    assert (spec["forms"], spec["schedule_of_events"]) == before and st["removed"] == []
    assert "protocol_basis_removed" not in spec["study_meta"]
    rec = {r["form_id"]: r for r in st["forms"]}["WASH"]
    assert rec["action"] == pb.WOULD_REMOVE and "reused, not freshly analysed" in rec["reason"]
    assert any(m.startswith("WASH (Washout Period): no protocol text asks for this data") for m in spec["review_flags"][pb.FLAG])
    assert pb.section(spec)[3][0][3] == "Kept (reused specification)"


def test_an_ai_failure_or_an_untrustworthy_answer_removes_nothing():
    for response in ("", "the model wrote prose", json.dumps({"forms": [{"form_id": "NOPE", "quotes": []}]})):
        spec = _spec(_vs(), _invented())
        st = pb.apply(spec, response, PROTOCOL, fresh=True)
        assert [f["form_id"] for f in spec["forms"]] == ["VS", "WASH"] and st["status"] == "not_run"
        assert pb.summary_lines(spec) == ["Protocol basis check did not run (the AI answer could not be read): no form was removed."]
    # most forms without a basis is not a credible answer
    forms = [_invented(f"F{i}", f"Form {i}") for i in range(5)]
    spec = _spec(_vs(), *forms)
    st = pb.apply(spec, _resp(VS=_req(Q_VS), **{f"F{i}": NONE for i in range(5)}), PROTOCOL, fresh=True)
    assert len(spec["forms"]) == 6 and st["status"] == "not_run" and "not credible" in st["note"]
    # a form the answer does not cover is kept and flagged
    spec = _spec(_vs(), _invented())
    st = pb.apply(spec, _resp(VS=_req(Q_VS)), PROTOCOL, fresh=True)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "WASH"]
    assert {r["form_id"]: r for r in st["forms"]}["WASH"]["action"] == pb.KEPT_NOT_JUDGED


def test_pipeline_and_convention_forms_are_not_candidates_but_completeness_check_additions_are():
    dov = _invented("DOV", "Date of Visit")
    sae = _invented("SAE", "Serious Adverse Events")
    conv = _invented("SAEX", "SAE Report")
    conv["convention_required"] = {"convention": "SAE_FORM", "content_source": "CDASHIG"}
    added = _invented("WID2", "Widget Calibration")
    added["protocol_required"] = {"assessment": "Widget calibration", "section": "6.2", "quote": Q_WID, "content_source": pf.SRC_CDASH}
    spec = _spec(_vs(), dov, sae, conv, added)
    answers = {"SAE_FORM": {"value": "yes"}, "DEATH_DETAILS_FORM": {"value": "protocol"}}
    assert [f["form_id"] for f in pb.candidates(spec, PROTOCOL, ["DOV"], answers)] == ["VS", "WID2"]
    assert [f["form_id"] for f in pb.candidates(spec, PROTOCOL, [], {})] == ["VS", "DOV", "SAE", "WID2"]
    prompt, extra = pb.build_request(spec, PROTOCOL, ["DOV"], answers)
    assert "FORM VS | Vital Signs" in extra and "VSDAT: Date" in extra and "FORM DOV" not in extra and "PROTOCOL TEXT:" in extra
    assert "FORM WID2 | Widget Calibration" in extra and '"heading"' in prompt
    assert "PROTOCOL TEXT:" not in pb.build_request(spec, PROTOCOL, ["DOV"], answers, with_text=False)[1]
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WID2=NONE), PROTOCOL, True, ["DOV"], answers)
    assert [f["form_id"] for f in spec["forms"]] == ["VS", "DOV", "SAE", "SAEX"]
    rec = {r["form_id"]: r for r in st["forms"]}
    assert rec["DOV"]["basis"] == "required by a pipeline rule" and not rec["DOV"]["checked"]
    assert [r["form_id"] for r in st["removed"]] == ["WID2"]


def test_an_event_left_without_forms_is_reported():
    spec = _spec(_vs(), _invented(visits=("SE_WASHOUT",)))
    st = pb.apply(spec, _resp(VS=_req(Q_VS), WASH=MENTION), PROTOCOL, fresh=True)
    assert st["events_emptied"] == ["SE_WASHOUT"]
    assert any("event SE_WASHOUT has no form left" in m for m in spec["review_flags"][pb.FLAG])


def test_kill_switch_and_no_protocol_text(monkeypatch):
    spec = _spec(_vs(), _invented())
    assert pb.needs_check(spec, PROTOCOL) and not pb.needs_check(spec, "") and pb.build_request(spec, "") is None
    monkeypatch.setenv("PROTOCOL_BASIS_CHECK", "0")
    assert not pb.needs_check(spec, PROTOCOL)


def test_the_pipeline_step_runs_after_the_completeness_check_and_survives_a_failed_call(monkeypatch):
    src = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pipeline.py")).read()
    body = src[src.index("async def _post_match_forms_step("):src.index("async def _form_conventions_step(")]
    assert body.index("_protocol_forms_step(") < body.index("_protocol_basis_step(") < body.index("_enforce_common_visit(")
    assert "answers, fresh)" in body
    step = src[src.index("async def _protocol_basis_step("):src.index("def _protocol_forms_refresh(")]
    assert "no form removed" in step and "return struct_json" in step.split("except Exception as _ce:")[1].split("_pb.apply(")[0]
    assert "PROTOCOL_BASIS" not in open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "prompts.py")).read()

"""
Offline re-score of the 2026-10-06 BioIVT UAT runs (Detroit, Precision).

Runs the loader's own _evaluate_uat_cases over the saved result workbooks with
the per-participant data OpenClinica actually stored
(docs/uat_analysis/participant_data.json). See
docs/UAT_FAILURE_ANALYSIS_2026-10-06.md. No network, nothing is created.
"""
import collections
import io
import json
import os
import sys

import openpyxl
import pytest

_REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, _REPO)

import uat_loader  # noqa: E402

_ANALYSIS = os.path.join(_REPO, "docs", "uat_analysis")
_RUNS = {
    "Detroit":   "detroit_bioivt_uat_results.xlsx",
    "Precision": "precision_bioivt_uat_results.xlsx",
}
_RESULT_COLS = ("Status", "Actual Result", "Test Result", "Execution Date", "Notes")


def _header(ws):
    for row in ws.iter_rows():
        if row[0].value == "UAT Case ID":
            return row[0].row, {str(c.value).strip(): c.column for c in row if c.value}
    raise AssertionError("UAT_Cases header row not found")


def _cases(wb_bytes):
    ws = openpyxl.load_workbook(io.BytesIO(wb_bytes), data_only=True)["UAT_Cases"]
    hdr, col = _header(ws)
    return [{k: r[c - 1].value for k, c in col.items()}
            for r in ws.iter_rows(min_row=hdr + 1) if r[0].value]


def _load_run(label):
    """(workbook bytes with result columns cleared, clinical_data, job_failures)."""
    with open(os.path.join(_ANALYSIS, "participant_data.json")) as f:
        stored = json.load(f)[label]
    clinical_data = {}
    for pkey, rows in stored.items():
        uat_loader._add_participant_data(
            clinical_data, pkey, {tuple(r[:4]): r[4] for r in rows})

    wb = openpyxl.load_workbook(os.path.join(_ANALYSIS, _RUNS[label]))
    ws = wb["UAT_Cases"]
    hdr, col = _header(ws)
    job_failures = {}
    for row in ws.iter_rows(min_row=hdr + 1):
        if not row[0].value:
            continue
        actual = str(row[col["Actual Result"] - 1].value or "")
        if actual.startswith("Import failed: "):
            pkey = str(row[col["Participant_Key"] - 1].value or "").strip()
            item = str(row[col["Item_OID"] - 1].value or "").strip().upper()
            job_failures[(pkey, item)] = actual[len("Import failed: "):]
        for name in _RESULT_COLS:      # back to a freshly stamped workbook
            row[col[name] - 1].value = None
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), clinical_data, job_failures


def _score(label):
    wb_bytes, clinical_data, job_failures = _load_run(label)
    scored = uat_loader._evaluate_uat_cases(wb_bytes, {}, clinical_data, job_failures)
    counts = collections.Counter()
    for c in _cases(scored):
        actual = str(c["Actual Result"] or "")
        if c["Test Result"] == "Fail" and actual.startswith("Import failed"):
            counts["Fail (import)"] += 1
        else:
            counts[str(c["Test Result"])] += 1
    return counts


@pytest.mark.parametrize("label,expected_pass", [("Detroit", 196), ("Precision", 178)])
def test_participant_aware_rescore_matches_analysis(label, expected_pass):
    """Each case is compared with its own participant's stored value."""
    assert _score(label)["Pass"] == expected_pass


def test_cases_for_one_item_are_scored_per_participant():
    """Detroit CM.CMSTDTC: P001 and P002 each stored what was loaded."""
    wb_bytes, clinical_data, job_failures = _load_run("Detroit")
    scored = uat_loader._evaluate_uat_cases(wb_bytes, {}, clinical_data, job_failures)
    by_id = {c["UAT Case ID"]: c for c in _cases(scored)}
    assert by_id["UAT-052"]["Test Result"] == "Pass"
    assert by_id["UAT-053"]["Test Result"] == "Pass"
    assert by_id["UAT-052"]["Actual Result"] != by_id["UAT-053"]["Actual Result"]


def test_participant_key_falls_back_to_stamp_map():
    """A workbook without a stamped Participant_Key still resolves the participant."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UAT_Cases"
    cols = ["UAT Case ID", "Status", "Actual Result", "Test Result", "Execution Date",
            "Notes", "Expected Result", "Scenario", "Participant_Key", "Study_Event_OID",
            "Form_OID", "Item_Group_OID", "Item_OID", "Participant_ID", "Load_Value"]
    ws.append(cols)
    base = ["", "", "", "", "", "", "No constraint error. Form saves.", "Happy path",
            "", "SE_A", "F_X", "IG_X_G", "I_X_AGE"]
    ws.append(["UAT-001"] + base[1:] + ["UAT-P001", "40"])
    ws.append(["UAT-002"] + base[1:] + ["UAT-P002", "50"])
    buf = io.BytesIO()
    wb.save(buf)

    clinical_data = {}
    coords = ("SE_A", "F_X", "IG_X_G", "I_X_AGE")
    uat_loader._add_participant_data(clinical_data, "RUN-P001", {coords: "40"})
    uat_loader._add_participant_data(clinical_data, "RUN-P002", {coords: "50"})
    stamp_map = {"UAT-P001": {"participant_key": "RUN-P001"},
                 "UAT-P002": {"participant_key": "RUN-P002"}}
    scored = uat_loader._evaluate_uat_cases(buf.getvalue(), stamp_map, clinical_data, {})
    assert [c["Test Result"] for c in _cases(scored)] == ["Pass", "Pass"]


def test_job_failures_are_keyed_by_participant():
    log = "ParticipantID,StudyEventOID,FormOID,ItemOID,Status,Message\n" \
          "x,SE_A,F_X,I_X_AGE,Failed,errorCode.bad\n" \
          "x,SE_A,F_X,I_X_SEX,Inserted,\n"
    failures = uat_loader._job_failures_from_imports([
        {"participant_key": "RUN-P001", "result": {"log": log}},
        {"participant_key": "RUN-P002", "result": {"log": ""}},
    ])
    assert failures == {("RUN-P001", "I_X_AGE"): "errorCode.bad"}

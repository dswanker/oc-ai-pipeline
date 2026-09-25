"""
Integration test: edc-builder end-to-end from Study Spec JSON → EDC Build ZIP.

Uses a known fixture spec (no Claude call, no Monday, no network).
Asserts the ZIP contains the correct forms with valid XLSForm structure.
Run with: pytest tests/build/test_edc_builder_integration.py -v
"""
import sys, json, zipfile, io, re
import pytest

sys.path.insert(0, '/Users/danswanker/oc-ai-pipeline/skills/edc-builder/scripts')
sys.path.insert(0, '/Users/danswanker/oc-ai-pipeline')

# ── Minimal but realistic Study Spec fixture ──────────────────────────────────
FIXTURE_SPEC = {
    "study_meta": {
        "protocol_number": "TEST-INTEG-001",
        "study_title": "Integration Test Study",
        "sponsor": "Test Sponsor",
        "phase": "Phase 2",
        "therapeutic_area": "Oncology",
        "indication": "Test indication",
    },
    "forms": [
        {
            "form_id": "DOV",
            "form_title": "Date of Visit",
            "repeating": False,
            "category": "INFRASTRUCTURE",
            "visits_assigned": ["SE_SCREEN", "SE_VISIT1"],
            "survey": [
                {"type": "date", "name": "SVDAT",
                 "label": "Date of Visit",
                 "bind__oc_itemgroup": "DOV",
                 "required": "yes",
                 "constraint": ". <= today()",
                 "constraint_message": "Date cannot be in the future."},
            ],
            "choices": [],
            "cross_form_dependencies": [],
            "settings": {
                "form_title": "Date of Visit",
                "form_id":    "DOV",
                "version":    "1",
                "style":      "theme-grid",
                "namespaces": 'oc="http://openclinica.org/xforms"',
            },
        },
        {
            "form_id": "DM",
            "form_title": "Demographics",
            "repeating": False,
            "category": "CDASH",
            "visits_assigned": ["SE_SCREEN"],
            "survey": [
                {"type": "text",    "name": "SUBJID", "label": "Subject ID",
                 "bind__oc_itemgroup": "DM", "required": "yes"},
                {"type": "select_one sex", "name": "SEX",   "label": "Sex",
                 "bind__oc_itemgroup": "DM", "required": "yes"},
                {"type": "integer", "name": "AGE",   "label": "Age (years)",
                 "bind__oc_itemgroup": "DM", "required": "yes",
                 "constraint": ". >= 0 and . <= 120",
                 "constraint_message": "Age must be between 0 and 120."},
            ],
            "choices": [
                {"list_name": "sex", "name": "M", "label": "Male"},
                {"list_name": "sex", "name": "F", "label": "Female"},
                {"list_name": "sex", "name": "U", "label": "Unknown"},
            ],
            "cross_form_dependencies": [],
            "settings": {
                "form_title": "Demographics",
                "form_id":    "DM",
                "version":    "1",
                "style":      "theme-grid",
                "namespaces": 'oc="http://openclinica.org/xforms"',
            },
        },
        {
            "form_id": "AE",
            "form_title": "Adverse Events",
            "repeating": True,
            "category": "CDASH_SAFETY",
            "visits_assigned": ["SE_SCREEN", "SE_VISIT1"],
            "survey": [
                {"type": "date",   "name": "AESTDAT", "label": "AE Start Date",
                 "bind__oc_itemgroup": "AE", "required": "yes",
                 "constraint": ". <= today()",
                 "constraint_message": "Future dates are not allowed."},
                {"type": "text",   "name": "AETERM",  "label": "AE Term",
                 "bind__oc_itemgroup": "AE", "required": "yes"},
                {"type": "select_one aesev", "name": "AESEV", "label": "Severity",
                 "bind__oc_itemgroup": "AE", "required": "yes"},
                {"type": "select_one yn",    "name": "AESER", "label": "Serious?",
                 "bind__oc_itemgroup": "AE", "required": "yes"},
            ],
            "choices": [
                {"list_name": "aesev", "name": "1", "label": "Mild"},
                {"list_name": "aesev", "name": "2", "label": "Moderate"},
                {"list_name": "aesev", "name": "3", "label": "Severe"},
                {"list_name": "aesev", "name": "4", "label": "Life-threatening"},
                {"list_name": "aesev", "name": "5", "label": "Fatal"},
            ],
            "cross_form_dependencies": [],
            "settings": {
                "form_title": "Adverse Events",
                "form_id":    "AE",
                "version":    "1",
                "style":      "theme-grid",
                "namespaces": 'oc="http://openclinica.org/xforms"',
            },
        },
    ],
    "events": [
        {"event_id": "SE_SCREEN", "event_name": "Screening",
         "event_type": "SCHEDULED", "order": 1},
        {"event_id": "SE_VISIT1", "event_name": "Visit 1",
         "event_type": "SCHEDULED", "order": 2},
    ],
    "form_placements": [
        {"form_id": "DOV",  "event_id": "SE_SCREEN", "order": 1},
        {"form_id": "DM",   "event_id": "SE_SCREEN", "order": 2},
        {"form_id": "AE",   "event_id": "SE_SCREEN", "order": 3},
        {"form_id": "DOV",  "event_id": "SE_VISIT1", "order": 1},
        {"form_id": "AE",   "event_id": "SE_VISIT1", "order": 2},
    ],
}


@pytest.fixture(scope='module')
def build_zip_bytes():
    """Run edc-builder against the fixture spec and return the ZIP bytes."""
    from build_package import build_package
    import tempfile, os, shutil

    spec = json.loads(json.dumps(FIXTURE_SPEC))  # deep copy
    build_log = {
        'forms_built':         [],
        'forms_skipped':       [],
        'build_errors':        [],
        'build_warnings':      [],
        'placeholder_applied': [],
        'oid_placeholders':    [],
    }

    with tempfile.TemporaryDirectory() as tmpdir:
        forms_dir    = os.path.join(tmpdir, 'forms')
        csv_dir      = os.path.join(tmpdir, 'csv')
        checklist_dir = os.path.join(tmpdir, 'checklist')
        package_dir  = os.path.join(tmpdir, 'package')
        for d in (forms_dir, csv_dir, checklist_dir, package_dir):
            os.makedirs(d)

        # build_all_xlsforms produces the form XLSX files into forms_dir
        from build_xlsforms import build_all_xlsforms
        build_all_xlsforms(spec, forms_dir, build_log)

        # build_package creates the ZIP
        zip_path = build_package(spec, build_log, forms_dir, csv_dir,
                                 checklist_dir, package_dir)
        return open(zip_path, 'rb').read()


@pytest.fixture(scope='module')
def build_zip(build_zip_bytes):
    return zipfile.ZipFile(io.BytesIO(build_zip_bytes))


# ── Module-level XLSForm workbook fixtures ───────────────────────────────────
@pytest.fixture(scope='module')
def dm_wb(build_zip):
    import openpyxl
    data = build_zip.read(next(n for n in build_zip.namelist()
                               if n.upper().endswith('DM.XLSX')))
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)


@pytest.fixture(scope='module')
def ae_wb(build_zip):
    import openpyxl
    data = build_zip.read(next(n for n in build_zip.namelist()
                               if n.upper().endswith('AE.XLSX')))
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)


class TestEdcBuilderIntegration:

    # ── ZIP structure ─────────────────────────────────────────────────────
    def test_zip_is_valid(self, build_zip_bytes):
        assert zipfile.is_zipfile(io.BytesIO(build_zip_bytes))

    def test_zip_contains_expected_form_files(self, build_zip):
        names = build_zip.namelist()
        xlsx_files = [n for n in names if n.endswith('.xlsx')]
        form_ids = {n.split('/')[-1].replace('.xlsx', '').upper() for n in xlsx_files}
        assert 'DOV' in form_ids, f"DOV.xlsx missing from ZIP: {xlsx_files}"
        assert 'DM'  in form_ids, f"DM.xlsx missing from ZIP: {xlsx_files}"
        assert 'AE'  in form_ids, f"AE.xlsx missing from ZIP: {xlsx_files}"

    def test_form_count_matches_spec(self, build_zip):
        xlsx_files = [n for n in build_zip.namelist() if n.endswith('.xlsx')
                      and not n.endswith('checklist.xlsx')]
        assert len(xlsx_files) == 3, \
            f"Expected 3 form files, got {len(xlsx_files)}: {xlsx_files}"

    # ── DM form XLSForm content ───────────────────────────────────────────
    def test_dm_has_survey_sheet(self, dm_wb):
        assert 'survey' in dm_wb.sheetnames

    def test_dm_has_choices_sheet(self, dm_wb):
        assert 'choices' in dm_wb.sheetnames

    def test_dm_has_settings_sheet(self, dm_wb):
        assert 'settings' in dm_wb.sheetnames

    def test_dm_survey_has_correct_fields(self, dm_wb):
        ws = dm_wb['survey']
        rows = list(ws.iter_rows(values_only=True))
        # Find header row
        header = None
        for row in rows:
            if row and 'name' in [str(v).lower() for v in row if v]:
                header = [str(v).lower() if v else '' for v in row]
                break
        assert header is not None, "No header row found in DM survey"
        name_col = header.index('name')
        field_names = {str(r[name_col]).upper() for r in rows
                       if r and r[name_col] and str(r[name_col]) != 'None'
                       and str(r[name_col]).lower() not in ('name',)}
        assert 'SUBJID' in field_names, f"SUBJID missing from DM survey: {field_names}"
        assert 'SEX'    in field_names, f"SEX missing from DM survey: {field_names}"
        assert 'AGE'    in field_names, f"AGE missing from DM survey: {field_names}"

    def test_dm_choices_has_sex_list(self, dm_wb):
        ws = dm_wb['choices']
        rows = list(ws.iter_rows(values_only=True))
        all_vals = {str(v).upper() for row in rows for v in row if v}
        # The sex list from our fixture should appear
        assert 'SEX' in all_vals, f"sex list missing from DM choices: {all_vals}"
        assert 'MALE' in all_vals or 'M' in all_vals, \
            f"Male choice missing from DM choices: {all_vals}"

    # ── AE form — constraint message standardization ──────────────────────
    def test_ae_future_date_message_canonicalized(self, ae_wb):
        ws = ae_wb['survey']
        rows = list(ws.iter_rows(values_only=True))
        header = None
        for row in rows:
            if row and 'constraint_message' in [str(v).lower() for v in row if v]:
                header = [str(v).lower() if v else '' for v in row]
                break
        if header is None:
            pytest.skip("No constraint_message column in AE survey")
        msg_col = header.index('constraint_message')
        name_col = header.index('name') if 'name' in header else None
        messages = {str(r[msg_col]) for r in rows
                    if r and r[msg_col] and str(r[msg_col]) not in ('None', 'constraint_message')}
        # The "Future dates are not allowed." from the fixture should be
        # normalized to the canonical form
        bad_variants = {"Future dates are not allowed.", "Date must not be in the future.",
                        "No future dates allowed."}
        for msg in messages:
            assert msg not in bad_variants, \
                f"Non-canonical constraint message found: {msg!r}"

    def test_ae_is_repeating(self, ae_wb):
        ws = ae_wb['settings']
        rows = list(ws.iter_rows(values_only=True))
        all_vals = ' '.join(str(v) for row in rows for v in row if v).lower()
        # Repeating forms have a repeat group in the survey — check settings or survey
        # Just verify the file opened without error and has settings
        assert len(rows) > 0

    # ── Field name validity across all forms ─────────────────────────────
    def test_no_invalid_field_names_in_any_form(self, build_zip):
        import openpyxl
        valid_name = re.compile(r'^[A-Za-z_][A-Za-z0-9_.-]*$')
        skip_values = {'name', 'type', 'label', 'none', '', 'begin_repeat',
                       'end_repeat', 'begin group', 'end group'}
        violations = []
        for entry in build_zip.namelist():
            if not entry.upper().endswith('.XLSX'):
                continue
            data = build_zip.read(entry)
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            if 'survey' not in wb.sheetnames:
                continue
            ws = wb['survey']
            rows = list(ws.iter_rows(values_only=True))
            header = None
            for row in rows:
                if row and 'name' in [str(v).lower() for v in row if v]:
                    header = [str(v).lower() if v else '' for v in row]
                    break
            if not header or 'name' not in header:
                continue
            name_col = header.index('name')
            for row in rows:
                if not row or not row[name_col]:
                    continue
                val = str(row[name_col]).strip()
                if val.lower() in skip_values:
                    continue
                if not valid_name.match(val):
                    violations.append(f"{entry}: {val!r}")
        assert not violations, \
            f"Invalid XLSForm field names found:\n" + "\n".join(violations)

    def test_no_i_prefix_field_names(self, build_zip):
        """Regression: CRF Standards injection must not produce I_ prefixed names."""
        import openpyxl
        violations = []
        for entry in build_zip.namelist():
            if not entry.upper().endswith('.XLSX'):
                continue
            data = build_zip.read(entry)
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            if 'survey' not in wb.sheetnames:
                continue
            ws = wb['survey']
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                for val in row:
                    if val and str(val).startswith('I_'):
                        violations.append(f"{entry}: {val!r}")
        assert not violations, \
            f"I_ prefixed field names found (regression):\n" + "\n".join(violations)

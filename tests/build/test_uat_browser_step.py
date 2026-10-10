"""
The browser (Playwright) UAT step, playwright_uat.py.

Run of 2026-10-10: 204 "field left blank" cases failed with "No required-field error shown" on four forms. They
were scored on a page where the form had never opened (several participant pages opened at once raced on the
sign-in state cookie and landed on InvalidStateCookieWarning). The step also opened the first form of a name
whatever the visit, never emptied a field that held a loaded value, never set a gate value, and took the form
title for an error message, so its passes were unreliable too.

Part 1 tests the decisions without a browser. Part 2 drives the real step against fixture pages served from
this file (no network): a participant page with visits and a common-event table, and a form whose markup follows
what OpenClinica's Enketo renders (question classes, message elements, /data/NAME inputs, a hidden date input
behind its widget). The fixture's script is a small model of the form engine: it validates a field on its change
event, as the live form was seen to do for required fields. It is skipped when no Chromium is installed.
"""
import asyncio
import io
import json
import os
import sys

import openpyxl
import pytest

_REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, _REPO)

import playwright_uat as pw  # noqa: E402
import uat_loader  # noqa: E402


# ── Part 1: decisions ─────────────────────────────────────────────────────────

def test_case_kinds():
    kind = pw._classify_pw_row
    assert kind({"Load_Value": "(leave blank)", "Expected Result": "Required-field error shown."}) == "leave_blank"
    assert kind({"Load_Value": "GATE=Y", "Expected Result": "Field 'X' is VISIBLE."}) == "visibility"
    assert kind({"Load_Value": "43", "Expected Result": "Constraint fires. Message: m"}) == "constraint"
    assert kind({"Load_Value": "x", "Expected Result": "Value is accepted"}) is None


def test_field_name_prefers_the_item_name_column():
    assert pw._field_name({"Item_Name": "DIN", "Item_OID": "I_DEMOG_DIN_2803"}) == "DIN"
    assert pw._field_name({"Item_OID": "I_DEMOG_AGE_DISP"}) == "AGE_DISP"


def test_gates_are_parsed_from_the_input():
    assert pw._parse_gates("GATE=Y") == [("GATE", "Y")]
    assert pw._parse_gates("A=1, B=two words") == [("A", "1"), ("B", "two words")]
    assert pw._parse_gates("GATE=(blank)") == [("GATE", "")]
    assert pw._parse_gates("Set referenced fields so expression is true (e.g. 1)") is None
    assert pw._parse_gates("") is None


def test_value_under_test():
    assert pw._value_under_test({"Load_Value": "43"}) == "43"
    assert pw._value_under_test({"Load_Value": "A=1, then this date=2026-01-15", "Setup_Steps": "[{}]",
                                 "Test_Value": "2026-01-15"}) == "2026-01-15"
    assert pw._value_under_test({"Load_Value": "A=1, then this date=2026-01-15"}) is None    # no structured setup
    assert pw._value_under_test({"Load_Value": "A=1"}) is None


def test_simple_gate_reads_one_comparison_only():
    assert pw._simple_gate(" /data/GRP/GATE ='Y'") == ("GATE", "Y")
    assert pw._simple_gate("/data/GATE = 'N'") == ("GATE", "N")
    assert pw._simple_gate(" /data/A ='Y' or  /data/B ='Z'") is None
    assert pw._simple_gate(" /data/A  != 'Baseline'") is None
    assert pw._simple_gate("") is None


def test_required_result_comes_from_the_field_itself():
    assert pw._judge_required({"required": True, "message": "This field is required"}) == \
        ("Pass", "Required message shown: This field is required")
    result, actual = pw._judge_required({"required": False, "message": ""})
    assert result == "Fail" and "No required-field error shown" in actual
    assert pw._judge_required({})[0] == "Fail"


def test_constraint_results():
    shown = {"constraint": True, "message": "Must be 1 to 42"}
    assert pw._judge_constraint(True, shown, "43") == ("Pass", "Constraint message shown for '43': Must be 1 to 42")
    assert pw._judge_constraint(True, {"constraint": False}, "43")[0] == "Fail"
    assert pw._judge_constraint(False, {"constraint": False}, "21") == \
        ("Pass", "Value '21' accepted — no constraint message")
    assert pw._judge_constraint(False, shown, "21")[0] == "Fail"


def test_a_field_kept_on_screen_with_the_no_longer_relevant_flag_counts_as_not_shown():
    flagged = {"visible": True, "irrelevant_flag": True}
    assert pw._judge_visibility(False, flagged, "GATE=N")[0] == "Pass"
    assert pw._judge_visibility(True, flagged, "GATE=Y")[0] == "Fail"
    assert pw._judge_visibility(True, {"visible": True}, "GATE=Y") == ("Pass", "Field shown with GATE=Y")
    result, actual = pw._judge_visibility(True, {"visible": False}, "GATE=Y")
    assert result == "Fail" and actual == "Field hidden with GATE=Y — expected shown"


def test_a_multi_step_case_is_tested_on_its_own_participant(monkeypatch):
    assert pw._case_participant({"Participant_ID": "UAT-P005", "Setup_Steps": "[{}]"}) == "UAT-P005"
    assert pw._case_participant({"Participant_ID": "UAT-P003"}) == "UAT-P001"
    monkeypatch.setenv("UAT_SETUP_OWN_PARTICIPANT", "0")
    assert pw._case_participant({"Participant_ID": "UAT-P005", "Setup_Steps": "[{}]"}) == "UAT-P001"


def test_study_names_come_from_the_metadata():
    xml = ('<ODM xmlns="http://www.cdisc.org/ns/odm/v1.3"><Study OID="S_T"><GlobalVariables><StudyName>T</StudyName>'
           '</GlobalVariables><MetaDataVersion OID="v1">'
           '<StudyEventDef OID="SE_ONE" Name="Visit One (Day 0)" Repeating="No" Type="Unscheduled"/>'
           '<StudyEventDef OID="SE_COMMON" Name="Common" Repeating="Yes" Type="Common"/>'
           '<FormDef OID="F_AA" Name="Form A"/></MetaDataVersion></Study></ODM>')
    names = uat_loader._parse_study_names(xml)
    assert names["events"] == {"SE_ONE": "Visit One (Day 0)", "SE_COMMON": "Common"}
    assert names["forms"] == {"F_AA": "Form A"} and names["common_events"] == {"SE_COMMON"}
    assert names["study_name"] == "T"
    assert uat_loader._parse_study_names("")["events"] == {}


def test_legacy_engine_is_still_importable():
    import playwright_uat_legacy
    assert callable(playwright_uat_legacy.run_playwright_uat)


# ── Part 2: the step against fixture pages ────────────────────────────────────

_HOST = "uatfixture"
_TOP = f"https://{_HOST}.eu.openclinica.io/OpenClinica/ParticipantDetailsPage"
_APP = f"https://{_HOST}.eu.openclinica.io/study-runner-ui/index.html"
_FORM = "https://form.eu.openclinica.io/edit/fs/c/i/"

_TOP_HTML = f'<html><body><iframe src="{_APP}" style="width:1400px;height:900px"></iframe></body></html>'

_APP_HTML = """<html><body>
<p-accordion><p-accordiontab id="visits"><div id="visitsElement">
  <div class="p-grid visit"><div class="visit-event"><h5 class="event-name" title="Visit Two (Week 4)">Visit Two (W...</h5></div>
    <div class="visit-forms"><div class="visit-form visit-form-card" title="Edit Form A" data-open="two-F_AA">Form A</div></div></div>
  <div class="p-grid visit"><div class="visit-event"><h5 class="event-name" title="Visit One (Day 0)">Visit One (D...</h5></div>
    <div class="visit-forms"><div class="visit-form visit-form-card" title="Edit Form A" data-open="one-F_AA">Form A</div>
      <div class="visit-form visit-form-card" title="Edit Form B" data-open="one-F_BB">Form B</div></div></div>
</div></p-accordiontab>
<p-accordiontab><a class="p-accordion-header-link" aria-expanded="false" id="commonhdr">Common</a>
  <div id="commonbody" style="display:none">
   <p-table id="SE_COMMONF_CC-table"><table><tbody><tr id="SE_COMMONF_CC"><td>
     <button type="button" class="p-button form-menu">...</button></td></tr></tbody></table></p-table>
   <p-table id="SE_COMMONF_DD-table"><div><button class="p-button add-new-button">Add New</button></div>
     <table><tbody><tr><td>No data available in table</td></tr></tbody></table></p-table>
  </div></p-accordiontab></p-accordion>
<script>
function openForm(id) { var f = document.createElement('iframe'); f.src = '__FORM__' + id;
  f.style.cssText = 'width:1200px;height:700px'; document.body.appendChild(f); }
document.querySelectorAll('.visit-form-card').forEach(function (c) {
  c.addEventListener('click', function () { openForm(c.getAttribute('data-open')); }); });
document.getElementById('commonhdr').addEventListener('click', function () {
  this.setAttribute('aria-expanded', 'true'); document.getElementById('commonbody').style.display = 'block'; });
document.querySelector('.form-menu').addEventListener('click', function () {
  setTimeout(function () { var a = document.createElement('a'); a.className = 'p-menuitem-link'; a.textContent = 'Edit';
    a.addEventListener('click', function () { openForm('common-F_CC'); }); document.body.appendChild(a); }, 300); });
document.querySelector('.add-new-button').addEventListener('click', function () { openForm('common-F_DD'); });
</script></body></html>""".replace("__FORM__", _FORM)

# A form in the markup OpenClinica's Enketo renders. values: what the participant already holds.
_FORM_HTML = """<html><head><style>
.or-required-msg,.or-constraint-msg,.or-relevant-msg{display:none}
.invalid-required .or-required-msg,.invalid-constraint .or-constraint-msg,.invalid-relevant .or-relevant-msg{display:block}
.or-branch.disabled{display:none} .or-branch.disabled.invalid-relevant{display:block}
</style></head><body><form class="or theme-grid">
<h3 id="form-title">UAT-1: A long form title that is not an error</h3>
<fieldset class="question simple-select"><fieldset><legend><span class="question-label active">Performed?</span></legend>
  <label><input type="radio" name="/data/GRP/PERF" value="N" data-required="true()"><span>No</span></label>
  <label><input type="radio" name="/data/GRP/PERF" value="Y" data-required="true()"><span>Yes</span></label></fieldset>
  <span class="or-required-msg active">This field is required</span></fieldset>
<label class="question non-select"><span class="question-label active">Date</span>
  <div class="widget date"><input class="ignore" type="text"></div>
  <input type="text" name="/data/GRP/VISDAT" data-required="true()" data-type-xml="date" style="display:none">
  <span class="or-required-msg active">This field is required</span></label>
<label class="question non-select"><span class="question-label active">Optional note</span>
  <input type="text" name="/data/GRP/NOTE"></label>
<label class="question non-select"><span class="question-label active">Volume</span>
  <input type="number" name="/data/GRP/VOL" data-min="1" data-max="42">
  <span class="or-constraint-msg active">Volume must be 1 to 42</span></label>
<label class="question non-select or-branch disabled"><span class="question-label active">Detail</span>
  <input type="text" name="/data/GRP/DETAIL" data-relevant=" /data/GRP/PERF ='Y'" data-required="true()">
  <span class="or-required-msg active">This field is required</span>
  <span class="or-relevant-msg active">No longer relevant</span></label>
<label class="question non-select or-branch disabled"><span class="question-label active">More</span>
  <input type="text" name="/data/GRP/MORE" data-relevant=" /data/GRP/DETAIL ='x'"></label>
<label class="question non-select or-branch disabled"><span class="question-label active">By calc</span>
  <input type="text" name="/data/GRP/BYCALC" data-relevant=" /data/GRP/TPT  != 'Baseline'"></label>
<label class="calculation non-select"><input type="hidden" name="/data/GRP/TPT" data-calculate="x" value="Baseline"></label>
<select name="/data/GRP/PICK" class="question"><option value="">...</option><option value="A">A</option></select>
</form><script>
var VALUES = __VALUES__;
window.__changes = [];
function els(n) { return Array.prototype.slice.call(document.querySelectorAll('[name="/data/GRP/' + n + '"]')); }
function val(n) { var e = els(n); if (!e.length) return '';
  if (e[0].type === 'radio') { var c = e.filter(function (x) { return x.checked; }); return c.length ? c[0].value : ''; }
  return e[0].value; }
function q(e) { return e.closest('.question') || e.parentElement; }
function relevance() {
  document.querySelectorAll('[data-relevant]').forEach(function (e) {
    var m = e.getAttribute('data-relevant').match(/\\/data\\/GRP\\/(\\w+)\\s*(!=|=)\\s*'([^']*)'/);
    var on = m[2] === '=' ? val(m[1]) === m[3] : val(m[1]) !== m[3];
    var b = q(e); b.classList.toggle('disabled', !on);
    b.classList.toggle('invalid-relevant', !on && e.value !== '');
  }); }
function validate(e) {
  var b = q(e), n = e.name.split('/').pop(), v = val(n);
  b.classList.toggle('invalid-required', e.hasAttribute('data-required') && v === '');
  if (e.hasAttribute('data-min')) b.classList.toggle('invalid-constraint',
      v !== '' && (Number(v) < Number(e.getAttribute('data-min')) || Number(v) > Number(e.getAttribute('data-max')))); }
Object.keys(VALUES).forEach(function (n) { els(n).forEach(function (e) {
  if (e.type === 'radio') e.checked = (e.value === VALUES[n]); else e.value = VALUES[n]; }); });
relevance(); relevance();
document.addEventListener('change', function (ev) { var e = ev.target; if (!e.name) return;
  window.__changes.push(e.name.split('/').pop() + '=' + val(e.name.split('/').pop()));
  setTimeout(function () { validate(e); relevance(); relevance(); }, 120); });
</script></body></html>"""

_BUILD_HTML = "<html><body><script>localStorage.setItem('jhi-idtoken', 'x')</script></body></html>"
_VALUES = {"one-F_AA": {"PERF": "N", "VISDAT": "2026-01-20", "VOL": "21"}, "two-F_AA": {}, "one-F_BB": {},
           "common-F_CC": {"PERF": "Y", "DETAIL": "abc"}, "common-F_DD": {}}


def _routes(state):
    """A context hook serving the fixture pages; state["race"] = how many participant-page loads still land on
    the sign-in race page; state["opened"] collects the forms the step opened."""
    async def hook(context):
        async def handle(route):
            url = route.request.url
            if url.startswith(_TOP):
                if state.get("race", 0) > 0:
                    state["race"] -= 1
                    return await route.fulfill(status=302, headers={
                        "location": f"https://{_HOST}.eu.openclinica.io/OpenClinica/InvalidStateCookieWarning"})
                return await route.fulfill(content_type="text/html", body=_TOP_HTML)
            if "InvalidStateCookieWarning" in url:
                return await route.fulfill(content_type="text/html", body="<html><body><input><input></body></html>")
            if url.startswith(_APP):
                return await route.fulfill(content_type="text/html", body=_APP_HTML)
            if url.startswith(_FORM):
                form = url[len(_FORM):]
                state.setdefault("opened", []).append(form)
                return await route.fulfill(content_type="text/html", body=_FORM_HTML.replace(
                    "__VALUES__", json.dumps(_VALUES.get(form, {}))))
            if ".build.openclinica.io" in url:
                return await route.fulfill(content_type="text/html", body=_BUILD_HTML)
            return await route.abort()
        await context.route("**/*", handle)
    return hook


def _chromium_available():
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            p.chromium.launch(headless=True).close()
        return True
    except Exception:
        return False


browser = pytest.mark.skipif(not _chromium_available(), reason="no Chromium for Playwright on this machine")

_COLS = ["UAT Case ID", "Status", "Scenario", "Expected Result", "Actual Result", "Test Result", "Execution Date",
         "Notes", "Study_Event_OID", "Form_OID", "Item_OID", "Participant_ID", "Load_Value", "Item_Name",
         "Setup_Steps", "Test_Value"]
_REQ = "Required-field error shown. Form does not save."
_FIRES = "Constraint fires. Message: Volume must be 1 to 42"
_OK = "No constraint error. Form saves."


def _case(uid, form, event, item, load, expected, **extra):
    row = {"UAT Case ID": uid, "Form_OID": form, "Study_Event_OID": event, "Item_Name": item, "Load_Value": load,
           "Expected Result": expected, "Participant_ID": "UAT-P001", "Actual Result": "Not Testable via ODM",
           "Test Result": "Not Run"}
    row.update(extra)
    return row


_CASES = [
    _case("UAT-001", "F_AA", "SE_ONE", "PERF", "(leave blank)", _REQ),            # radio that holds a loaded value
    _case("UAT-002", "F_AA", "SE_ONE", "VISDAT", "(leave blank)", _REQ),          # date behind its widget
    _case("UAT-003", "F_AA", "SE_ONE", "NOTE", "(leave blank)", _REQ),            # not required: a real failure
    _case("UAT-004", "F_AA", "SE_ONE", "VOL", "43", _FIRES),
    _case("UAT-005", "F_AA", "SE_ONE", "VOL", "0", _FIRES),
    _case("UAT-006", "F_AA", "SE_ONE", "DETAIL", "PERF=Y", "Field 'Detail' is VISIBLE."),
    _case("UAT-007", "F_AA", "SE_ONE", "DETAIL", "PERF=N", "Field 'Detail' is HIDDEN."),
    _case("UAT-008", "F_AA", "SE_ONE", "DETAIL", "(leave blank)", _REQ),          # hidden: its gate is set first
    _case("UAT-009", "F_AA", "SE_ONE", "MORE", "DETAIL=x", "Field 'More' is VISIBLE."),   # gate behind a gate
    _case("UAT-010", "F_AA", "SE_ONE", "BYCALC", "TPT=0", "Field 'By calc' is VISIBLE."),  # calculated gate
    _case("UAT-011", "F_AA", "SE_ONE", "MISSING", "(leave blank)", _REQ),
    _case("UAT-012", "F_AA", "SE_ONE", "PICK", "ZZZ_INVALID", _FIRES),
    _case("UAT-013", "F_AA", "SE_ONE", "VOL", "A=1, then this value=5", _OK),     # multi-step without setup
    _case("UAT-014", "F_AA", "SE_TWO", "PERF", "(leave blank)", _REQ),            # same form, another visit
    _case("UAT-015", "F_CC", "SE_COMMON", "DETAIL", "PERF=N", "Field 'Detail' is HIDDEN."),   # holds a value
    _case("UAT-016", "F_DD", "SE_COMMON", "VISDAT", "(leave blank)", _REQ),       # no entry yet: Add New
    _case("UAT-017", "F_BB", "SE_THREE", "PERF", "(leave blank)", _REQ),          # visit not scheduled
    _case("UAT-018", "F_EE", "SE_ONE", "PERF", "(leave blank)", _REQ),            # form not on that visit
    _case("UAT-019", "F_IE", "SE_ONE", "PERF", "(leave blank)", _REQ),            # skipped by default, as before
    _case("UAT-020", "F_AA", "SE_ONE", "VOL", "5", _OK, **{"Test Result": "Blocked",
                                                           "Actual Result": "Blocked: setup not in place — x"}),
    _case("UAT-021", "F_AA", "SE_ONE", "VOL", "21", "No required-field error. Form saves.",
          **{"Test Result": "Pass", "Actual Result": "21"}),                      # scored by the data import
    _case("UAT-022", "F_AA", "SE_ONE", "VOL", "A=1, then this value=5", _OK,
          **{"Setup_Steps": '[{"form": "F_AA"}]', "Test_Value": "5", "Participant_ID": "UAT-P005"}),
    _case("UAT-023", "F_AA", "SE_ONE", "VOL", "A=1, then this value=5", _OK,
          **{"Setup_Steps": '[{"form": "F_AA"}]', "Test_Value": "5", "Participant_ID": "UAT-P009"}),
]


def _workbook(cases):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UAT_Cases"
    ws.append(["UAT cases"])
    ws.append(_COLS)
    for c in cases:
        ws.append([c.get(k, "") for k in _COLS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _results(b):
    ws = openpyxl.load_workbook(io.BytesIO(b))["UAT_Cases"]
    rows = list(ws.iter_rows(values_only=True))
    return {r[0]: dict(zip(rows[1], r)) for r in rows[2:] if r and r[0]}


def _run(tmp_path, cases, state, **env):
    session = tmp_path / "s.json"
    session.write_text('{"cookies": [], "origins": []}')
    stamp = {"UAT-P001": {"oc_oid": "SS_1"}, "UAT-P005": {"oc_oid": "SS_5"}}
    out = asyncio.run(pw.run_playwright_uat(
        _workbook(cases), _HOST, "SS_1", "user@example.com", stamp,
        fo_titles={"F_AA": "Form A", "F_BB": "Form B", "F_CC": "Form C", "F_DD": "Form D", "F_EE": "Form E"},
        ev_titles={"SE_ONE": "Visit One (Day 0)", "SE_TWO": "Visit Two (Week 4)", "SE_THREE": "Visit Three"},
        common_events={"SE_COMMON"}, session_path=str(session), _context_hook=_routes(state)))
    return _results(out)


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    state = {"race": 1}          # the warm-up load lands on the race page once, then the session exists
    return _run(tmp_path_factory.mktemp("pw"), _CASES, state), state


@browser
def test_required_message_is_read_from_the_field_after_it_is_emptied(run):
    res, _ = run
    for uid in ("UAT-001", "UAT-002"):
        assert res[uid]["Test Result"] == "Pass", res[uid]["Actual Result"]
        assert res[uid]["Actual Result"] == "Required message shown: This field is required"
        assert res[uid]["Notes"] == "Playwright" and res[uid]["Status"] == "Pass"


@browser
def test_a_field_without_a_required_rule_fails_and_the_form_title_is_never_an_error(run):
    res, _ = run
    assert res["UAT-003"]["Test Result"] == "Fail"
    assert "No required-field error shown" in res["UAT-003"]["Actual Result"]
    assert not any("long form title" in str(r["Actual Result"]) for r in res.values())


@browser
def test_constraint_message_is_read_from_the_field(run):
    res, _ = run
    for uid, value in (("UAT-004", "43"), ("UAT-005", "0")):
        assert res[uid]["Test Result"] == "Pass", res[uid]["Actual Result"]
        assert res[uid]["Actual Result"] == f"Constraint message shown for '{value}': Volume must be 1 to 42"


@browser
def test_gate_values_are_set_in_the_form_before_the_field_is_read(run):
    res, _ = run
    assert (res["UAT-006"]["Test Result"], res["UAT-006"]["Actual Result"]) == ("Pass", "Field shown with PERF=Y")
    assert (res["UAT-007"]["Test Result"], res["UAT-007"]["Actual Result"]) == ("Pass", "Field hidden with PERF=N")
    assert res["UAT-009"]["Test Result"] == "Pass", res["UAT-009"]["Actual Result"]     # PERF=Y, then DETAIL=x
    assert res["UAT-015"]["Test Result"] == "Pass"
    assert "not relevant (flagged because it holds a value)" in res["UAT-015"]["Actual Result"]


@browser
def test_a_hidden_required_field_is_brought_into_view_for_its_blank_test(run):
    res, _ = run
    assert res["UAT-008"]["Test Result"] == "Pass", res["UAT-008"]["Actual Result"]


@browser
def test_what_cannot_be_run_is_blocked_with_its_reason_never_failed(run):
    res, _ = run
    reasons = {
        "UAT-010": "gate TPT is calculated or read-only (now 'Baseline')",
        "UAT-011": "field MISSING not found in the published form",
        "UAT-012": "'ZZZ_INVALID' is not an option of PICK",
        "UAT-013": "not a single value the browser can enter",
        "UAT-017": "visit SE_THREE (Visit Three) is not scheduled for this participant",
        "UAT-018": "form F_EE is not on visit Visit One (Day 0)",
        "UAT-023": "participant UAT-P009 was not created",
    }
    for uid, reason in reasons.items():
        assert res[uid]["Test Result"] == "Blocked", (uid, res[uid])
        assert res[uid]["Actual Result"].startswith("Blocked: ") and reason in res[uid]["Actual Result"], res[uid]


@browser
def test_the_form_is_opened_in_the_visit_the_case_names(run):
    res, state = run
    assert res["UAT-014"]["Test Result"] == "Pass"          # Visit Two's copy is empty: required shows
    opened = state["opened"]
    assert opened.count("one-F_AA") == 2 and "two-F_AA" in opened       # UAT-P001's and the multi-step participant's
    assert "common-F_CC" in opened and "common-F_DD" in opened          # Edit from the row menu; Add New when empty
    assert res["UAT-016"]["Test Result"] == "Pass"
    assert res["UAT-022"]["Test Result"] == "Pass"          # multi-step: its Test_Value, on its own participant
    assert res["UAT-022"]["Actual Result"] == "Value '5' accepted — no constraint message"


@browser
def test_cases_already_scored_or_blocked_are_left_alone(run):
    res, _ = run
    assert (res["UAT-020"]["Test Result"], res["UAT-020"]["Notes"]) == ("Blocked", None)
    assert res["UAT-020"]["Actual Result"] == "Blocked: setup not in place — x"
    assert (res["UAT-021"]["Test Result"], res["UAT-021"]["Actual Result"], res["UAT-021"]["Notes"]) == \
        ("Pass", "21", None)
    assert res["UAT-019"]["Test Result"] == "Skip"


@browser
def test_every_value_a_case_changed_is_put_back():
    """Lower level: one form, the cases above, then the values the participant held are still there."""
    from playwright.async_api import async_playwright

    async def go():
        async with async_playwright() as p:
            b = await p.chromium.launch(headless=True)
            page = await (await b.new_context()).new_page()
            await page.set_content(_FORM_HTML.replace("__VALUES__", json.dumps(_VALUES["one-F_AA"])))
            for c in _CASES[:9]:
                await pw._run_case(page.main_frame, c, pw._classify_pw_row(c))
            await asyncio.sleep(0.4)
            held = {n: (await page.evaluate(pw._JS_STATE, n))["value"] for n in ("PERF", "VISDAT", "VOL", "DETAIL")}
            changes = await page.evaluate("window.__changes")
            await b.close()
            return held, changes

    held, changes = asyncio.run(go())
    assert held == {"PERF": "N", "VISDAT": "2026-01-20", "VOL": "21", "DETAIL": ""}
    assert "PERF=" in changes and "PERF=N" in changes and "VOL=43" in changes      # emptied, entered, restored


@browser
def test_a_form_that_never_opens_blocks_its_cases(tmp_path):
    """Every load of the participant page lands on the race page: nothing is scored on it."""
    res = _run(tmp_path, _CASES[:3], {"race": 99})
    for uid in ("UAT-001", "UAT-002", "UAT-003"):
        assert res[uid]["Test Result"] == "Blocked", res[uid]
        assert "participant page did not open (landed on InvalidStateCookieWarning)" in res[uid]["Actual Result"]
    assert pw.last_status["ran"] is True


def test_without_a_saved_login_nothing_is_run_and_the_reason_is_kept(tmp_path):
    wb = _workbook(_CASES[:2])
    out = asyncio.run(pw.run_playwright_uat(wb, _HOST, "SS_1", "nobody@example.com", {},
                                            session_path=str(tmp_path / "missing.json")))
    assert out == wb and pw.last_status == {"ran": False, "reason": "no saved browser login for nobody@example.com"}


# ── Part 3: what takes a question off the screen (run of 2026-10-10: 415 cases "hidden by its display rule") ──
# The form engine puts the questions in the page first and keeps the whole form hidden behind its loader until
# it has finished loading. In a container whose system locale is C/POSIX the browser reported a language the
# engine rejects, so it never finished: every question was in the page, none was on screen, and each was
# reported as hidden by a rule. The fixture below follows that structure (loader before .main, pages, a group
# with a header, a group and a question made non-relevant) with a small model of the engine: a change event is
# acted on only when the value really changed.

_FORM2_URL = "https://form.eu.openclinica.io/edit/fs/c/i/fixture2"
_FORM2_HTML = """<html><head><title>OpenClinica</title><style>
.main-loader ~ .main{display:none}
.or-required-msg{display:none} .invalid-required .or-required-msg{display:block}
.or-branch.disabled,.calculation{display:none}
.or.pages [role=page]{display:none} .or.pages [role=page].current{display:block}
.shut > :not(h4){display:none} .offscreen{display:none}
</style></head><body class="iframe"><div class="main-loader"></div><div class="main"><article class="paper">
<form class="or theme-grid __PAGES__">
<section class="or-group" role="page" name="/data/P1">
  <label class="question non-select"><span class="question-label active">First</span>
    <input type="text" name="/data/P1/FIRST" data-required="true()" data-type-xml="string">
    <span class="or-required-msg active">This field is required</span></label>
  <fieldset class="question simple-select"><fieldset><legend>Done?</legend>
    <label><input type="radio" name="/data/P1/DONE" value="N" data-type-xml="string"><span>No</span></label>
    <label><input type="radio" name="/data/P1/DONE" value="Y" data-type-xml="string"><span>Yes</span></label></fieldset></fieldset>
  <label class="calculation non-select"><input type="hidden" name="/data/P1/CALC" data-calculate="1 + 1" value="2"></label>
</section>
<section class="or-group" role="page" name="/data/P2">
  <label class="question non-select"><span class="question-label active">Second</span>
    <input type="number" name="/data/P2/SECOND" data-required="true()" data-type-xml="int">
    <span class="or-required-msg active">This field is required</span></label>
  <section class="or-group shut" name="/data/P2/BOX"><h4>More</h4>
    <label class="question non-select"><span class="question-label active">Inside</span>
      <input type="text" name="/data/P2/BOX/INSIDE" data-required="true()" data-type-xml="string">
      <span class="or-required-msg active">This field is required</span></label></section>
  <section class="or-group or-branch disabled" name="/data/P2/GATED" data-relevant=" /data/P1/DONE ='Y'">
    <label class="question non-select"><span class="question-label active">In a gated group</span>
      <input type="text" name="/data/P2/GATED/DEEP" data-required="true()" data-type-xml="string">
      <span class="or-required-msg active">This field is required</span></label></section>
  <label class="question non-select or-branch disabled"><span class="question-label active">Either</span>
    <input type="text" name="/data/P2/EITHER" data-required="true()" data-type-xml="string"
           data-relevant=" /data/P1/CALC ='9' or  /data/P1/DONE ='N'">
    <span class="or-required-msg active">This field is required</span></label>
  <label class="question non-select or-branch disabled"><span class="question-label active">Both</span>
    <input type="text" name="/data/P2/BOTH" data-relevant=" /data/P1/DONE ='Y' and  /data/P1/FIRST ='a'"></label>
  <div class="offscreen"><label class="question non-select"><span class="question-label active">Odd</span>
    <input type="text" name="/data/P2/ODD" data-required="true()" data-type-xml="string"></label></div>
</section>
<button type="button" class="previous-page">Back</button><button type="button" class="next-page">Next</button>
</form></article></div><script>
var MODEL = {};
function named(n) { return Array.prototype.slice.call(document.querySelectorAll('input')).filter(function (e) {
  return e.name.split('/').pop() === n; }); }
function shown(n) { var e = named(n); if (!e.length) return '';
  if (e[0].type === 'radio') { var c = e.filter(function (x) { return x.checked; }); return c.length ? c[0].value : ''; }
  return e[0].value; }
function holds(expr) { return expr.split(/\\s+or\\s+/).some(function (alt) {
  return alt.split(/\\s+and\\s+/).every(function (c) { var m = c.match(/(\\w+)\\s*=\\s*'([^']*)'/);
    return shown(m[1]) === m[2]; }); }); }
function relevance() { document.querySelectorAll('[data-relevant]').forEach(function (e) {
  var b = e.matches('section') ? e : e.closest('.question'); b.classList.toggle('disabled', !holds(e.getAttribute('data-relevant'))); }); }
document.addEventListener('change', function (ev) { var e = ev.target, n = e.name.split('/').pop(), v = shown(n);
  if ((MODEL[n] || '') === v) return;                 // nothing changed: the engine does nothing
  MODEL[n] = v;
  e.closest('.question').classList.toggle('invalid-required', e.hasAttribute('data-required') && v === '');
  relevance(); });
var pages = Array.prototype.slice.call(document.querySelectorAll('[role=page]'));
function turn(by) { var i = pages.findIndex(function (p) { return p.classList.contains('current'); });
  if (by > 0 && pages[i].querySelector('.invalid-required')) return;      // Next refuses a page with an error
  if (pages[i + by]) { pages[i].classList.remove('current'); pages[i + by].classList.add('current'); } }
if (document.querySelector('form.pages')) pages[0].classList.add('current');
document.querySelector('.next-page').addEventListener('click', function () { turn(1); });
document.querySelector('.previous-page').addEventListener('click', function () { turn(-1); });
document.querySelector('h4').addEventListener('click', function () { this.parentElement.classList.toggle('shut'); });
relevance();
var WANT = '__LANGUAGE__';
setTimeout(function () {
  if (WANT && navigator.language !== WANT) { var a = document.createElement('div'); a.className = 'vex-content';
    a.textContent = 'Incorrect locale information provided'; document.body.appendChild(a); return; }
  document.querySelector('.main-loader').remove(); document.title = 'P-1: Fixture'; }, __LOAD_MS__);
</script></body></html>"""


def _form2(pages=True, load_ms=0, language=""):
    return (_FORM2_HTML.replace("__PAGES__", "pages" if pages else "").replace("__LOAD_MS__", str(load_ms))
            .replace("__LANGUAGE__", language))


def _on_form2(html, work, **context):
    """Serve the fixture at a form URL, wait for it the way the step does, then run work(frame)."""
    from playwright.async_api import async_playwright

    async def go():
        async with async_playwright() as p:
            b = await p.chromium.launch(headless=True)
            ctx = await b.new_context(**context)
            page = await ctx.new_page()

            async def serve(route):
                await route.fulfill(content_type="text/html", body=html)
            await page.route("**/*", serve)
            await page.goto(_FORM2_URL)
            try:
                frame = await pw._wait_form_frame(page)
                return await work(frame) if work else frame is not None
            finally:
                await b.close()
    return asyncio.run(go())


def _blank(item):
    return {"Item_Name": item, "Load_Value": "(leave blank)", "Expected Result": _REQ}


async def _cases(frame, *rows):
    return [await pw._run_case(frame, r, pw._classify_pw_row(r)) for r in rows]


def test_a_rule_with_alternatives_gives_each_gate_and_anything_else_gives_none():
    assert pw._gate_options(" /data/G/A ='Y' or  /data/B ='Baseline'") == [("A", "Y"), ("B", "Baseline")]
    assert pw._gate_options(" /data/A ='Y'") == [("A", "Y")]
    assert pw._gate_options(" /data/A ='Y' and  /data/B ='1'") == []
    assert pw._gate_options(" /data/A  != 'Baseline'") == [] and pw._gate_options("") == []


def test_the_value_entered_and_removed_in_an_empty_field_follows_its_type():
    assert pw._blank_probe({"type": "radio", "options": ["N", "Y"]}) == "N"
    assert pw._blank_probe({"type": "select", "options": ["", "A"]}) == "A"
    assert pw._blank_probe({"type": "text", "xml_type": "int"}) == "1"
    assert pw._blank_probe({"type": "text", "xml_type": "date"}) == "2020-01-01"
    assert pw._blank_probe({"type": "text", "xml_type": "string"}) == "x"
    assert pw._blank_probe({"type": "radio", "options": []}) is None
    assert pw._blank_probe({"type": "text", "xml_type": "geopoint"}) is None


def test_the_reason_a_field_is_not_on_screen_names_what_hides_it():
    assert "calculated item" in pw._not_on_screen("X", {"hidden_by": "calculated"})
    assert "still loading" in pw._not_on_screen("X", {"hidden_by": "loading"})
    assert "another page" in pw._not_on_screen("X", {"hidden_by": "page", "hidden_el": "section.or-group"})
    assert "collapsed section" in pw._not_on_screen("X", {"hidden_by": "collapsed", "hidden_el": "section.or-group"})
    text = pw._not_on_screen("X", {"hidden_by": "other", "hidden_el": "div.offscreen"})
    assert "no display rule hides it" in text and "div.offscreen" in text


def test_the_browser_language_is_set_explicitly(monkeypatch):
    monkeypatch.delenv("PW_LOCALE", raising=False)
    assert pw._locale() == "en-US"
    monkeypatch.setenv("PW_LOCALE", "de-CH")
    assert pw._locale() == "de-CH"


@browser
def test_cases_wait_for_the_form_to_finish_loading(monkeypatch):
    """The questions are in the page 2.5 s before the form is on screen. Nothing is read before that."""
    async def work(frame):
        early = await frame.evaluate("() => document.title")
        return early, await _cases(frame, _blank("FIRST"))
    title, results = _on_form2(_form2(pages=False, load_ms=2500), work)
    assert title == "P-1: Fixture"
    assert results == [("Pass", "Required message shown: This field is required")]


@browser
def test_a_form_that_never_finishes_loading_blocks_with_what_the_form_says(monkeypatch):
    monkeypatch.setenv("PW_FORM_LOAD_S", "5")
    with pytest.raises(pw._Blocked) as e:
        _on_form2(_form2(pages=False, language="xx-XX"), None, locale="en-US")
    assert "was still loading after 5s" in str(e.value)
    assert "the form says: Incorrect locale information provided" in str(e.value)


@browser
def test_the_form_loads_with_the_language_the_step_sets(monkeypatch):
    """The fixture finishes loading only for the language given: the step's own setting reaches the page."""
    monkeypatch.setenv("PW_LOCALE", "de-CH")
    monkeypatch.setenv("PW_FORM_LOAD_S", "5")
    assert _on_form2(_form2(pages=False, language="de-CH"), None, locale=pw._locale()) is True


@browser
def test_a_question_on_another_page_or_in_a_collapsed_section_is_opened_not_blocked():
    async def work(frame):
        first = await pw._state(frame, "SECOND")
        results = await _cases(frame, _blank("SECOND"), _blank("INSIDE"))
        return first, results
    first, results = _on_form2(_form2(), work)
    assert (first["visible"], first["rule_hidden"], first["hidden_by"]) == (False, False, "page")
    assert results == [("Pass", "Required message shown: This field is required")] * 2      # empty fields: entered, removed


@browser
def test_the_page_is_opened_even_when_next_refuses_to_leave_a_page_with_an_error():
    async def work(frame):
        await pw._set(frame, "FIRST", "a")
        await pw._set(frame, "FIRST", "")                    # page one now shows a required error
        return await _cases(frame, _blank("SECOND"))
    assert _on_form2(_form2(), work) == [("Pass", "Required message shown: This field is required")]


@browser
def test_hidden_by_a_rule_means_non_relevant_itself_or_through_its_group():
    async def work(frame):
        deep, either = await pw._state(frame, "DEEP"), await pw._state(frame, "EITHER")
        results = await _cases(
            frame, _blank("DEEP"),                           # its group's rule: DONE=Y is set, then it is tested
            _blank("EITHER"),                                # first alternative is calculated: the second is set
            {"Item_Name": "DEEP", "Load_Value": "DONE=N", "Expected Result": "Field 'x' is HIDDEN."},
            {"Item_Name": "DEEP", "Load_Value": "DONE=Y", "Expected Result": "Field 'x' is VISIBLE."},
            {"Item_Name": "SECOND", "Load_Value": "DONE=N", "Expected Result": "Field 'x' is HIDDEN."},
            _blank("BOTH"), _blank("CALC"), _blank("ODD"))
        return deep, either, results, await pw._state(frame, "DONE")
    deep, either, results, done = _on_form2(_form2(), work)
    assert deep["rule_hidden"] and deep["relevant"].strip() == "/data/P1/DONE ='Y'"
    assert either["rule_hidden"] and either["hidden_by"] == "rule"
    assert results[0][0] == "Pass" and results[1][0] == "Pass", results[:2]
    assert results[2] == ("Pass", "Field hidden with DONE=N")
    assert results[3] == ("Pass", "Field shown with DONE=Y")
    assert results[4] == ("Fail", "Field shown with DONE=N — expected hidden")      # on another page is not hidden
    assert results[5][0] == "Blocked" and "not field = value comparisons the test can set" in results[5][1]
    assert results[6][0] == "Blocked" and "calculated item" in results[6][1]
    assert results[7][0] == "Blocked" and "no display rule hides it" in results[7][1] and "div.offscreen" in results[7][1]
    assert done["value"] == ""                               # every gate value was put back

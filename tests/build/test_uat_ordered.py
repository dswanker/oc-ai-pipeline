"""
Ordered (multi-step) UAT: a case that needs a prerequisite value (a cross-form source such as the consent date, or a
same-form start date) gets structured Setup_Steps; the loader imports those values into UAT-P001 BEFORE the browser
step, confirms them by read-back, and marks a case Not Run ("Setup failed: ...") when its setup did not store.
"""
import contextlib
import io
import json
import os
import sys
import tempfile

import openpyxl

_HERE = os.path.dirname(__file__)
_REPO = os.path.abspath(os.path.join(_HERE, "..", ".."))
sys.path.insert(0, _REPO)
sys.path.insert(0, os.path.join(_REPO, "tests", "conventions"))
sys.path.insert(0, os.path.join(_REPO, "skills", "dvs-specification", "scripts"))

import uat_loader  # noqa: E402
from conventions_engine import apply_conventions  # noqa: E402
from extract_dvs_from_forms import extract_dvs_data  # noqa: E402
from generate_dvs import build_dvs  # noqa: E402
from test_consent_floor import _spec as consent_spec  # noqa: E402


def _dvs_bytes():
    spec = consent_spec()
    for f in spec["forms"]:
        f["form_title"] = {"ICF": "Informed Consent", "F_AE": "Adverse Events", "DM": "Demographics"}[f["form_id"]]
        f.setdefault("visits_assigned", {"F_AE": ["SE_COMMON"], "DM": ["SE_SCREEN"]}.get(f["form_id"], []))
    with contextlib.redirect_stdout(io.StringIO()):
        apply_conventions(spec, study_id="T", customer_subdomain="", client_name="")
    forms = {f"{f['form_id']}.xlsx": {"survey": [{("bind::oc:itemgroup" if k == "bind__oc_itemgroup" else k): v
                                                    for k, v in r.items()} for r in f["survey"]],
                                       "choices": f.get("choices") or []} for f in spec["forms"]}
    path = os.path.join(tempfile.mkdtemp(), "dvs.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        build_dvs(extract_dvs_data(spec, {"forms": forms}), path)
    return open(path, "rb").read()


def _cases(b):
    ws = openpyxl.load_workbook(io.BytesIO(b))["UAT_Cases"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:6]) if r and "UAT Case ID" in r)
    return [dict(zip(rows[hi], r)) for r in rows[hi + 1:] if r and r[0]]


def test_multi_step_cases_carry_structured_setup():
    cases = _cases(_dvs_bytes())
    floor = [c for c in cases if c["Item_Name"] == "ONSET" and "then" in str(c["Load_Value"])]
    assert floor, "consent-floor cases expected"
    for c in floor:
        steps = json.loads(c["Setup_Steps"])
        assert steps[0]["form"] == "F_ICF" and steps[0]["item"] == "ICFDAT" and steps[0]["event"] == "SE_SCREEN"
        assert c["Test_Value"] and c["Participant_ID"] == "UAT-P001"


def test_setup_rows_are_loaded_into_uat_p001_once():
    rows = uat_loader._parse_uat_cases(_dvs_bytes())
    setup = uat_loader._setup_rows(rows)
    keys = [(r["Form_OID"], r["Item_Name"]) for r in setup]
    assert ("F_ICF", "ICFDAT") in keys and len(keys) == len(set(keys))
    assert all(r["Participant_ID"] == "UAT-P001" for r in setup)
    xml = uat_loader._build_odm_xml("S_T", "SITE", "SS_P1", "UAT-P001", setup)
    assert "ICFDAT" in xml and setup[0]["Load_Value"] in xml        # imported before the browser step


def test_confirmation_marks_failed_setup_not_run_and_records_confirmed_setup():
    dvs = _dvs_bytes()
    rows = uat_loader._parse_uat_cases(dvs)
    setup = uat_loader._setup_rows(rows)
    stamp = {"UAT-P001": {"participant_key": "SS_P1"}}
    stored = {("SS_P1", r["Study_Event_OID"].upper(), r["Form_OID"].upper(), r["Item_Group_OID"].upper(),
               r["Item_OID"].upper()): r["Load_Value"] for r in setup}
    ok = uat_loader._confirm_setup(setup, stamp, stored)
    assert all(v[0] for v in ok.values())
    marked = _cases(uat_loader._mark_setup_results(dvs, ok))
    floor = [c for c in marked if c["Setup_Steps"]]
    assert floor and all(str(c["Preconditions"]).startswith("Setup: ") and "loaded and confirmed" in c["Preconditions"]
                         for c in floor)
    assert all(str(c["Preconditions"]).startswith("Setup: F_ICF.ICFDAT=") for c in floor if c["Item_Name"] == "ONSET")
    bad = uat_loader._confirm_setup(setup, stamp, {})                    # nothing stored
    failed = _cases(uat_loader._mark_setup_results(dvs, bad))
    assert all(c["Test Result"] == "Not Run" and str(c["Actual Result"]).startswith("Setup failed")
               for c in failed if c["Setup_Steps"])


def test_evidence_shows_setup_then_test():
    dvs = _dvs_bytes()
    rows = uat_loader._parse_uat_cases(dvs)
    setup = uat_loader._setup_rows(rows)
    stamp = {"UAT-P001": {"participant_key": "SS_P1"}}
    stored = {("SS_P1", r["Study_Event_OID"].upper(), r["Form_OID"].upper(), r["Item_Group_OID"].upper(),
               r["Item_OID"].upper()): r["Load_Value"] for r in setup}
    marked = uat_loader._mark_setup_results(dvs, uat_loader._confirm_setup(setup, stamp, stored))
    wb = openpyxl.load_workbook(io.BytesIO(marked))
    ws, hrow, col = uat_loader._case_sheet(wb)
    for r in range(hrow + 1, ws.max_row + 1):                              # simulate the browser result
        if ws.cell(row=r, column=col["Setup_Steps"]).value:
            ws.cell(row=r, column=col["Notes"], value="Playwright")
            ws.cell(row=r, column=col["Test Result"], value="Pass")
            ws.cell(row=r, column=col["Actual Result"], value="Constraint: Date cannot be before the informed consent date.")
    buf = io.BytesIO(); wb.save(buf)
    final, _ = uat_loader._finalize_test_methods(buf.getvalue(), "ran")
    done = [c for c in _cases(final) if c.get("Setup_Steps")]
    assert done and all(c["Evidence"].startswith("Setup: ") and "loaded and confirmed. Test: entered" in c["Evidence"]
                        for c in done)
    assert all(c["Evidence"].startswith("Setup: F_ICF.ICFDAT=") for c in done if c["Item_Name"] == "ONSET")

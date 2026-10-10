"""
Two UAT generator / loader rules, both from the run of 2026-10-10.

1. A test value has the item's real data type. The year of a partial date is an integer item; the sample picked by
   field-name prefix was a full date and the import rejected it (errorCode.valueTypeMismatch).
2. A multi-step test (Setup_Steps) has a participant of its own. Its setup value (a consent date) was loaded into
   UAT-P001 and then replaced by the happy-path load of the same item for another case, so the dependent cases ran
   against the wrong value. The reservation that should have prevented it predicted the wrong item OID for a
   cross-form source. After loading, a setup value that is not in place makes the case Blocked, never Pass or Fail.
"""
import contextlib
import io
import json
import os
import sys

import openpyxl
import pytest

_REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, _REPO)
sys.path.insert(0, os.path.join(_REPO, "skills", "dvs-specification", "scripts"))

import extract_dvs_from_forms as gen  # noqa: E402
import uat_loader  # noqa: E402
from generate_dvs import build_dvs  # noqa: E402

_CF = ("instance('clinicaldata')/ODM/ClinicalData/SubjectData/StudyEventData/FormData[@FormOID='F_CONS']"
       "/ItemGroupData[@OpenClinica:ItemGroupName='CONS']/ItemData[@OpenClinica:ItemName='CONSDAT']/@Value")

_STRUCT = {
    "study_meta": {"protocol_number": "TEST-001"},
    "forms": [
        {"form_id": "CONS", "form_title": "Informed Consent", "visits_assigned": ["SE_SCR"]},
        {"form_id": "EXAM", "form_title": "Physical Examination", "visits_assigned": ["SE_SCR"]},
        {"form_id": "LATE", "form_title": "Later Assessment", "visits_assigned": ["SE_W4"]},
        {"form_id": "MED", "form_title": "Medication", "visits_assigned": ["SE_COMMON"]},
    ],
}
_FORMS = {"forms": {
    "CONS.xlsx": {"survey": [
        {"type": "date", "name": "CONSDAT", "label": "Consent date", "bind::oc:itemgroup": "CONS",
         "required": "yes", "constraint": ". <= today()"}], "choices": []},
    "EXAM.xlsx": {"survey": [
        {"type": "calculate", "name": "CONSDAT_CF", "calculation": _CF, "bind::oc:external": "clinicaldata"},
        {"type": "date", "name": "EXAMDAT", "label": "Exam date", "bind::oc:itemgroup": "EXAM", "required": "yes",
         "constraint": "(. <= today()) and (. = '' or ${CONSDAT_CF} = '' or . >= ${CONSDAT_CF})"}], "choices": []},
    "LATE.xlsx": {"survey": [
        {"type": "calculate", "name": "CONSDAT_CF", "calculation": _CF, "bind::oc:external": "clinicaldata"},
        {"type": "text", "name": "LATENOTE", "label": "Note", "bind::oc:itemgroup": "LATE", "required": "yes"},
        {"type": "date", "name": "LATEDAT", "label": "Assessment date", "bind::oc:itemgroup": "LATE",
         "constraint": ". = '' or ${CONSDAT_CF} = '' or . >= ${CONSDAT_CF}"}], "choices": []},
    "MED.xlsx": {"survey": [
        {"type": "integer", "name": "CMSTDAT_YEAR", "label": "Start Year", "bind::oc:itemgroup": "MED",
         "required": "yes"},
        {"type": "integer", "name": "CMSTDAT_MONTH", "label": "Start Month", "bind::oc:itemgroup": "MED",
         "required": "yes"},
        {"type": "integer", "name": "CMSTDAT_DAY", "label": "Start Day", "bind::oc:itemgroup": "MED",
         "required": "yes"},
        {"type": "decimal", "name": "CMDOSU", "label": "Dose (number)", "bind::oc:itemgroup": "MED",
         "required": "yes"}], "choices": []},
}}


def _generate():
    with contextlib.redirect_stdout(io.StringIO()):
        return gen.extract_dvs_data(_STRUCT, _FORMS)


@pytest.fixture(scope="module")
def cases():
    return _generate()["uat_cases"]


def _happy(cases, item):
    return next(c for c in cases if c["Item_Name"] == item and c["Scenario"].startswith("Happy path: field populated"))


# ── 1. test values in the item's data type ────────────────────────────────────

def test_year_month_and_day_parts_of_a_partial_date_get_integers(cases):
    assert _happy(cases, "CMSTDAT_YEAR")["Load_Value"] == "2026"       # was "2026-01-15"
    assert _happy(cases, "CMSTDAT_MONTH")["Load_Value"] == "1"
    assert _happy(cases, "CMSTDAT_DAY")["Load_Value"] == "15"


def test_a_numeric_item_never_gets_a_text_sample(cases):
    assert _happy(cases, "CMDOSU")["Load_Value"] == "1.5"              # the name prefix gave the unit "mg"
    for c in cases:
        if c["Form_OID"] == "F_MED" and c["Load_Value"] != "(leave blank)":
            float(c["Load_Value"])


def test_typed_sample_rules():
    assert gen._typed_sample("2026-01-15", "integer", "ANY_YEAR") == "2026"
    assert gen._typed_sample("2026-01-15", "integer", "ANYDAT") == "1"          # no date part in the name
    assert gen._typed_sample("13.5", "integer", "X") == "13"
    assert gen._typed_sample("13.5", "decimal", "X") == "13.5"
    assert gen._typed_sample("2026-01-15", "date", "ANY_YEAR") == "2026-01-15"  # other types are untouched
    assert gen._typed_sample("Sample", "text", "X") == "Sample"


def test_typed_samples_kill_switch(monkeypatch):
    monkeypatch.setenv("UAT_TYPED_SAMPLES", "0")
    assert gen._typed_sample("2026-01-15", "integer", "ANY_YEAR") == "2026-01-15"


# ── 2. multi-step tests have their own participant ────────────────────────────

def _multi(cases):
    return [c for c in cases if c.get("Setup_Steps")]


def test_setup_step_names_the_slot_the_source_form_really_uses(cases):
    source = next(c for c in cases if c["Item_Name"] == "CONSDAT")
    for c in _multi(cases):
        st = json.loads(c["Setup_Steps"])[0]
        assert (st["form"], st["item_oid"], st["item_group"]) == \
            ("F_CONS", source["Item_OID"], source["Item_Group_OID"])


def test_no_other_case_shares_the_participant_of_a_multi_step_test(cases):
    multi = _multi(cases)
    assert multi, "multi-step cases expected"
    own = {c["Participant_ID"] for c in multi}
    assert "UAT-P001" not in own
    assert not [c["UAT Case ID"] for c in cases if not c.get("Setup_Steps") and c["Participant_ID"] in own]


def test_steps_of_one_ordered_test_share_a_participant_and_different_tests_do_not(cases):
    by_check = {}
    for c in _multi(cases):
        by_check.setdefault(c["Related Check ID"], set()).add(c["Participant_ID"])
    assert len(by_check) == 2                                          # EXAMDAT and LATEDAT against the consent date
    assert all(len(p) == 1 for p in by_check.values())
    assert len(set.union(*by_check.values())) == 2


def test_two_cases_never_write_one_item_on_one_participant(cases):
    """Loaded rows plus setup rows: one value per (participant, form, event, item)."""
    rows = [dict(c) for c in cases]
    writes = {}
    for r in rows + uat_loader._setup_rows(rows):
        if r.get("_setup") or uat_loader._plain_loadable(r):
            key = (r["Participant_ID"], r["Form_OID"], r["Study_Event_OID"], r["Item_Name"])
            writes.setdefault(key, set()).add(str(r["Load_Value"]))
    assert writes and all(len(v) == 1 for v in writes.values())


def test_kill_switch_keeps_the_setup_on_uat_p001_with_its_slot_reserved(monkeypatch):
    monkeypatch.setenv("UAT_SETUP_OWN_PARTICIPANT", "0")
    cases = _generate()["uat_cases"]
    assert all(c["Participant_ID"] == "UAT-P001" for c in _multi(cases))
    loaded_on_p1 = [c for c in cases if c["Item_Name"] == "CONSDAT" and c["Participant_ID"] == "UAT-P001"
                    and uat_loader._plain_loadable(c)]
    assert not loaded_on_p1                                            # the reservation now matches the real slot
    assert all(r["Participant_ID"] == "UAT-P001" for r in uat_loader._setup_rows([dict(c) for c in cases]))


# ── loader: an older workbook, read-back, Blocked ─────────────────────────────

def _old_style_rows():
    """The 2026-10-10 pattern: both dependent cases on UAT-P001 next to a happy-path load of the same item."""
    step = json.dumps([{"form": "F_CONS", "event": "SE_SCR", "item_group": "IG_CONS_CONS", "item": "CONSDAT",
                        "item_oid": "I_CONS_CONSDAT", "value": "2026-01-20"}])
    base = {"Study_Event_OID": "SE_SCR", "Event_Repeat_Key": "1", "Participant_ID": "UAT-P001"}
    return [
        dict(base, **{"UAT Case ID": "UAT-001", "Related Check ID": "DVS-001", "Form_OID": "F_CONS",
                      "Item_Group_OID": "IG_INFOR_CONS", "Item_OID": "I_INFOR_CONSDAT", "Item_Name": "CONSDAT",
                      "Load_Value": "2026-10-10", "Scenario": "Happy path: today's date",
                      "Expected Result": "No constraint error. Form saves."}),
        dict(base, **{"UAT Case ID": "UAT-002", "Related Check ID": "DVS-001", "Form_OID": "F_CONS",
                      "Item_Group_OID": "IG_INFOR_CONS", "Item_OID": "I_INFOR_CONSDAT", "Item_Name": "CONSDAT",
                      "Load_Value": "2026-10-09", "Scenario": "Happy path: a date in the past",
                      "Expected Result": "No constraint error. Form saves.", "Participant_ID": "UAT-P002"}),
        dict(base, **{"UAT Case ID": "UAT-003", "Related Check ID": "DVS-002", "Form_OID": "F_EXAM",
                      "Item_Group_OID": "IG_PHYSI_EXAM", "Item_OID": "I_PHYSI_EXAMDAT", "Item_Name": "EXAMDAT",
                      "Load_Value": "CONSDAT on form F_CONS=2026-01-20, then this date=2026-01-25",
                      "Setup_Steps": step, "Test_Value": "2026-01-25", "Scenario": "Happy path",
                      "Expected Result": "No constraint error. Form saves."}),
        dict(base, **{"UAT Case ID": "UAT-004", "Related Check ID": "DVS-002", "Form_OID": "F_EXAM",
                      "Item_Group_OID": "IG_PHYSI_EXAM", "Item_OID": "I_PHYSI_EXAMDAT", "Item_Name": "EXAMDAT",
                      "Load_Value": "CONSDAT on form F_CONS=2026-01-20, then this date=2026-01-15",
                      "Setup_Steps": step, "Test_Value": "2026-01-15", "Scenario": "Sad path",
                      "Expected Result": "Constraint fires. Message: x"}),
    ]


def test_loader_moves_a_shared_multi_step_test_to_its_own_participant():
    rows = _old_style_rows()
    moved = uat_loader._isolate_setup_cases(rows)
    assert moved == {"UAT-003": "UAT-P003", "UAT-004": "UAT-P003"}
    assert rows[0]["Participant_ID"] == "UAT-P001"
    setup = uat_loader._setup_rows(rows)
    assert [(r["Participant_ID"], r["Form_OID"], r["Item_Name"], r["Load_Value"]) for r in setup] == \
        [("UAT-P003", "F_CONS", "CONSDAT", "2026-01-20")]
    assert uat_loader._isolate_setup_cases(rows) == {}                 # already on its own: nothing to do


def test_a_case_tested_in_another_visit_gets_that_visit_opened_on_its_participant():
    rows = [dict(c) for c in _generate()["uat_cases"]]
    setup = uat_loader._setup_rows(rows)
    late_pid = next(c["Participant_ID"] for c in rows if c["Item_Name"] == "LATEDAT" and c.get("Setup_Steps"))
    mine = [r for r in setup if r["Participant_ID"] == late_pid]
    assert {(r["Form_OID"], r["Study_Event_OID"]) for r in mine} == {("F_CONS", "SE_SCR"), ("F_LATE", "SE_W4")}
    seed = next(r for r in mine if r["Form_OID"] == "F_LATE")
    assert seed["Item_Name"] == "LATENOTE" and seed["Scenario"].startswith("Seed:")
    exam_pid = next(c["Participant_ID"] for c in rows if c["Item_Name"] == "EXAMDAT" and c.get("Setup_Steps"))
    assert [r["Form_OID"] for r in setup if r["Participant_ID"] == exam_pid] == ["F_CONS"]   # same visit: no seed


def _workbook(rows):
    cols = ["UAT Case ID", "Status", "Related Check ID", "Scenario", "Preconditions", "Expected Result",
            "Actual Result", "Test Result", "Execution Date", "Notes", "Site_OID", "Participant_Key",
            "Study_Event_OID", "Form_OID", "Item_Group_OID", "Item_OID", "Participant_ID", "Load_Value",
            "Item_Name", "Setup_Steps", "Test_Value"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UAT_Cases"
    ws.append(["UAT cases"])
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read(b):
    ws = openpyxl.load_workbook(io.BytesIO(b))["UAT_Cases"]
    rows = list(ws.iter_rows(values_only=True))
    return {r[0]: dict(zip(rows[1], r)) for r in rows[2:] if r and r[0]}


def _stored(setup, stamp, value=None):
    return {(stamp[r["Participant_ID"]]["participant_key"], r["Study_Event_OID"].upper(), r["Form_OID"].upper(),
             r["Item_Group_OID"].upper(), r["Item_OID"].upper()): (value or r["Load_Value"]) for r in setup}


def test_confirmed_setup_is_recorded_and_the_case_stays_runnable():
    rows = _old_style_rows()
    wb = uat_loader._write_participant_ids(_workbook(rows), uat_loader._isolate_setup_cases(rows))
    setup = uat_loader._setup_rows(rows)
    stamp = {"UAT-P001": {"participant_key": "K1"}, "UAT-P003": {"participant_key": "K3"}}
    conf = uat_loader._confirm_setup(setup, stamp, _stored(setup, stamp), rows)
    assert all(ok for ok, _ in conf.values())
    after = _read(uat_loader._mark_setup_results(wb, conf))
    for uid in ("UAT-003", "UAT-004"):
        assert after[uid]["Participant_ID"] == "UAT-P003"
        assert after[uid]["Preconditions"] == "Setup: F_CONS.CONSDAT=2026-01-20 loaded and confirmed"
        assert not after[uid]["Test Result"]


def test_setup_not_in_place_blocks_the_case_and_never_scores_it():
    rows = _old_style_rows()
    wb = uat_loader._write_participant_ids(_workbook(rows), uat_loader._isolate_setup_cases(rows))
    setup = uat_loader._setup_rows(rows)
    stamp = {"UAT-P003": {"participant_key": "K3"}}
    for stored, text in (({}, "not stored"), (_stored(setup, stamp, "2026-03-03"), "read back '2026-03-03'")):
        conf = uat_loader._confirm_setup(setup, stamp, stored, rows)
        after = _read(uat_loader._mark_setup_results(wb, conf))
        for uid in ("UAT-003", "UAT-004"):
            assert after[uid]["Test Result"] == "Blocked" and after[uid]["Status"] == "Blocked"
            assert after[uid]["Actual Result"].startswith("Blocked: setup not in place") and \
                text in after[uid]["Actual Result"]
    blocked = uat_loader._mark_setup_results(wb, uat_loader._confirm_setup(setup, stamp, {}, rows))
    rescored = _read(uat_loader._evaluate_uat_cases(blocked, stamp, {}, {}))
    assert rescored["UAT-003"]["Test Result"] == "Blocked"             # a later scoring pass leaves it alone
    final, summary = uat_loader._finalize_test_methods(blocked, "ran")
    done = _read(final)
    assert done["UAT-004"]["Test Result"] == "Blocked" and done["UAT-004"]["Test Method"] == uat_loader.METHOD_BROWSER
    assert done["UAT-004"]["Evidence"].startswith("Blocked: setup not in place")
    assert summary[uat_loader.METHOD_BROWSER]["Blocked"] == 2


def test_overwritten_setup_names_the_case_that_loaded_the_other_value(monkeypatch):
    """With the setup on a shared participant (the old assignment), the value read back is another case's load."""
    monkeypatch.setenv("UAT_SETUP_OWN_PARTICIPANT", "0")
    rows = _old_style_rows()
    assert uat_loader._isolate_setup_cases(rows) == {}
    setup = uat_loader._setup_rows(rows)
    stamp = {"UAT-P001": {"participant_key": "K1"}}
    conf = uat_loader._confirm_setup(setup, stamp, _stored(setup, stamp, "2026-10-10"), rows)
    (ok, note), = conf.values()
    assert not ok
    assert "setup value overwritten by another test case (UAT-001 loaded '2026-10-10'" in note
    assert "OpenClinica holds" not in note
    after = _read(uat_loader._mark_setup_results(_workbook(rows), conf))
    assert after["UAT-003"]["Test Result"] == "Blocked"


def test_generated_workbook_round_trips_through_the_loader(tmp_path):
    path = str(tmp_path / "dvs.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        build_dvs(_generate(), path)
    rows = uat_loader._parse_uat_cases(open(path, "rb").read())
    assert uat_loader._isolate_setup_cases(rows) == {}                 # the generator's assignment is kept

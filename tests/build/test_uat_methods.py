"""
UAT results are trustworthy evidence: browser (Playwright) results are never overwritten by the data-import
scoring, and every case states how it was tested (Test Method) and what proves it (Evidence).
Uses the saved 2026-10-06 BioIVT Detroit run (docs/uat_analysis). No network, nothing is created.
"""
import collections
import io
import os
import sys

import openpyxl

_HERE = os.path.dirname(__file__)
sys.path.insert(0, _HERE)
import uat_loader  # noqa: E402
from test_uat_rescore import _load_run, _cases  # noqa: E402


def _set_playwright_results(wb_bytes, n):
    """Simulate the browser step: record a result on the first n cases the evaluation left untestable via ODM."""
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws, hrow, col = uat_loader._case_sheet(wb)
    done = []
    for r in range(hrow + 1, ws.max_row + 1):
        if len(done) >= n:
            break
        if ws.cell(row=r, column=col["Actual Result"]).value == "Not Testable via ODM":
            ws.cell(row=r, column=col["Test Result"], value="Pass")
            ws.cell(row=r, column=col["Actual Result"], value="Constraint fired: Date cannot be in the future.")
            ws.cell(row=r, column=col["Notes"], value="Playwright")
            done.append(ws.cell(row=r, column=col["UAT Case ID"]).value)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue(), done


def _scored(label="Detroit"):
    wb_bytes, clinical, failures = _load_run(label)[:3]
    stamp = {}
    return uat_loader._evaluate_uat_cases(wb_bytes, stamp, clinical, failures), clinical, failures


def test_browser_results_survive_a_second_scoring_pass():
    scored, clinical, failures = _scored()
    with_pw, ids = _set_playwright_results(scored, 25)
    assert len(ids) == 25
    rescored = uat_loader._evaluate_uat_cases(with_pw, {}, clinical, failures)
    after = {c["UAT Case ID"]: c for c in _cases(rescored)}
    assert all(after[i]["Test Result"] == "Pass" and after[i]["Notes"] == "Playwright" for i in ids)


def test_every_case_has_a_method_and_evidence():
    scored, _c, _f = _scored()
    with_pw, ids = _set_playwright_results(scored, 25)
    final, summary = uat_loader._finalize_test_methods(with_pw, "was skipped: no saved browser login for x@y.com")
    cases = _cases(final)
    assert all(c.get("Test Method") for c in cases) and all(c.get("Evidence") for c in cases)
    by = collections.Counter(c["Test Method"] for c in cases)
    assert by[uat_loader.METHOD_BROWSER] == 25
    assert by[uat_loader.METHOD_ODM] >= 190                      # the re-scored data-import passes and fails
    odm = next(c for c in cases if c["Test Method"] == uat_loader.METHOD_ODM)
    assert "loaded" in odm["Evidence"] and "OpenClinica stored" in odm["Evidence"]
    manual = [c for c in cases if c["Test Method"] == uat_loader.METHOD_MANUAL]
    assert manual and all(c["Test Result"] == "Not Run" for c in manual)
    assert any("no saved browser login" in c["Evidence"] for c in manual)
    assert sum(sum(v.values()) for v in summary.values()) == len(cases)
    assert "Data import (ODM):" in uat_loader._method_summary_text(summary)


def test_finalize_is_idempotent():
    scored, _c, _f = _scored()
    once, s1 = uat_loader._finalize_test_methods(scored, "was skipped")
    twice, s2 = uat_loader._finalize_test_methods(once, "was skipped")
    assert s1 == s2
    assert [(c["Test Method"], c["Evidence"]) for c in _cases(once)] == \
        [(c["Test Method"], c["Evidence"]) for c in _cases(twice)]

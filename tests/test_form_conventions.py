"""Customer form conventions (form_conventions.py): SAE_FORM and DEATH_DETAILS_FORM, asked like every CQ question
and applied to the final forms. Synthetic forms only."""
import copy
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "skills", "protocol-analysis", "scripts"))
import form_conventions as fc
import protocol_forms as pf
import standards_match as sm
from tests.standards import fixtures as fx

SAE_Q = "Collect Serious Adverse Events on a separate SAE form?"
DEATH_Q = "Collect death details (cause of death, autopsy) on a separate Death Details form?"


@pytest.fixture(autouse=True)
def env(monkeypatch):
    monkeypatch.setitem(sm._CDASH, "v", (None, {}, {}, set()))
    monkeypatch.delenv("FORM_CONVENTIONS", raising=False)


def _ae(serious=True, standard=False):
    rows = [("text", "AETERM", "Adverse event term", {}), ("date", "AESTDAT", "Start date", {})]
    if serious:
        rows += [("select_one NY", "AESER", "Serious?", {}), ("select_multiple CRIT", "AESCRIT", "Seriousness criteria", {}),
                 ("select_one OUT", "AEOUT", "Outcome", {})]
    f = fx._f("AE", "Adverse Event", "AE", rows, visits=("SE_COMMON",))
    if standard:
        f["customer_standard"] = {"form_oid": "AE", "domain": "AE", "source": "uploaded XLSForm"}
    return f


def _sae():
    f = fx._f("AESAE", "Serious Adverse Events", "AE", [("text", "SAETERM", "SAE term", {}), ("text", "SAEDESC", "Narrative", {})],
              visits=("SE_COMMON",))
    return f


def _spec(*forms):
    other = fx._f("VS", "Vital Signs", "VS", [("date", "VSDAT", "Date", {})])
    other["cross_form_dependencies"] = [{"source_form": "AESAE", "source_field": "SAETERM"}, {"source_form": "AE", "source_field": "AETERM"}]
    return {"study_meta": {}, "forms": list(forms) + [other],
            "timepoint_csv": {"rows": [{"event": "SE_SCREENING", "timepoint": "Screening"}, {"event": "SE_COMMON", "timepoint": "Common"}]},
            "schedule_of_events": {"form_placements": [{"target_visit_oid": "SE_COMMON", "form_id": "AESAE"},
                                                       {"target_visit_oid": "SE_COMMON", "form_id": "AE"}]}}


def test_unanswered_questions_take_their_defaults_and_answers_are_read_like_any_cq_column():
    a = fc.answers({})
    assert (a["SAE_FORM"]["value"], a["SAE_FORM"]["source"]) == ("no", "default")
    assert (a["DEATH_DETAILS_FORM"]["value"], a["DEATH_DETAILS_FORM"]["source"]) == ("protocol", "default")
    a = fc.answers({SAE_Q: "Yes - add a separate SAE form (date became serious, hospitalization dates, narrative, sponsor notification)",
                    DEATH_Q: "Never - death is captured in Disposition and Adverse Events",
                    "Do you collect Date of Visit (DOV) at scheduled visits?": "No"})
    assert (a["SAE_FORM"]["value"], a["SAE_FORM"]["source"]) == ("yes", "customer answer")
    assert a["DEATH_DETAILS_FORM"]["value"] == "never" and set(a) == {"SAE_FORM", "DEATH_DETAILS_FORM"}
    assert fc.answers({DEATH_Q: "Always"})["DEATH_DETAILS_FORM"]["value"] == "always"
    assert fc.answers({DEATH_Q: "Only when the protocol asks for them"})["DEATH_DETAILS_FORM"]["value"] == "protocol"
    assert fc.answers({"SAE Form": "No - capture seriousness, criteria and outcome on the AE form"})["SAE_FORM"]["source"] == "customer answer"
    odd = fc.answers({SAE_Q: "maybe"})["SAE_FORM"]
    assert odd["value"] == "no" and odd["source"] == "default" and "not understood" in odd["note"]


def test_sae_no_removes_the_analysis_sae_form_when_the_ae_form_carries_seriousness():
    spec = _spec(_ae(), _sae())
    lines = fc.apply_sae(spec, fc.answers({}))
    assert [f["form_id"] for f in spec["forms"]] == ["AE", "VS"]
    assert spec["schedule_of_events"]["form_placements"] == [{"target_visit_oid": "SE_COMMON", "form_id": "AE"}]
    assert spec["forms"][1]["cross_form_dependencies"] == [{"source_form": "AE", "source_field": "AETERM"}]
    assert any("AESAE (Serious Adverse Events): separate SAE form from the analysis removed (SAE_FORM = No)" in m
               and "1 cross-form reference(s)" in m for m in spec["review_flags"][fc.FLAG])
    assert lines[0] == "Customer convention SAE_FORM: No - capture seriousness, criteria and outcome on the AE form (default)"
    assert spec["study_meta"]["form_conventions_removed"]["AESAE"]["form_title"] == "Serious Adverse Events"
    # once per set of answers
    assert fc.needs_apply(spec, fc.answers({})) is False and fc.needs_apply(spec, fc.answers({SAE_Q: "Yes"})) is True


def test_sae_no_adds_seriousness_criteria_and_outcome_to_an_ae_form_that_lacks_them():
    spec = _spec(_ae(serious=False), _sae())
    fc.apply_sae(spec, fc.answers({SAE_Q: "No"}))
    ae = spec["forms"][0]
    names = [r["name"] for r in ae["survey"]]
    assert names[2:] == ["AESER", "AESDTH", "AESLIFE", "AESHOSP", "AESDISAB", "AESCONG", "AESMIE", "AEOUT"]
    assert {c["list_name"] for c in ae["choices"]} == {"NY", "AEOUT"} and ae["survey"][-1]["completion_status"] == "FLAGGED"
    assert "AESAE" not in [f["form_id"] for f in spec["forms"]]
    # an AE form that already has them is left alone
    spec = _spec(_ae())
    before = copy.deepcopy(spec["forms"][0])
    fc.apply_sae(spec, fc.answers({}))
    assert spec["forms"][0] == before and fc.FLAG not in spec.get("review_flags", {})


def test_sae_no_never_changes_a_customer_standard_ae_form_it_proposes_and_keeps_the_sae_form():
    spec = _spec(_ae(serious=False, standard=True), _sae())
    before = copy.deepcopy(spec["forms"][0])
    fc.apply_sae(spec, fc.answers({}))
    assert spec["forms"][0] == before and [f["form_id"] for f in spec["forms"]] == ["AE", "AESAE", "VS"]
    flags = spec["review_flags"][fc.FLAG]
    assert any("PROPOSAL: add seriousness, SAE criteria, outcome" in m for m in flags)
    assert any(m.startswith("AESAE (Serious Adverse Events): kept although SAE_FORM = No") for m in flags)


def test_sae_no_keeps_a_combined_ae_sae_form():
    combined = fx._f("AE", "Adverse Events / Serious Adverse Events", "AE", [("select_one NY", "AESER", "Serious?", {}),
                                                                           ("text", "AESCRIT", "Seriousness criteria", {}),
                                                                           ("text", "AEOUT", "Outcome", {})])
    spec = _spec(combined)
    fc.apply_sae(spec, fc.answers({}))
    assert [f["form_id"] for f in spec["forms"]] == ["AE", "VS"] and fc.is_sae_only("SAE Report") and not fc.is_sae_only(combined["form_title"])


def test_sae_yes_keeps_an_existing_sae_form_or_adds_one_from_cdashig():
    spec = _spec(_ae(), _sae())
    fc.apply_sae(spec, fc.answers({SAE_Q: "Yes"}))
    assert [f["form_id"] for f in spec["forms"]] == ["AE", "AESAE", "VS"]
    spec = _spec(_ae())
    fc.apply_sae(spec, fc.answers({SAE_Q: "Yes"}))
    sae = spec["forms"][-1]
    assert sae["form_id"] == "SAE" and sae["visits_assigned"] == ["SE_COMMON"] and sae["has_repeating_group"] is True
    assert [r["name"] for r in sae["survey"]] == ["AESPID", "AETERM", "AESERDAT", "AESDTH", "AESLIFE", "AESHOSP", "AESDISAB",
                                                  "AESCONG", "AESMIE", "HOSTDAT", "HOENDAT", "AEOUT", "AENARR", "AESPNDAT"]
    assert sae["convention_required"]["content_source"] == "CDASHIG"
    assert {"target_visit_oid": "SE_COMMON", "form_id": "SAE", "required": False, "repeating": True,
            "notes": "Customer convention SAE_FORM = Yes"} in spec["schedule_of_events"]["form_placements"]


def test_sae_yes_uses_the_customers_standard_sae_form_when_there_is_one():
    settings = {"form_title": "Serious Adverse Event Report", "form_id": "SAESTD", "version": "2", "style": "theme-grid"}
    survey = [["date", "SAEAWDAT", "Date site became aware", "SAE", "", "", "", "yes", "", "", "", "", "", ""],
              ["text", "SAENARR", "Narrative", "SAE", "", "", "", "", "", "", "", "", "", ""]]
    src = sm.load_sources([("SAESTD.xlsx", fx.xlsform_bytes(survey, [], settings))])
    spec = sm.apply(_spec(_ae()), src)
    fc.apply_sae(spec, fc.answers({SAE_Q: "Yes"}), src)
    assert sm.splice_added(spec, src)[0]["basis"] == "created for this standard form (customer convention)"
    sae = spec["forms"][-1]
    assert sae["form_id"] == "SAESTD" and [r["name"] for r in sae["survey"]] == ["SAEAWDAT", "SAENARR"]
    assert sae["customer_standard"]["form_oid"] == "SAESTD" and sae["visits_assigned"] == ["SE_COMMON"]


def test_kill_switch_leaves_the_forms_alone(monkeypatch):
    monkeypatch.setenv("FORM_CONVENTIONS", "0")
    assert fc.answers({SAE_Q: "Yes", DEATH_Q: "Always"})["SAE_FORM"]["value"] == "no"
    assert fc.needs_apply(_spec(_ae(), _sae()), None) is False


def test_both_answers_are_recorded_with_their_source_and_shown_in_the_study_specification():
    spec = _spec(_ae(), _sae(), fx._f("DS", "Disposition", "DS", [("date", "DSSTDAT", "Date", {})]))
    ans = fc.answers({DEATH_Q: "Never"})
    fc.apply_sae(spec, ans)
    death = {"name": "Death", "domain": "DD", "quote": "The cause of death will be recorded.", "log": True, "section": "8.2", "events": []}
    pf.apply(spec, [death], None, None, "The cause of death will be recorded.", death=ans["DEATH_DETAILS_FORM"]["value"])
    fc.record_death(spec, ans)
    title, _note, headers, rows, _w = fc.section(spec)
    assert title == "CUSTOMER FORM CONVENTIONS" and headers == ["Convention", "Question", "Answer", "Source", "Applied"]
    assert rows[0][0] == "SAE_FORM" and rows[0][3] == "default" and "removed (SAE_FORM = No)" in rows[0][4]
    assert rows[1][:4] == ["DEATH_DETAILS_FORM", DEATH_Q, "Never - death is captured in Disposition and Adverse Events", "customer answer"]
    assert "DEATH_DETAILS_FORM = Never" in rows[1][4] and "(recorded on DS)" in rows[1][4]
    import openpyxl
    from generate_study_spec_xlsx import build_customer_conventions_sheet
    from generate_study_spec_pdf import build_form_conventions_block, make_styles
    wb = openpyxl.Workbook()
    build_customer_conventions_sheet(wb, spec)
    ws = wb["CUSTOMER_CONVENTIONS"]
    assert ws.cell(row=3, column=1).value == "SAE_FORM" and ws.cell(row=4, column=4).value == "customer answer"
    assert len(build_form_conventions_block(spec, make_styles())) == 4
    wb2 = openpyxl.Workbook()
    build_customer_conventions_sheet(wb2, {"study_meta": {}})
    assert "CUSTOMER_CONVENTIONS" not in wb2.sheetnames and build_form_conventions_block({"study_meta": {}}, make_styles()) == []


def test_a_reused_specification_keeps_its_forms_the_answer_is_recorded_and_flagged():
    spec = _spec(_ae(serious=False), _sae())
    before = copy.deepcopy(spec["forms"])
    fc.apply_sae(spec, fc.answers({}), fresh=False)
    assert spec["forms"] == before
    rec = fc.state(spec)["answers"][0]
    assert rec["applied"] and "SAE_FORM = No, but there is a separate SAE form (AESAE); not applied" in rec["applied"][0]
    assert spec["review_flags"][fc.FLAG] == rec["applied"]
    spec = _spec(_ae())
    fc.apply_sae(spec, fc.answers({SAE_Q: "Yes"}), fresh=False)
    assert [f["form_id"] for f in spec["forms"]] == ["AE", "VS"] and "there is no separate SAE form" in fc.state(spec)["answers"][0]["applied"][0]


def test_only_the_fresh_analysis_path_lets_the_convention_change_forms():
    src = open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pipeline.py")).read()
    assert src.count("await _post_match_forms_step(") == 4 and src.count("_oc_files, cols, fresh=True)") == 1

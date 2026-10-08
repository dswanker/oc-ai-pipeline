"""Edited DVS round trip: Approve / Reject / Delete / Change / Add read back from DVS_OC4 and applied
deterministically; deletions survive the rules engine re-running; builder keeps combined and DM messages."""
import copy, io, json, os, sys, tempfile, contextlib
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "conventions"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "skills", "dvs-specification", "scripts"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "skills", "edc-builder", "scripts"))
import openpyxl
from conventions_engine import apply_conventions
import ai_edit_checks as ai
import dvs_edits
from extract_dvs_from_forms import extract_dvs_data
from generate_dvs import build_dvs
from build_xlsforms import _normalize_constraint_messages
from test_consent_floor import _spec as consent_spec

PROPOSALS = json.dumps({"checks": [
    {"target_form": "F_AE", "target_field": "ONSET", "operator": ">=", "source_form": "DM", "source_field": "VISDAT",
     "message": "AE start must be on or after the visit date.", "category": "cross_form"},
    {"target_form": "DM", "target_field": "VISDAT", "operator": ">", "source_form": "DM", "source_field": "BRTHDAT",
     "message": "Visit after birth.", "category": "sequential_dates"}]})


def _engine(spec):
    with contextlib.redirect_stdout(io.StringIO()):
        apply_conventions(spec, study_id="T", customer_subdomain="", client_name="")
    return spec


def _dvs_bytes(spec):
    fj = {"forms": {f"{f['form_id']}.xlsx": {"survey": [{("bind::oc:itemgroup" if k == "bind__oc_itemgroup" else k): v
                                                          for k, v in r.items()} for r in f["survey"]],
                                               "choices": f.get("choices") or []} for f in spec["forms"]}}
    path = os.path.join(tempfile.mkdtemp(), "dvs.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        build_dvs(extract_dvs_data(spec, fj), path)
    return open(path, "rb").read()


def _edit(xbytes, edits, add=None):
    wb = openpyxl.load_workbook(io.BytesIO(xbytes)); ws = wb["DVS_OC4"]
    col = {c.value: c.column for c in ws[3]}
    for r in range(4, ws.max_row + 1):
        key = (ws.cell(r, col["Target Item Name"]).value, ws.cell(r, col["Rule / Proposal ID"]).value)
        if key in edits:
            action, msg = edits[key]
            ws.cell(r, col["Action"]).value = action
            if msg:
                ws.cell(r, col["Constraint / Required / Relevant Message"]).value = msg
    if add:
        nr = ws.max_row + 1
        for k, v in add.items():
            ws.cell(nr, col[k]).value = v
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _row(spec, form, item):
    return next(r for f in spec["forms"] if f["form_id"] == form for r in f["survey"] if r.get("name") == item)


def test_round_trip():
    spec = _engine(consent_spec())
    spec["study_meta"]["ai_edit_checks"] = {"proposals": ai.validate_response(spec, PROPOSALS)["proposals"]}
    edited = _edit(_dvs_bytes(spec), {
        ("ONSET", "AI.001"): ("Approve", None),
        ("VISDAT", "AI.002"): ("Reject", None),
        ("ONSET", "CONSENT_FLOOR"): ("Change", "AE start is before consent: please check."),
        ("VISDAT", "CONSENT_FLOOR"): ("Delete", None)},
        add={"Action": "Add", "Target Form OID": "F_AE", "Plain-English Description": "AE start on or after enrollment."})
    acts = dvs_edits.parse_actions(edited)
    res = dvs_edits.apply_actions(spec, acts)
    assert (res["applied"], res["pending"], res["failed"]) == (4, 1, 0)
    _engine(spec)                                                    # the engine re-runs on every build
    onset = _row(spec, "F_AE", "ONSET")
    assert "AI.001" in onset["edit_checks"] and "${VISDAT_CF}" in onset["constraint"]
    assert "AE start is before consent: please check." in onset["constraint_message"]
    assert onset["constraint_message_locked"] is True
    vis = _row(spec, "DM", "VISDAT")
    assert "RFICDAT" not in str(vis.get("constraint")) and "CONSENT_FLOOR" in vis["edit_checks_suppressed"]
    assert [p["id"] for p in spec["study_meta"]["ai_edit_checks"]["proposals"]] == []      # approved + rejected
    assert "AI.002" in spec["study_meta"]["ai_edit_checks"]["rejected_ids"]
    again = dvs_edits.apply_actions(copy.deepcopy(spec), acts)                             # idempotent
    assert again["applied"] == 0 and again["failed"] == 0


def test_old_format_dvs_is_not_read_by_new_path():
    wb = openpyxl.Workbook(); wb.active.title = "DVS_OC4"; wb.active.append(["Check ID", "Status"])
    buf = io.BytesIO(); wb.save(buf)
    assert dvs_edits.parse_actions(buf.getvalue()) is None


def test_builder_keeps_combined_and_dm_messages():
    rows = [{"name": "A", "constraint": "(. <= today()) and (. >= ${X})", "constraint_message": "Not future. After X."},
            {"name": "B", "constraint": ". <= today()", "constraint_message": "Future dates are not allowed."},
            {"name": "C", "constraint": ". <= today()", "constraint_message": "DM wording.", "constraint_message_locked": True}]
    with contextlib.redirect_stdout(io.StringIO()):
        out = {r["name"]: r["constraint_message"] for r in _normalize_constraint_messages(rows)}
    assert out == {"A": "Not future. After X.", "B": "Date cannot be in the future.", "C": "DM wording."}


def test_dm_plain_english_adds():
    spec = _engine(consent_spec())
    rows = [{"Action": "add", "Target Form OID": "F_AE", "Target Item Name": "ONSET",
             "Plain-English Description": "AE start must be on or after the visit date."},
            {"Action": "add", "Target Form OID": "DM", "Target Item Name": "",
             "Plain-English Description": "Birth year must be after 1900."},
            {"Action": "add", "Target Form OID": "F_AE", "Target Item Name": "",
             "Plain-English Description": "AE start must equal the moon landing date."}]
    adds = dvs_edits.pending_adds(spec, rows)
    assert len(adds) == 3
    prompt, extra = dvs_edits.build_add_request(spec, adds)
    assert "REQUESTS:" in extra and "1. form F_AE | item ONSET" in extra
    scripted = json.dumps({"results": [
        {"index": 1, "target_form": "F_AE", "target_field": "ONSET", "operator": ">=", "source_form": "DM",
         "source_field": "VISDAT", "message": "AE start must be on or after the visit date."},
        {"index": 2, "cannot": "needs a constant comparison"},
        {"index": 3, "target_form": "F_AE", "target_field": "ONSET", "operator": "=", "source_form": "DM",
         "source_field": "MOONDAT", "message": "x"}]})
    by_idx = dvs_edits.parse_add_response("```json\n" + scripted + "\n```", adds)
    trans = {dvs_edits._add_key(r): by_idx.get(i) for i, r in enumerate(adds, 1)}
    res = dvs_edits.apply_actions(spec, rows, trans)
    assert [r["status"] for r in res["results"]] == ["applied", "needs_build_team", "needs_build_team"]
    onset = _row(spec, "F_AE", "ONSET")
    assert "DM.001" in onset["edit_checks"] and "${VISDAT_CF}" in onset["constraint"]
    assert any(d.get("source") == "DM-Added" for d in onset["edit_check_details"])
    assert dvs_edits.pending_adds(spec, rows) == []                       # nothing re-sent to the AI
    again = dvs_edits.apply_actions(spec, rows, {})
    assert [r["status"] for r in again["results"]] == ["already", "needs_build_team", "needs_build_team"]
    wb = openpyxl.load_workbook(io.BytesIO(_dvs_bytes(spec)))
    ws = wb["DVS_OC4"]; col = {c.value: c.column for c in ws[3]}
    dm = [(ws.cell(r, col["Check Source"]).value, ws.cell(r, col["Status"]).value,
           ws.cell(r, col["Plain-English Description"]).value) for r in range(4, ws.max_row + 1)
          if ws.cell(r, col["Check Source"]).value == "DM-Added"]
    assert ("DM-Added", "Needs Build Team", "Birth year must be after 1900.") in dm
    assert any(s == "Draft" for _, s, _ in dm)                               # DM.001 is in the build

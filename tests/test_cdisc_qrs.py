"""cdisc_qrs.py: QRS instruments from CDISC CT. Synthetic CT (PHQ-9-style instrument, customer field names):
no network, no member files. Includes an end-to-end pass through the real conventions engine, EDC builder and
DVS extractor."""
import ast, contextlib, copy, io, json, os, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "skills", "dvs-specification", "scripts"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "skills", "edc-builder", "scripts"))
import cdisc_cdash as cd
import cdisc_concepts as cc
import cdisc_ct as c
import cdisc_qrs as q
import sdtm_mapping as sm

HDR = ("Code\tCodelist Code\tCodelist Extensible (Yes/No)\tCodelist Name\tCDISC Submission Value\t"
       "CDISC Synonym(s)\tCDISC Definition\tNCI Preferred Term\n")
_P = "Patient Health Questionnaire - 9"
ROWS = [
    ("C66742", "", "No", "No Yes Response", "NY", "", "", ""),
    ("C49487", "C66742", "", "", "N", "No", "", "No"),
    ("C49488", "C66742", "", "", "Y", "Yes", "", "Yes"),
    ("C100129", "", "Yes", "Category of Questionnaire", "QSCAT", "", "", ""),
    ("C103519", "C100129", "", "", "PHQ-9", "PHQ01", "", f"{_P} Item Questionnaire"),
    ("C118971", "", "Yes", "Category of Clinical Classification", "CCCAT", "", "", ""),
    ("C102408", "C118971", "", "", "ECOG", "ECOG1", "", "Eastern Cooperative Oncology Group Performance Status"),
    ("C103481", "", "No", f"{_P} Item Questionnaire Test Code", "PHQ01TC", "", "", ""),
    ("C103705", "C103481", "", "", "PHQ0101", "PHQ01-Little Interest", "", "PHQ-9 - Little Interest"),
    ("C103706", "C103481", "", "", "PHQ0102", "PHQ01-Feeling Down", "", "PHQ-9 - Feeling Down"),
    ("C103707", "C103481", "", "", "PHQ0103", "PHQ01-Trouble Sleeping", "", "PHQ-9 - Trouble Sleeping"),
    ("C103714", "C103481", "", "", "PHQ0110", "PHQ01-Difficult to Work", "", "PHQ-9 - How Difficult"),
    ("C113887", "C103481", "", "", "PHQ0111", "PHQ01-Total Score", "", "PHQ-9 - Total Score"),
    ("C103480", "", "No", f"{_P} Item Questionnaire Test Name", "PHQ01TN", "", "", ""),
    ("C103705", "C103480", "", "", "PHQ01-Little Interest", "", "", "PHQ-9 - Little Interest"),
    ("C103706", "C103480", "", "", "PHQ01-Feeling Down", "", "", "PHQ-9 - Feeling Down"),
    ("C103707", "C103480", "", "", "PHQ01-Trouble Sleeping", "", "", "PHQ-9 - Trouble Sleeping"),
    ("C103714", "C103480", "", "", "PHQ01-Difficult to Work", "", "", "PHQ-9 - How Difficult"),
    ("C113887", "C103480", "", "", "PHQ01-Total Score", "", "", "PHQ-9 - Total Score"),
    ("C213939", "", "No", f"{_P} Questionnaire ORRES for PHQ0101 Through PHQ0103 TN/TC", "PHQ0101T03OR", "", "", ""),
    ("C214825", "C213939", "", "", "More than half the days", "", f"{_P} original result for PHQ0101 through PHQ0103-More than half the days.", ""),
    ("C214826", "C213939", "", "", "Nearly every day", "", f"{_P} original result for PHQ0101 through PHQ0103-Nearly every day.", ""),
    ("C214823", "C213939", "", "", "Not at all", "", f"{_P} original result for PHQ0101 through PHQ0103-Not at all.", ""),
    ("C214824", "C213939", "", "", "Several days", "", f"{_P} original result for PHQ0101 through PHQ0103-Several days.", ""),
    ("C213941", "", "No", f"{_P} Questionnaire STRESC for PHQ0101 Through PHQ0103 TN/TC", "PHQ0101T03STR", "", "", ""),
    ("C214831", "C213941", "", "", "0", "", f"{_P} standardized character result for PHQ0101 through PHQ0103-Not at all.", ""),
    ("C214832", "C213941", "", "", "1", "", f"{_P} standardized character result for PHQ0101 through PHQ0103-Several days.", ""),
    ("C214833", "C213941", "", "", "2", "", f"{_P} standardized character result for PHQ0101 through PHQ0103-More than half the days.", ""),
    ("C214834", "C213941", "", "", "3", "", f"{_P} standardized character result for PHQ0101 through PHQ0103-Nearly every day.", ""),
    ("C213940", "", "No", f"{_P} Questionnaire ORRES for PHQ0110 TN/TC", "PHQ0110OR", "", "", ""),
    ("C214827", "C213940", "", "", "Not difficult at all", "", f"{_P} original result for PHQ0110-Not difficult at all.", ""),
    ("C214828", "C213940", "", "", "Somewhat difficult", "", f"{_P} original result for PHQ0110-Somewhat difficult.", ""),
    ("C213942", "", "No", f"{_P} Questionnaire STRESC for PHQ0110 TN/TC", "PHQ0110STR", "", "", ""),
    ("C214835", "C213942", "", "", "Not difficult at all", "", f"{_P} standardized character result for PHQ0110-Not difficult at all.", ""),
    ("C214836", "C213942", "", "", "Somewhat difficult", "", f"{_P} standardized character result for PHQ0110-Somewhat difficult.", ""),
    ("C101815", "", "No", "Eastern Cooperative Oncology Group Performance Status Clinical Classification Test Code", "ECOG1TC", "", "", ""),
    ("C102408", "C101815", "", "", "ECOG101", "ECOG1-Performance Status", "", "ECOG - Performance Status"),
    ("C101816", "", "No", "Eastern Cooperative Oncology Group Performance Status Clinical Classification Test Name", "ECOG1TN", "", "", ""),
    ("C102408", "C101816", "", "", "ECOG1-Performance Status", "", "", "ECOG - Performance Status"),
    ("C179943", "", "No", "ECOG Clinical Classification ORRES for ECOG101 TN/TC", "ECOG101OR", "", "", ""),
    ("C180249", "C179943", "", "", "Fully active, no restriction", "", "ECOG original result for ECOG101-Fully active, no restriction.", ""),
    ("C180250", "C179943", "", "", "Restricted in strenuous activity", "", "ECOG original result for ECOG101-Restricted in strenuous activity.", ""),
    ("C180254", "C179943", "", "", "Dead", "", "ECOG original result for ECOG101-Dead.", ""),
    ("C179944", "", "No", "ECOG Clinical Classification STRESC for ECOG101 TN/TC", "ECOG101STR", "", "", ""),
    ("C180255", "C179944", "", "", "0", "", "ECOG standardized character result for ECOG101-Fully active, no restriction.", ""),
    ("C180256", "C179944", "", "", "1", "", "ECOG standardized character result for ECOG101-Restricted in strenuous activity.", ""),
    ("C180260", "C179944", "", "", "5", "", "ECOG standardized character result for ECOG101-Dead.", ""),
]
CT = c.CTPackage(c._parse_evs(HDR + "".join("\t".join(r) + "\n" for r in ROWS)), "test", "2026-01-01")
STD = c.Standards(CT, {}, {})
IX = q.build_index(CT)
FREQ = [("a", "Not at all"), ("b", "Several days"), ("c", "More than half the days"), ("d", "Nearly every day")]


def setup_module():
    cd._MEMO["fields"] = {("RS", "RSORRES"): {"domain": "RS", "variable": "RSORRES", "label": "", "question": "",
                                              "prompt": "", "instruction": "", "core": "HR", "type": "Char",
                                              "sdtm_target": "RSORRES", "mapping": ""}}


def teardown_module():
    cd._MEMO.clear()


def _spec():
    """A PHQ-9-style form with customer field names and codes, and an ECOG form coded by grade."""
    ch = [{"list_name": "freq", "name": n, "label": l} for n, l in FREQ]
    ch += [{"list_name": "diff", "name": "1", "label": "Not difficult at all"},
           {"list_name": "diff", "name": "2", "label": "Somewhat difficult"},
           {"list_name": "yn", "name": "Y", "label": "Yes"}, {"list_name": "yn", "name": "N", "label": "No"}]
    return {"study_meta": {"protocol_number": "QRS-1"}, "forms": [
        {"form_id": "MOOD", "form_title": "PHQ-9 Depression Questionnaire", "visits_assigned": ["SE_V1"],
         "choices": ch, "survey": [
             {"type": "date", "name": "MOODDT", "label": "Date completed", "required": "yes"},
             {"type": "select_one freq", "name": "MOOD_Q1", "label": "Question 1", "required": "yes"},
             {"type": "select_one freq", "name": "MOOD_Q2", "label": "Question 2", "required": "yes"},
             {"type": "select_one freq", "name": "MOOD_Q3", "label": "Question 3", "required": "yes"},
             {"type": "select_one diff", "name": "MOOD_Q10", "label": "Question 10"},
             {"type": "text", "name": "MOODWHY", "label": "Follow-up", "relevant": "${MOOD_Q3} = 'd'"}]},
        {"form_id": "PERF", "form_title": "ECOG Performance Status", "visits_assigned": ["SE_V1"],
         "choices": [{"list_name": "ecog", "name": "0", "label": "0 - Fully active"},
                     {"list_name": "ecog", "name": "1", "label": "1 - Restricted"}],
         "survey": [{"type": "date", "name": "PERFDT", "label": "Date"},
                    {"type": "select_one ecog", "name": "PERFGR", "label": "Performance status", "required": "yes"}]},
        {"form_id": "VS", "form_title": "Vital Signs", "visits_assigned": ["SE_V1"],
         "choices": [{"list_name": "yn", "name": "Y", "label": "Yes"}, {"list_name": "yn", "name": "N", "label": "No"}],
         "survey": [{"type": "select_one yn", "name": "VSDONE", "label": "Vitals taken?"}]},
    ]}


def _rows(spec):
    return {r["name"]: r for f in spec["forms"] for r in f["survey"]}


def _lists(spec, fid):
    out = {}
    for ch in next(f for f in spec["forms"] if f["form_id"] == fid)["choices"]:
        out.setdefault(ch["list_name"], []).append(ch)
    return out


def _run(spec, protected=frozenset()):
    n = q.tag_by_order(spec, IX)
    return n, q.apply_to_spec(spec, STD, protected)


def test_index_items_responses_and_scores_come_from_ct():
    phq = IX.instruments["PHQ01"]
    assert (phq["domain"], phq["category"]) == ("QS", "PHQ-9") and list(phq["items"])[0] == "PHQ0101"
    assert phq["items"]["PHQ0101"]["test"] == "PHQ01-Little Interest"
    assert IX.instruments["ECOG1"]["domain"] == "RS"
    r = IX.responses["PHQ0102"]  # inside "PHQ0101 Through PHQ0103"
    assert r["orres"]["short_name"] == "PHQ0101T03OR" and r["score_basis"] == "ct_stresc_definition"
    assert r["scores"] == {"NOT AT ALL": 0, "SEVERAL DAYS": 1, "MORE THAN HALF THE DAYS": 2, "NEARLY EVERY DAY": 3}
    assert IX.responses["PHQ0110"]["scores"] == {}        # STRESC is text: no numeric score is invented
    assert "PHQ0111" not in IX.responses                   # total score: no response list


def test_form_names_its_instrument():
    assert q.candidates({"form_title": "PHQ-9 Depression Questionnaire"}, IX) == {"PHQ01": True}
    assert q.candidates({"form_title": "ECOG Performance Status"}, IX) == {"ECOG1": False}   # short key: weak
    assert q.candidates({"form_title": "Vital Signs", "form_id": "VS"}, IX) == {}
    assert q.candidates({"form_title": "Mood", "survey": [{"type": "note", "label": "Patient Health Questionnaire - 9"}]},
                        IX) == {"PHQ01": True}


def test_customer_named_phq_form_gets_codes_lists_and_scores():
    spec = _spec()
    n, dec = _run(spec)
    rows = _rows(spec)
    assert n["tagged"] == 5
    assert [(rows[k]["concept"], rows[k]["concept_qualifier"], rows[k]["concept_source"])
            for k in ("MOOD_Q1", "MOOD_Q3", "MOOD_Q10")] == [
        ("QSORRES", "PHQ0101", "qrs_instrument"), ("QSORRES", "PHQ0103", "qrs_instrument"),
        ("QSORRES", "PHQ0110", "qrs_instrument")]
    assert "concept" not in rows["MOODDT"] and "concept" not in rows["VSDONE"]
    # response codelist applied: instrument order, score as the code, CT text as the label, score recorded
    assert rows["MOOD_Q1"]["type"] == "select_one qrs_phq0101t03or"
    lst = _lists(spec, "MOOD")["qrs_phq0101t03or"]
    assert [(x["name"], x["label"], x["cdisc_code"]) for x in lst] == [
        ("0", "Not at all", "C214823"), ("1", "Several days", "C214824"),
        ("2", "More than half the days", "C214825"), ("3", "Nearly every day", "C214826")]
    assert set(lst[0]) == {"list_name", "name", "label", "source", "cdisc_submission_value", "cdisc_code",
                           "cdisc_codelist"}                      # same choice columns as the CT layer, no new one
    assert all(x["source"] == "CDISC_QRS" and x["cdisc_codelist"] == "C213939" for x in lst)
    assert rows["MOOD_Q1"]["qrs"]["scores"] == {"0": 0, "1": 1, "2": 2, "3": 3}
    assert (rows["MOOD_Q1"]["qrs"]["category"], rows["MOOD_Q1"]["qrs"]["testcd"]) == ("PHQ-9", "PHQ0101")
    # text responses (no numeric standard result): XLSForm-safe names, no score
    assert [x["name"] for x in _lists(spec, "MOOD")["qrs_phq0110or"]] == ["Not_difficult_at_all", "Somewhat_difficult"]
    assert "scores" not in rows["MOOD_Q10"]["qrs"]
    # skip logic follows the renamed code; the question label is never touched; replaced lists are removed
    assert rows["MOODWHY"]["relevant"] == "${MOOD_Q3} = '3'"
    assert rows["MOOD_Q1"]["label"] == "Question 1"
    assert "freq" not in _lists(spec, "MOOD") and "diff" not in _lists(spec, "MOOD")
    assert next(f for f in spec["forms"] if f["form_id"] == "MOOD")["qrs_instruments"][0]["category"] == "PHQ-9"
    assert q.summarize(dec)["actions"]["response_codelist_applied"] == 5


def test_ecog_matched_by_grade_keeps_study_subset():
    spec = _spec()
    _run(spec)
    r = _rows(spec)["PERFGR"]
    assert (r["concept"], r["concept_qualifier"]) == ("RSORRES", "ECOG101")
    lst = _lists(spec, "PERF")[r["type"].split(" ")[1]]
    # the study shows the grade in front of the response: kept; "Dead" (not collected by the study) is not added
    assert [(x["name"], x["label"]) for x in lst] == [("0", "0 - Fully active, no restriction"),
                                                      ("1", "1 - Restricted in strenuous activity")]
    assert [x["cdisc_submission_value"] for x in lst] == ["Fully active, no restriction", "Restricted in strenuous activity"]


def test_protected_fields_untouched_and_unknown_tier_is_metadata_only():
    for protected in ({"MOOD_Q1", "MOOD_Q2", "MOOD_Q3", "MOOD_Q10", "PERFGR"}, None):
        spec = _spec()
        before = copy.deepcopy(spec)
        _run(spec, protected)
        for f, b in zip(spec["forms"], before["forms"]):
            assert f["choices"] == b["choices"]
            assert [(r["type"], r.get("relevant")) for r in f["survey"]] == [(r["type"], r.get("relevant")) for r in b["survey"]]
        r = _rows(spec)["MOOD_Q2"]
        assert r["concept_qualifier"] == "PHQ0102" and r["qrs"]["orres_codelist"] == "PHQ0101T03OR"
        assert r["qrs"]["scores"] == {"a": 0, "b": 1, "c": 2, "d": 3}    # scores for the customer's own codes
    spec = _spec()
    spec["forms"][0]["choices"][0]["source"] = "crf_standards_injection"   # a customer-standards list
    _run(spec)
    assert _rows(spec)["MOOD_Q1"]["type"] == "select_one freq"


def test_not_guessed_when_ambiguous():
    spec = _spec()   # one extra answer field: counts no longer line up, nothing is tagged
    spec["forms"][0]["survey"].insert(2, {"type": "select_one freq", "name": "EXTRA", "label": "Extra"})
    _run(spec)
    assert not any(r.get("concept") for r in spec["forms"][0]["survey"])
    spec = _spec()   # responses that are not the instrument's: no order match
    spec["forms"][0]["choices"][0]["label"] = "Never ever"
    _run(spec)
    assert not any(r.get("concept") for r in spec["forms"][0]["survey"])
    spec = _spec()   # weakly named instrument, no response evidence: not matched on the name alone
    spec["forms"][1]["choices"] = []
    spec["forms"][1]["survey"][1]["type"] = "integer"
    _run(spec)
    assert "concept" not in _rows(spec)["PERFGR"]


def test_idempotent():
    spec = _spec()
    _run(spec)
    once = copy.deepcopy(spec)
    n, dec = _run(spec)
    assert spec == once and n["tagged"] == 0
    c.apply_to_spec(spec, STD)             # the CT layer leaves questionnaire response lists alone
    for f, o in zip(spec["forms"][:2], once["forms"][:2]):
        assert f["choices"] == o["choices"] and [r["type"] for r in f["survey"]] == [r["type"] for r in o["survey"]]


def test_ai_tags_validated_against_instrument_codelist_then_order_fills_gaps():
    spec = _spec()
    spec["forms"][0]["form_title"] = "Mood (PHQ-9)"
    spec["forms"][0]["survey"].append({"type": "integer", "name": "MOODTOT", "label": "Total"})
    req = cc.build_request(spec, STD)
    assert "QRS INSTRUMENT ITEMS" in req[1] and "PHQ0101 | PHQ-9 - Little Interest" in req[1]
    assert "Questionnaires, ratings and scales" in req[0]
    assert cc.build_request(spec, STD, qrs=False)[0] == cc.PROMPT
    resp = json.dumps({"tags": [
        {"form_id": "MOOD", "field": "MOOD_Q1", "concept": "QSORRES", "qualifier": "PHQ0101", "confidence": "high"},
        {"form_id": "MOOD", "field": "MOOD_Q3", "concept": "RSORRES", "qualifier": "PHQ0103", "confidence": "high"},
        {"form_id": "MOOD", "field": "MOODTOT", "concept": "QSORRES", "qualifier": "PHQ0111", "confidence": "high"},
        {"form_id": "MOOD", "field": "MOOD_Q10", "concept": "QSORRES", "qualifier": "PHQ0199", "confidence": "high"},
        {"form_id": "MOOD", "field": "MOODDT", "concept": "QSORRES", "qualifier": "PHQ0102", "confidence": "high"},
        {"form_id": "VS", "field": "VSDONE", "concept": "QSORRES", "qualifier": "PHQ0102", "confidence": "high"},
        {"form_id": "PERF", "field": "PERFGR", "concept": "RSORRES", "qualifier": "ECOG101", "confidence": "medium"}]})
    out = cc.apply_ai_response(spec, STD, resp)
    rows = _rows(spec)
    assert out["accepted"] == 3
    assert out["rejected"] == {"qrs_test_code_not_in_ct": 1, "type_mismatch": 1,
                               "qrs_instrument_not_named_on_form": 1, "not_high_confidence": 1}
    assert (rows["MOOD_Q3"]["concept"], rows["MOOD_Q3"]["concept_source"]) == ("QSORRES", "claude")  # domain corrected
    assert rows["MOODTOT"]["concept_qualifier"] == "PHQ0111"
    n, _ = _run(spec)
    assert rows["MOOD_Q2"]["concept_qualifier"] == "PHQ0102" and rows["MOOD_Q2"]["concept_source"] == "qrs_instrument"
    assert rows["MOOD_Q1"]["concept_source"] == "claude"       # an existing tag is never overridden
    assert rows["MOODTOT"]["qrs"]["testcd"] == "PHQ0111" and "orres_codelist" not in rows["MOODTOT"]["qrs"]
    # with the layer off the questionnaire tags are not accepted
    off_spec = _spec()
    cc.apply_ai_response(off_spec, STD, resp, qrs=False)
    assert not any(r.get("concept") in ("QSORRES", "FTORRES") for r in _rows(off_spec).values())


def test_sdtm_mapping_rows():
    spec = _spec()
    _run(spec)
    rows = {r["Field"]: r for r in sm.build_rows(spec)}
    r = rows["MOOD_Q1"]
    assert (r["SDTM Domain"], r["SDTM Variable"]) == ("QS", "QSORRES")
    assert r["Value-Level Detail"] == "QSTESTCD = PHQ0101; QSCAT = PHQ-9"
    assert r["Controlled Terminology"] == "PHQ0101T03OR (C213939)"
    assert (r["Confidence"], r["Basis"]) == ("Medium", "QRS instrument (item order)")
    assert "QSSTRESN = response score (0=0, 1=1, 2=2, 3=3)" in r["Notes"] and "Review" in r["Notes"]
    assert rows["PERFGR"]["Value-Level Detail"] == "RSTESTCD = ECOG101; RSCAT = ECOG"
    assert rows["MOODDT"]["Confidence"] == "None"
    files = sm.build_files(spec)
    assert files["pdf"][:4] == b"%PDF" and files["summary"]["by_confidence_basis"]["Medium / QRS instrument (item order)"] == 5


def _build(spec):
    from build_xlsforms import build_all_xlsforms
    import openpyxl
    log = {k: [] for k in ("forms_built", "forms_skipped", "build_errors", "build_warnings",
                           "placeholder_applied", "oid_placeholders")}
    forms = {}
    with tempfile.TemporaryDirectory() as t, contextlib.redirect_stdout(io.StringIO()):
        build_all_xlsforms(copy.deepcopy(spec), t, log)
        for fn in sorted(os.listdir(t)):
            if fn.endswith(".xlsx"):
                wb = openpyxl.load_workbook(os.path.join(t, fn), data_only=True)
                sheets = {}
                for sh in ("survey", "choices"):
                    rows = list(wb[sh].iter_rows(values_only=True))
                    hdr = [str(h or "") for h in rows[0]]
                    sheets[sh] = [{h: v for h, v in zip(hdr, r) if v is not None} for r in rows[1:] if any(r)]
                forms[fn] = sheets
    return log, {"forms": forms}


def test_end_to_end_engine_builder_dvs():
    from conventions_engine import apply_conventions
    from extract_dvs_from_forms import extract_dvs_data
    spec = _spec()
    _run(spec)
    with contextlib.redirect_stdout(io.StringIO()):
        apply_conventions(spec, study_id="T", customer_subdomain="", client_name="")
    log, forms_json = _build(spec)
    assert not log["build_errors"] and len(log["forms_built"]) == 3
    mood = next(v for k, v in forms_json["forms"].items() if "MOOD" in k)
    assert [r["type"] for r in mood["survey"] if r.get("name") == "MOOD_Q1"] == ["select_one qrs_phq0101t03or"]
    built = [(str(x["name"]), x["label"]) for x in mood["choices"] if x["list_name"] == "qrs_phq0101t03or"]
    assert built == [("0", "Not at all"), ("1", "Several days"), ("2", "More than half the days"), ("3", "Nearly every day")]
    assert not any(k.startswith("qrs") for x in mood["choices"] for k in x)   # scores stay in spec metadata
    assert not any(k in ("qrs", "concept", "concept_qualifier") for r in mood["survey"] for k in r)
    with contextlib.redirect_stdout(io.StringIO()):
        dvs = extract_dvs_data(spec, forms_json)
    row = next(r for r in dvs["dvs_oc4"] if r["Target Item Name"] == "MOOD_Q1" and r["Check Type"] == "Required")
    assert row["Notes"] == ("QRS instrument: PHQ-9 (PHQ01), item PHQ0101; response codelist PHQ0101T03OR (C213939), "
                            "scored; CDISC CT 2026-01-01")
    other = next(r for r in dvs["dvs_oc4"] if r["Target Item Name"] == "MOODDT")
    assert other["Notes"] == ""


def _pipeline_fn(name):
    path = os.path.join(os.path.dirname(__file__), "..", "pipeline.py")
    src = open(path).read()
    node = next(n for n in ast.walk(ast.parse(src)) if isinstance(n, ast.FunctionDef) and n.name == name)
    ns = {"os": os}
    exec(compile(ast.get_source_segment(src, node), "<pipeline>", "exec"), ns)
    return ns[name]


def test_pipeline_layer_kill_switch_and_failure_leave_spec_unchanged(monkeypatch):
    apply_qrs = _pipeline_fn("_apply_qrs")
    spec = _spec()
    before = copy.deepcopy(spec)
    monkeypatch.setenv("CDISC_QRS", "0")
    assert apply_qrs(spec, STD, frozenset()) is None and spec == before
    monkeypatch.delenv("CDISC_QRS")
    monkeypatch.setattr(q, "apply_to_spec", lambda *a, **k: 1 / 0)
    with contextlib.redirect_stdout(io.StringIO()):
        assert apply_qrs(spec, STD, frozenset()) is None
    assert spec == before
    monkeypatch.undo()
    with contextlib.redirect_stdout(io.StringIO()):
        out = apply_qrs(spec, STD, frozenset())
    assert out["instruments"] == ["ECOG1", "PHQ01"] and out["scored_items"] == 4
    assert spec["study_meta"]["cdisc_standards"]["qrs"]["tagged_by_item_order"] == 5
    no_qrs = {"study_meta": {}, "forms": [copy.deepcopy(before["forms"][2])]}
    same = copy.deepcopy(no_qrs)
    assert apply_qrs(no_qrs, STD, frozenset()) is None and no_qrs == same   # no questionnaire: spec not touched

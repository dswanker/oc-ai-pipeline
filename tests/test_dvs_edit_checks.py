"""Edit Checks sheet + AI proposals: sources, plain English, cross-form dependencies, proposals not applied."""
import copy, io, json, os, sys, tempfile, contextlib
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "conventions"))
import openpyxl
from conventions_engine import apply_conventions
import ai_edit_checks as ai
import edit_checks_sheet as ecs
from test_consent_floor import _spec as consent_spec
from test_core_checks import _spec as core_spec


def _engine(spec):
    with contextlib.redirect_stdout(io.StringIO()):
        apply_conventions(spec, study_id="T", customer_subdomain="", client_name="")
    return spec


def _rows(spec):
    return {(r["Form"], r["Item"], r["Check Type"], r["Check ID"]): r for r in ecs.build_rows(spec)}


def test_sources_and_cross_form_dependency():
    rows = ecs.build_rows(_engine(consent_spec()))
    floor = next(r for r in rows if r["Item"] == "ONSET" and r["Check ID"] == "CONSENT_FLOOR")
    assert floor["Source"] == "Global Rule" and floor["Cross-Form Dependency"] == "F_ICF.ICFDAT (any visit)"
    assert floor["Plain-English Description"] == "AE start must be on or after Consent date (F_ICF)."
    assert floor["Status"] == "In build"


def test_cdisc_sources_and_plain_english():
    rows = ecs.build_rows(_engine(core_spec()))
    ser = next(r for r in rows if r["Item"] == "SERIOUS" and r["Check Type"] == "Constraint")
    assert ser["Source"] == "CDISC CORE (CORE-000022)"
    assert ser["Plain-English Description"] == ("If Serious? is Yes, at least one of Results in death, Hospitalisation, "
                                                "Life threatening must be Yes.")
    why = next(r for r in rows if r["Item"] == "TEMPWHY" and r["Check Type"] == "Show-when")
    assert why["Source"] == "CDISC CORE (CORE-000440)" and why["Plain-English Description"] == \
        "Shown only when Temperature taken? is No."
    out = next(r for r in rows if r["Item"] == "OUTCOME" and r["Check Type"] == "Constraint")
    assert out["Plain-English Description"] == "If Outcome is Fatal, Results in death must be Yes."
    lif = next(r for r in rows if r["Item"] == "AESLIFE" and r["Check Type"] == "Show-when")
    assert lif["Source"] == "Study Build"                                # author's relevance, not a rule


PROPOSALS = json.dumps({"checks": [
    {"target_form": "F_AE", "target_field": "ONSET", "operator": ">=", "source_form": "DM", "source_field": "VISDAT",
     "message": "AE start must be on or after the visit date.", "rationale": "test", "category": "cross_form"},
    {"target_form": "F_AE", "target_field": "NOPE", "operator": ">=", "source_form": "DM", "source_field": "VISDAT"},
    {"target_form": "F_AE", "target_field": "ONSET", "operator": "=", "source_form": "DM", "source_field": "BRTHDAT",
     "when": {"field": "MHSTDAT", "equals": "Y"}}]})


def test_proposals_are_listed_not_applied_then_approval_applies():
    spec = _engine(consent_spec())
    before = copy.deepcopy(spec)
    v = ai.validate_response(spec, PROPOSALS)
    assert spec == before                                                 # validation never changes the build
    assert [p["id"] for p in v["proposals"]] == ["AI.001"]
    assert v["rejected"] == {"unknown_field": 1, "bad_condition": 1}
    spec["study_meta"]["ai_edit_checks"] = {"proposals": v["proposals"]}
    prop_row = next(r for r in ecs.build_rows(spec) if r["Check ID"] == "AI.001")
    assert prop_row["Source"] == "AI-Proposed" and prop_row["Status"].startswith("Proposed")
    assert prop_row["Cross-Form Dependency"] == "DM.VISDAT (any visit)"
    assert json.loads(prop_row["Machine Data"])["source_field"] == "VISDAT"
    assert ai.apply_proposal(spec, v["proposals"][0]) is True
    onset = next(r for f in spec["forms"] for r in f["survey"] if r.get("name") == "ONSET")
    assert "${VISDAT_CF}" in onset["constraint"] and "AI.001" in onset["edit_checks"]
    applied = next(r for r in ecs.build_rows(spec) if r["Check ID"] == "AI.001" and r["Status"] == "In build")
    assert applied["Source"] == "AI-Proposed"


def test_sheet_written_with_dropdown_and_hidden_machine_data():
    spec = _engine(consent_spec())
    spec["study_meta"]["ai_edit_checks"] = {"proposals": ai.validate_response(spec, PROPOSALS)["proposals"]}
    path = os.path.join(tempfile.mkdtemp(), "dvs.xlsx")
    openpyxl.Workbook().save(path)
    n = ecs.add_sheet(path, spec)
    wb = openpyxl.load_workbook(path)
    ws = wb["Edit Checks"]
    assert n == ws.max_row - 1 and [c.value for c in ws[1]] == ecs.COLUMNS
    assert ws.column_dimensions["N"].hidden is True
    assert any("Delete,Change,Approve,Reject,Add" in str(dv.formula1) for dv in ws.data_validations.dataValidation)
    assert "Edit Checks Guide" in wb.sheetnames

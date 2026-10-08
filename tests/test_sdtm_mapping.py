"""Unit tests for sdtm_mapping.py. Synthetic CDASHIG rows (no member file)."""
import io
import cdisc_cdash as cd
import sdtm_mapping as sm


def _rec(dom, var, target, mapping=""):
    return {"domain": dom, "variable": var, "label": var, "question": "", "prompt": "", "instruction": "",
            "core": "HR", "type": "Char", "sdtm_target": target, "mapping": mapping}


FIELDS = {("DM", "SEX"): _rec("DM", "SEX", "SEX", "Maps directly to the SDTMIG variable listed."),
          ("DM", "RACEOTH"): _rec("DM", "RACEOTH", "QVAL", "This does not map directly to an SDTMIG variable. "
                                                          "Submit in SUPPDM. Extra."),
          ("VS", "VSORRES"): _rec("VS", "VSORRES", "VSORRES"),
          ("CP", "VISIT"): _rec("CP", "VISIT", "VISIT"), ("LB", "VISIT"): _rec("LB", "VISIT", ""),
          ("CP", "VISDAT"): _rec("CP", "VISDAT", "", "This field is not an SDTM variable. Populate CPDTC from it.")}


def setup_module():
    cd._MEMO["fields"] = FIELDS


def teardown_module():
    cd._MEMO.clear()


SPEC = {"study_meta": {"protocol_number": "P1", "cdisc_standards": {"ct_version": "2026-09-25", "cdashig": "2.3"}},
        "forms": [
            {"form_id": "DM", "cdash_domain": "DM", "choices": [
                {"list_name": "ct_sex", "name": "F", "label": "Female", "cdisc_submission_value": "F"}],
             "survey": [
                {"type": "select_one ct_sex", "name": "SEX", "label": "Sex", "concept": "SEX",
                 "concept_source": "cdash_name", "cdisc_ct": {"codelist": "SEX", "codelist_code": "C66731"}},
                {"type": "text", "name": "RACE_OTH", "label": "Other race", "concept": "RACEOTH",
                 "concept_source": "claude"},
                {"type": "select_one v", "name": "VISIT", "label": "Visit", "concept": "VISIT",
                 "concept_source": "cdash_name"},
                {"type": "date", "name": "COLLDT", "label": "Visit date", "concept": "VISDAT",
                 "concept_source": "customer_alias"},
                {"type": "text", "name": "DIN", "label": "Donor ID"},
                {"type": "begin group", "name": "G"}, {"type": "calculate", "name": "X_CF"}]},
            {"form_id": "VS", "cdash_domain": "VS", "survey": [
                {"type": "decimal", "name": "BPSYS", "label": "Systolic", "concept": "VSORRES",
                 "concept_qualifier": "SYSBP", "concept_source": "claude"}]}]}


def _rows():
    return {r["Field"]: r for r in sm.build_rows(SPEC)}


def test_rows_mapping_and_confidence():
    r = _rows()
    assert set(r) == {"SEX", "RACE_OTH", "VISIT", "COLLDT", "DIN", "BPSYS"}          # data fields only
    assert (r["SEX"]["SDTM Domain"], r["SEX"]["SDTM Variable"], r["SEX"]["Confidence"], r["SEX"]["Basis"]) == \
        ("DM", "SEX", "High", "CDASH variable name")
    assert r["SEX"]["Controlled Terminology"] == "SEX (C66731)" and r["SEX"]["Choices"] == "F=Female"
    assert (r["RACE_OTH"]["SDTM Domain"], r["RACE_OTH"]["SDTM Variable"], r["RACE_OTH"]["Value-Level Detail"]) == \
        ("SUPPDM", "QVAL", "QNAM = RACEOTH")
    assert r["RACE_OTH"]["Confidence"] == "Medium" and r["RACE_OTH"]["Notes"].startswith("Review:")
    assert "Submit in SUPPDM." in r["RACE_OTH"]["Notes"] and "does not map directly" not in r["RACE_OTH"]["Notes"]
    assert (r["VISIT"]["SDTM Domain"], r["VISIT"]["SDTM Variable"]) == ("(each domain)", "VISIT")
    assert r["COLLDT"]["SDTM Variable"] == "(not submitted as its own variable)"
    assert "--DTC" in r["COLLDT"]["Notes"] and r["COLLDT"]["Basis"] == "Customer alias"
    assert (r["BPSYS"]["SDTM Variable"], r["BPSYS"]["Value-Level Detail"]) == ("VSORRES", "VSTESTCD = SYSBP")
    assert (r["DIN"]["Confidence"], r["DIN"]["Basis"]) == ("None", "Not mapped") and "SUPPDM.QVAL" in r["DIN"]["Notes"]


def test_summary_and_files():
    s = sm.summarize(sm.build_rows(SPEC))
    assert s == {"fields": 6, "mapped": 5, "by_confidence_basis": {
        "High / CDASH variable name": 2, "Medium / AI (validated)": 2, "High / Customer alias": 1,
        "None / Not mapped": 1}}
    out = sm.build_files(SPEC)
    assert out["pdf"][:5] == b"%PDF-"
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(out["xlsx"]))
    assert wb.sheetnames == ["Summary", "Mapping"]
    ws = wb["Mapping"]
    assert [c.value for c in ws[1]] == sm.COLUMNS and ws.max_row == 7

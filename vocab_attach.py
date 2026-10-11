"""Which vocabulary CSVs does a built XLSForm actually reference?

OpenClinica requires a form that uses select_one_from_file / select_multiple_from_file to be uploaded
together with the CSV it names. The uploader used to attach EVERY csv sitting next to the form, which is
fine for two small lists but wasteful (and risky) once several large lists exist. This returns only the
sibling CSVs the form's survey sheet really references.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List, Set

_REF = re.compile(r"select_(?:one|multiple)_from_file\s+(\S+?\.csv)", re.I)


def referenced_csvs(xlsx_path) -> Set[str]:
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    try:
        names: Set[str] = set()
        if "survey" in wb.sheetnames:
            for row in wb["survey"].iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        names.update(m.group(1) for m in _REF.finditer(cell))
        return names
    finally:
        wb.close()


def csvs_to_attach(xlsx_path, sibling_csvs: Iterable[Path]) -> List[Path]:
    refs = referenced_csvs(xlsx_path)
    return sorted((p for p in sibling_csvs if p.name in refs), key=lambda p: p.name)


def files_to_attach(xlsx_path) -> List[Path]:
    """Every CSV a form's upload needs: the vocabulary lists beside it and the lookups it reads with
    pulldata() or an external instance, which the build keeps in csv/ (lookup_csv.files_for_form).
    OpenClinica resolves pulldata('name', ...) against the media file name.csv of that form, so a lookup
    that is not uploaded with the form shows "Can't find name.csv" when the form opens."""
    import lookup_csv
    xlsx_path = Path(xlsx_path)
    out = {p.name: p for p in csvs_to_attach(xlsx_path, sorted(xlsx_path.parent.glob("*.csv")))}
    for p in lookup_csv.files_for_form(xlsx_path):
        out.setdefault(p.name, p)
    return [out[n] for n in sorted(out)]

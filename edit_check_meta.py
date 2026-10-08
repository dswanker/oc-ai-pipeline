"""edit_checks_sheet.py: the "Edit Checks" sheet of the DVS workbook. One row per check in the study, organized by
form: what it does in plain English, its exact logic, where it came from (Source), what other form/item it reads
(Cross-Form Dependency), plus AI-proposed checks that are NOT in the build until approved. DMs edit the Action column
(Delete / Change / Approve / Reject) and add rows (Action = Add, plain-English description) for the next run.
"""
import json, os, re

COLUMNS = ["Action", "Source", "Check ID", "Form", "Item", "Item Label", "Cross-Form Dependency", "Check Type",
           "Plain-English Description", "Logic", "Query Message", "Rationale / Protocol Reference", "Status",
           "Machine Data"]  # hidden: structured definition used when the sheet is read back
ACTIONS = ["Delete", "Change", "Approve", "Reject", "Add"]

_REF = re.compile(r"\$\{(\w+)\}")
_NON_DATA = ("note", "begin group", "end group", "begin repeat", "end repeat")


def _global_ids():
    d = os.path.join(os.path.dirname(os.path.abspath(__file__)), "conventions_engine", "conventions", "global")
    try:
        return {f[:-5] for f in os.listdir(d) if f.endswith(".json")}
    except Exception:
        return set()


def source_of_check_id(cid):
    c = str(cid or "")
    if c.startswith("CORE-"):
        return f"CDISC CORE ({c})"
    if c.startswith(("SDTM.", "CDISC.")):
        return "CDISC Standard"
    if c.startswith("AI."):
        return "AI-Proposed"
    if c.startswith("DM."):
        return "DM-Added"
    return "Global Rule"


def source_of_convention(conv_id, global_ids):
    c = str(conv_id or "")
    if not c:
        return "Study Build"
    if "cdisc_reason_not_done" in c:
        return "CDISC CORE (CORE-000440)"
    if "cdisc_dose_only_when_occurred" in c:
        return "CDISC CORE (CORE-000004/137)"
    if "cdisc_" in c:
        return "CDISC Standard"
    return "Global Rule" if c in global_ids else "Customer Rule"


def _split_and(expr):
    try:
        import sys
        sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "skills", "dvs-specification", "scripts"))
        from extract_dvs_from_forms import _split_and_clauses
        return _split_and_clauses(expr) or [expr]
    except Exception:
        return [expr]


class _FormCtx:
    """Labels, choice labels and cross-form helpers for one form, used to describe logic in plain English."""

    def __init__(self, form, survey, all_forms):
        self.form = form
        self.survey = survey
        self.labels = {r.get("name"): str(r.get("label") or r.get("name")) for r in survey if isinstance(r, dict)}
        self.types = {r.get("name"): str(r.get("type") or "").split(" ")[0].lower() for r in survey if isinstance(r, dict)}
        self.choices = {}
        for r in survey:
            t = str((r or {}).get("type") or "")
            if isinstance(r, dict) and t.startswith("select") and " " in t:
                ln = t.split(" ", 1)[1].strip()
                self.choices[r.get("name")] = {str(c.get("name")): str(c.get("label") or c.get("name"))
                                               for c in form.get("choices") or [] if c.get("list_name") == ln}
        self.helpers = {}
        for r in survey:
            calc = str((r or {}).get("calculation") or "") if isinstance(r, dict) else ""
            m = re.search(r"FormOID='([^']+)'.*?ItemName='([^']+)'", calc)
            if m and "instance('clinicaldata')" in calc:
                ev = re.search(r"StudyEventOID='([^']+)'", calc)
                src_label = m.group(2)
                for f in all_forms:
                    fid = str(f.get("form_id") or "")
                    if m.group(1) in (fid, "F_" + fid):
                        src_label = next((str(x.get("label") or x.get("name")) for x in f.get("survey") or []
                                          if isinstance(x, dict) and x.get("name") == m.group(2)), src_label)
                self.helpers[r.get("name")] = (m.group(1), m.group(2), ev.group(1) if ev else "any visit", src_label)

    def lab(self, name):
        if name in self.helpers:
            fo, item, _ev, label = self.helpers[name]
            return f"{label} ({fo})"
        return self.labels.get(name, name)

    def choice(self, name, code):
        return self.choices.get(name, {}).get(code, code)

    def dependency(self, text):
        deps = [f"{fo}.{item} ({ev})" for n in _REF.findall(str(text or "")) if n in self.helpers
                for fo, item, ev, _l in [self.helpers[n]]]
        return "; ".join(dict.fromkeys(deps))


_WORDS = {">=": ("on or after", "at least"), ">": ("after", "greater than"), "<=": ("on or before", "at most"),
          "<": ("before", "less than"), "=": ("equal to", "equal to"), "!=": ("different from", "different from")}


def describe(clause, field, cx):
    """Plain-English description of one check clause on `field` (deterministic patterns; else a pointer)."""
    e = " ".join(str(clause or "").split())
    me = cx.lab(field)
    is_date = cx.types.get(field) in ("date", "datetime")
    m = re.match(r"^\$\{(\w+)\}\s*!=\s*'([^']*)'\s+or\s+\((.*)\)$", e)
    if m:
        return f"When {cx.lab(m.group(1))} is {cx.choice(m.group(1), m.group(2))}: " + describe(m.group(3), field, cx)
    m = re.match(r"^\$\{(\w+)\}\s*!=\s*'([^']*)'\s+or\s+\$\{(\w+)\}\s*=\s*''\s+or\s+\.\s*=\s*\$\{\3\}$", e)
    if m:
        return f"When {cx.lab(m.group(1))} is {cx.choice(m.group(1), m.group(2))}, {me} must equal {cx.lab(m.group(3))}."
    m = (re.match(r"^\.\s*=\s*''\s+or\s+\$\{(\w+)\}\s*=\s*''\s+or\s+\.\s*(>=|>|<=|<|=|!=)\s*\$\{\1\}$", e)
         or re.match(r"^\.\s*(>=|>|<=|<|=|!=)\s*\$\{(\w+)\}$", e))
    if m:
        ref, op = (m.group(1), m.group(2)) if m.re.pattern.startswith("^\\.\\s*=\\s*''") else (m.group(2), m.group(1))
        return f"{me} must be {_WORDS[op][0 if is_date else 1]} {cx.lab(ref)}."
    m = re.match(r"^\.\s*(<=|<|>=|>)\s*today\(\)$", e)
    if m:
        return {"<=": f"{me} cannot be in the future.", "<": f"{me} must be before today.",
                ">=": f"{me} must be today or later.", ">": f"{me} must be in the future."}[m.group(1)]
    m = re.match(r"^\.\s*(=|!=)\s*'([^']*)'$", e)
    if m:
        return f"{me} must {'be' if m.group(1) == '=' else 'not be'} {cx.choice(field, m.group(2))}."
    m = re.match(r"^\.\s*!=\s*'([^']*)'\s+or\s+\((.*)\)$", e)
    if m:
        parts = re.findall(r"\$\{(\w+)\}\s*=\s*'([^']*)'", m.group(2))
        if parts and " and " not in m.group(2):
            return (f"If {me} is {cx.choice(field, m.group(1))}, at least one of "
                    f"{', '.join(cx.lab(n) for n, _ in parts)} must be {cx.choice(parts[0][0], parts[0][1])}.")
    m = re.match(r"^\.\s*!=\s*'([^']*)'\s+or\s+\$\{(\w+)\}\s*=\s*'([^']*)'$", e)
    if m:
        return (f"If {me} is {cx.choice(field, m.group(1))}, {cx.lab(m.group(2))} must be "
                f"{cx.choice(m.group(2), m.group(3))}.")
    m = re.match(r"^\(?\.\s*(>=|>)\s*(-?[\d.]+)\s+and\s+\.\s*(<=|<)\s*(-?[\d.]+)\)?$", e)
    if m:
        return f"{me} must be between {m.group(2)} and {m.group(4)}."
    m = re.match(r"^\.\s*(>=|>|<=|<)\s*(-?[\d.]+)$", e)
    if m:
        return f"{me} must be {_WORDS[m.group(1)][1]} {m.group(2)}."
    return "Custom logic (see Logic column)."


def describe_relevant(rel, cx):
    e = " ".join(str(rel or "").split())
    m = re.match(r"^\$\{(\w+)\}\s*=\s*'([^']*)'$", e)
    if m:
        return f"Shown only when {cx.lab(m.group(1))} is {cx.choice(m.group(1), m.group(2))}."
    m = re.match(r"^selected\(\s*\$\{(\w+)\}\s*,\s*'([^']*)'\s*\)$", e)
    if m:
        return f"Shown only when {cx.lab(m.group(1))} includes {cx.choice(m.group(1), m.group(2))}."
    return "Shown only when its condition is met (see Logic column)."


def _built_rows(forms_json, form_id):
    forms = (forms_json or {}).get("forms") or {}
    for key in (f"{form_id}.xlsx", f"F_{form_id}.xlsx"):
        if key in forms:
            return {r.get("name"): r for r in (forms[key] or {}).get("survey") or [] if isinstance(r, dict)}
    return {}


def build_rows(struct_json, forms_json=None):
    """All checks in the study (by form), then AI proposals for that form. List of dicts keyed by COLUMNS."""
    gids = _global_ids()
    all_forms = [f for f in (struct_json or {}).get("forms") or [] if isinstance(f, dict)]
    proposals = (((struct_json or {}).get("study_meta") or {}).get("ai_edit_checks") or {}).get("proposals") or []
    rows = []
    for f in all_forms:
        fid = str(f.get("form_id") or "")
        built = _built_rows(forms_json, fid)
        survey = [{**r, **{k: built[r.get("name")][k] for k in ("constraint", "constraint_message", "relevant", "required")
                           if r.get("name") in built and k in built[r.get("name")]}}
                  for r in f.get("survey") or [] if isinstance(r, dict)]
        cx = _FormCtx(f, survey, all_forms)
        for r in survey:
            name, t = r.get("name"), str(r.get("type") or "").lower()
            if not name or t in _NON_DATA or t == "calculate":
                continue
            exprs = r.get("edit_check_exprs") or {}
            by_clause = {" ".join(str(v).split()): k for k, v in exprs.items() if v}
            details = {d.get("id"): d for d in r.get("edit_check_details") or [] if isinstance(d, dict)}
            set_by = r.get("edit_check_set_by") or {}
            base = {"Action": "", "Form": fid, "Item": name, "Item Label": str(r.get("label") or ""), "Status": "In build"}
            con = str(r.get("constraint") or "").strip()
            if con:
                for n, clause in enumerate(_split_and(con), 1):
                    cid = by_clause.get(" ".join(clause.split()))
                    det = details.get(cid) or {}
                    src = det.get("source") or (source_of_check_id(cid) if cid else
                                                source_of_convention(set_by.get("constraint"), gids))
                    rows.append({**base, "Source": src, "Check ID": cid or f"{fid}.{name}.C{n}",
                                 "Cross-Form Dependency": cx.dependency(clause), "Check Type": "Constraint",
                                 "Plain-English Description": describe(clause, name, cx), "Logic": clause,
                                 "Query Message": str(r.get("constraint_message") or ""),
                                 "Rationale / Protocol Reference": " ".join(
                                     x for x in (det.get("rationale"), det.get("protocol_reference")) if x)})
            rel = str(r.get("relevant") or "").strip()
            if rel:
                rows.append({**base, "Source": source_of_convention(set_by.get("relevant"), gids),
                             "Check ID": f"{fid}.{name}.SHOW", "Cross-Form Dependency": cx.dependency(rel),
                             "Check Type": "Show-when", "Plain-English Description": describe_relevant(rel, cx),
                             "Logic": rel, "Query Message": "", "Rationale / Protocol Reference": ""})
            if str(r.get("required") or "").strip().lower() in ("yes", "true", "1"):
                rows.append({**base, "Source": source_of_convention(set_by.get("required"), gids),
                             "Check ID": f"{fid}.{name}.REQ", "Cross-Form Dependency": "", "Check Type": "Required",
                             "Plain-English Description": f"{cx.lab(name)} is required" + (" when shown." if rel else "."),
                             "Logic": "required", "Query Message": str(r.get("required_message") or ""),
                             "Rationale / Protocol Reference": ""})
        for p in proposals:
            if str(p.get("target_form")) != fid:
                continue
            rows.append({"Action": "", "Source": "AI-Proposed", "Check ID": p.get("id"), "Form": fid,
                         "Item": p.get("target_field"), "Item Label": cx.labels.get(p.get("target_field"), ""),
                         "Cross-Form Dependency": (p.get("cross_form") or "") + (" (any visit)" if p.get("cross_form") else ""),
                         "Check Type": "Constraint", "Plain-English Description": p.get("message"),
                         "Logic": p.get("logic"), "Query Message": p.get("message"),
                         "Rationale / Protocol Reference": " ".join(
                             x for x in (p.get("rationale"), p.get("protocol_reference")) if x),
                         "Status": "Proposed: not in build until Approved",
                         "Machine Data": json.dumps({k: p.get(k) for k in ("id", "target_form", "target_field", "operator",
                                                     "source_form", "source_field", "when", "message", "rationale",
                                                     "protocol_reference", "category")})})
    return rows


def add_sheet(xlsx_path, struct_json, forms_json=None):
    """Append the 'Edit Checks' sheet (and a short guide) to the DVS workbook. Never raises."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
        rows = build_rows(struct_json, forms_json)
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.create_sheet("Edit Checks")
        ws.append(COLUMNS)
        head, band = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1B3A6B")
        fills = {"AI-Proposed": "FDEBD0", "DM-Added": "E8DAEF"}
        for c in ws[1]:
            c.font, c.fill = head, band
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append([r.get(c, "") for c in COLUMNS])
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if r["Source"] in fills:
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=fills[r["Source"]])
        for i, w in enumerate([10, 22, 22, 12, 16, 28, 26, 13, 52, 46, 40, 40, 22, 10], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.column_dimensions[get_column_letter(len(COLUMNS))].hidden = True
        dv = DataValidation(type="list", formula1='"' + ",".join(ACTIONS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"A2:A{max(ws.max_row + 200, 300)}")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS) - 1)}{max(ws.max_row, 2)}"
        g = wb.create_sheet("Edit Checks Guide")
        for line in [
            ["How to use the Edit Checks sheet"],
            ["Every check in the study is listed by form. Leave Action blank to keep a check as it is."],
            ["Delete: remove the check from the build.   Change: edit Plain-English Description and/or Query Message."],
            ["Approve / Reject: for AI-Proposed rows (orange). Nothing proposed is in the build until Approved."],
            ["Add: add a new row with Form, Item (optional), and a Plain-English Description. Logic is filled in by the pipeline."],
            ["Upload the edited DVS and re-run. Results appear in the Status column of the regenerated sheet."],
            ["Source: CDISC CORE (rule id), CDISC Standard, Global Rule, Customer Rule, Study Build, AI-Proposed, DM-Added."],
            ["Cross-Form Dependency: the other form and item (and visit) this check reads."]]:
            g.append(line)
        g["A1"].font = Font(bold=True, size=13)
        g.column_dimensions["A"].width = 120
        wb.save(xlsx_path)
        return len(rows)
    except Exception as e:
        print(f"Edit Checks sheet skipped: {e}", flush=True)
        return None

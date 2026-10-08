"""dvs_edits.py: read the DM's decisions from an edited DVS (DVS_OC4 "Action" column) and apply them to the Study
Spec deterministically, before the normal EDC builder runs. Replaces the old path where Claude regenerated every
form from an edited DVS. Only DVS files with the Action column use this path; older DVS files are untouched.

  Approve  AI-proposed row -> ai_edit_checks.apply_proposal (validated, same path as every engine check)
  Reject   AI-proposed row -> dropped from the proposals and never proposed again
  Delete   any check       -> the clause / show-when / required flag is removed; engine checks are suppressed so the
                              rules engine never re-adds them
  Change   any check       -> query message updated (a changed plain-English logic description needs translation)
  Add      new row         -> plain-English description translated by AI into a structured check, validated
                              against the study and applied as DM.nnn (Check Source DM-Added); anything that
                              cannot be expressed safely is "needs_build_team", never guessed

Every decision is recorded in study_meta.edit_check_decisions with its outcome.
"""
import io, json

ACTIONS = ("approve", "reject", "delete", "change", "add")


def parse_actions(xlsx_bytes):
    """Rows of DVS_OC4 that carry an Action. None when the DVS is not the new (Action column) format."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    except Exception:
        return None
    if "DVS_OC4" not in wb.sheetnames:
        return None
    ws = wb["DVS_OC4"]
    header_row = next((i for i in range(1, 6) if str(ws.cell(row=i, column=1).value or "").strip() == "Action"), None)
    if header_row is None:
        return None
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row]]
    out = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        row = {h: ("" if v is None else str(v).strip()) for h, v in zip(hdr, r) if h}
        act = row.get("Action", "").strip().lower()
        if act in ACTIONS:
            row["Action"] = act
            out.append(row)
    return out


def _forms(spec):
    return [f for f in spec.get("forms") or [] if isinstance(f, dict)]


def _find(spec, form_oid, item):
    for f in _forms(spec):
        fid = str(f.get("form_id") or "")
        if form_oid in (fid, "F_" + fid) or fid == "F_" + str(form_oid):
            row = next((r for r in f.get("survey") or [] if isinstance(r, dict) and r.get("name") == item), None)
            return f, row
    return None, None


def _norm(x):
    return " ".join(str(x or "").split())


def _split(expr):
    try:
        import os, sys
        sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "skills", "dvs-specification", "scripts"))
        from extract_dvs_from_forms import _split_and_clauses
        return _split_and_clauses(expr) or [expr]
    except Exception:
        return [expr]


def _delete(spec, row):
    f, r = _find(spec, row.get("Target Form OID"), row.get("Target Item Name"))
    if r is None:
        return "not_found", "form/item not in the study"
    ctype, rule = row.get("Check Type", ""), row.get("Rule / Proposal ID", "")
    if ctype in ("Constraint", "Cross-form") and "required" not in row.get("Expression / Calculation", ""):
        target = _norm(row.get("Expression / Calculation"))
        cur = str(r.get("constraint") or "")
        if not cur:
            return "already", "no constraint on this item"
        clauses = _split(cur)
        keep = [c for c in clauses if _norm(c) != target]
        if len(keep) == len(clauses):
            if _norm(cur) == target:
                keep = []
            else:
                return "not_found", "that logic is no longer on the item"
        removed = [c for c in clauses if c not in keep]
        # engine rules that stood aside because this author clause covered them must not come back
        for cid, ref in (r.get("edit_check_covered_by") or {}).items():
            if any(ref in c for c in removed) and cid not in (r.get("edit_checks_suppressed") or []):
                r.setdefault("edit_checks_suppressed", []).append(cid)
                if cid in (r.get("edit_checks") or []):
                    r["edit_checks"].remove(cid)
        r["constraint"] = keep[0] if len(keep) == 1 else " and ".join(f"({c})" for c in keep)
        msgs = r.get("edit_check_msgs") or {}
        if rule and msgs.get(rule) and r.get("constraint_message"):
            r["constraint_message"] = _norm(str(r["constraint_message"]).replace(msgs[rule], ""))
        if not r["constraint"]:
            r.pop("constraint", None)
            r.pop("constraint_message", None)
        if rule:
            for k in ("edit_checks",):
                if rule in (r.get(k) or []):
                    r[k].remove(rule)
            (r.get("edit_check_exprs") or {}).pop(rule, None)
            if rule not in (r.get("edit_checks_suppressed") or []):
                r.setdefault("edit_checks_suppressed", []).append(rule)
        return "applied", "check removed"
    if ctype in ("Conditional Display",) or (ctype == "Cross-form" and _norm(r.get("relevant")) == _norm(row.get("Expression / Calculation"))):
        if not r.get("relevant"):
            return "already", "no show-when on this item"
        r.pop("relevant", None)
        r.setdefault("edit_check_suppressed_paths", []).append("relevant")
        return "applied", "show-when removed (always shown)"
    if ctype == "Required":
        r["required"] = "no"
        r.setdefault("edit_check_suppressed_paths", []).append("required")
        return "applied", "no longer required"
    return "unsupported", f"delete not supported for check type {ctype!r}"


def _approve(spec, row):
    import ai_edit_checks as ai
    try:
        prop = json.loads(row.get("Machine Data") or "{}")
    except Exception:
        prop = {}
    if not prop.get("id"):
        return "unsupported", "Approve applies to proposed rows only"
    if prop.get("kind") == "standard_proposal":
        # logic proposed on a customer standard form (rules engine / AI-suggested): standards_match applies it
        import standards_match as sm
        full = prop if prop.get("ops") else next((p for p in sm.proposals(spec) if p.get("id") == prop["id"]), None)
        if full is None:
            f = next((x for x in _forms(spec) if sm.norm_id(x.get("form_id")) == sm.norm_id(prop.get("target_form"))), None)
            done = any(a.get("id") == prop["id"] for a in ((f or {}).get("customer_standard") or {}).get("approved") or [])
            return ("already", "already in the build") if done else ("not_found", "proposal no longer listed")
        return sm.apply_proposal(spec, full)
    f, r = _find(spec, prop.get("target_form"), prop.get("target_field"))
    if r is not None and prop["id"] in (r.get("edit_checks") or []):
        return "already", "already in the build"
    if ai.apply_proposal(spec, prop):
        props = ((spec.get("study_meta") or {}).get("ai_edit_checks") or {}).get("proposals") or []
        props[:] = [p for p in props if p.get("id") != prop["id"]]
        return "applied", "added to the build"
    return "failed", "could not be applied to the current study (fields changed or an equivalent check exists)"


def _reject(spec, row):
    pid = row.get("Rule / Proposal ID") or row.get("Check ID")
    try:
        machine = json.loads(row.get("Machine Data") or "{}")
    except Exception:
        machine = {}
    if machine.get("kind") == "standard_proposal":
        import standards_match as sm
        return ("applied", "proposal rejected") if sm.reject_proposal(spec, machine.get("id")) \
            else ("already", "proposal no longer listed")
    ae = spec.setdefault("study_meta", {}).setdefault("ai_edit_checks", {})
    props = ae.get("proposals") or []
    before = len(props)
    props[:] = [p for p in props if p.get("id") != pid]
    ae.setdefault("rejected_ids", [])
    if pid not in ae["rejected_ids"]:
        ae["rejected_ids"].append(pid)
    return ("applied", "proposal rejected") if len(props) < before else ("already", "proposal no longer listed")


def _change(spec, row):
    f, r = _find(spec, row.get("Target Form OID"), row.get("Target Item Name"))
    if r is None:
        return "not_found", "form/item not in the study"
    new_msg = row.get("Constraint / Required / Relevant Message", "")
    if new_msg and row.get("Check Type") in ("Constraint", "Cross-form") and new_msg != r.get("constraint_message"):
        rule = row.get("Rule / Proposal ID", "")
        if rule and (r.get("edit_check_msgs") or {}).get(rule) == new_msg:
            return "already", "message already updated"
        old_piece = (r.get("edit_check_msgs") or {}).get(rule)
        cur = str(r.get("constraint_message") or "")
        if old_piece and old_piece in cur:
            # this check's own sentence only; the field's other checks keep their messages
            r["constraint_message"] = cur.replace(old_piece, new_msg)
            r["edit_check_msgs"][rule] = new_msg
        else:
            r["constraint_message"] = new_msg
        r["constraint_message_locked"] = True  # the DM's wording: the EDC builder must not normalize it
        return "applied", "query message updated (logic unchanged)"
    return "pending", "logic changes from the plain-English description are translated in a later step"


def apply_actions(spec, rows, add_translations=None):
    """Apply DM decisions to the spec. add_translations: {add_key: translation} for Add rows (from the AI call in the
    pipeline); None means translation was not available this run. Returns counts + per-row results."""
    results = []
    done_adds = {}
    for d in (spec.get("study_meta") or {}).get("edit_check_decisions") or []:
        if d.get("action") == "add" and d.get("status") in ("applied", "already", "needs_build_team"):
            done_adds.setdefault(d.get("key"), d)
    for row in rows or []:
        act = row["Action"]
        if act == "add":
            key = _add_key(row)
            if key in done_adds:
                prev = done_adds[key]
                if prev.get("status") == "needs_build_team":  # same wording as before: same outcome, no new AI call
                    status, note = "needs_build_team", prev.get("note") or "needs the build team"
                else:
                    status, note = "already", "added on an earlier run"
            elif add_translations is None:
                status, note = "pending", "plain-English translation not available on this run"
            else:
                status, note = _apply_add(spec, row, add_translations.get(key))
            results.append({"action": act, "key": key, "check_id": "", "rule_id": "",
                            "form": row.get("Target Form OID"), "item": row.get("Target Item Name"),
                            "description": row.get("Plain-English Description"), "status": status, "note": note})
            continue
        try:
            if act == "approve":
                status, note = _approve(spec, row)
            elif act == "reject":
                status, note = _reject(spec, row)
            elif act == "delete":
                status, note = _delete(spec, row)
            elif act == "change":
                status, note = _change(spec, row)
            else:
                status, note = "unsupported", f"unknown action {act!r}"
        except Exception as e:
            status, note = "failed", f"error: {e}"
        results.append({"action": act, "check_id": row.get("Check ID"), "rule_id": row.get("Rule / Proposal ID"),
                        "form": row.get("Target Form OID"), "item": row.get("Target Item Name"),
                        "description": row.get("Plain-English Description"), "status": status, "note": note})
    dec = spec.setdefault("study_meta", {}).setdefault("edit_check_decisions", [])
    dec.extend(results)
    summ = {k: sum(1 for r in results if r["status"] == k) for k in ("applied", "already", "failed", "pending",
                                                                      "not_found", "unsupported", "needs_build_team")}
    return {**summ, "results": results}


# ── DM plain-English additions ──────────────────────────────────────────────────

ADD_PROMPT = """You translate data managers' plain-English edit-check requests into structured checks for an
OpenClinica EDC study. The STUDY (forms, visits, fields, choices, existing checks) and the numbered REQUESTS follow.

For each request return exactly one entry. If it can be expressed as "target_field OPERATOR source_field" (optionally
only when a field on the target form equals a choice code), return the check; otherwise return "cannot" with a short
reason (e.g. needs a calculation, a range of constants, or a field that does not exist).
Answer ONLY with JSON, no prose, no code fences:
{"results": [{"index": <request number>, "target_form": "...", "target_field": "...", "operator": ">=|>|<=|<|=|!=",
              "source_form": "...", "source_field": "...", "when": {"field": "...", "equals": "<code>"} (optional),
              "message": "<query text for the site, under 120 characters>"}
             | {"index": <request number>, "cannot": "<reason>"}]}
Use only form ids, field names and choice codes that appear in the STUDY. Dates compare with dates, numbers with numbers.
"""


def _add_key(row):
    return _norm(f"{row.get('Target Form OID')}|{row.get('Target Item Name')}|{row.get('Plain-English Description')}").lower()


def pending_adds(spec, rows):
    """Add rows that still need translation (not already applied from an earlier upload of the same DVS)."""
    done = {d.get("key") for d in (spec.get("study_meta") or {}).get("edit_check_decisions") or []
            if d.get("action") == "add" and d.get("status") in ("applied", "already", "needs_build_team")}
    return [r for r in rows or [] if r["Action"] == "add" and r.get("Plain-English Description")
            and _add_key(r) not in done]


def build_add_request(spec, adds):
    import ai_edit_checks as ai
    _p, study = ai.build_request(spec)
    reqs = "\n".join(f"{i}. form {r.get('Target Form OID') or '(any)'} | item {r.get('Target Item Name') or '(any)'} | "
                     f"{r.get('Plain-English Description')}" for i, r in enumerate(adds, 1))
    return ADD_PROMPT, study + "\n\nREQUESTS:\n" + reqs


def parse_add_response(text, adds):
    """{index: proposal-dict | {"cannot": reason}} for the add rows (1-based index)."""
    t = str(text or "").strip()
    t = t[t.find("{"): t.rfind("}") + 1] if "{" in t else "{}"
    try:
        res = json.loads(t).get("results") or []
    except Exception:
        return {}
    return {int(r["index"]): r for r in res if isinstance(r, dict) and str(r.get("index", "")).isdigit()
            and 1 <= int(r["index"]) <= len(adds)}


def _apply_add(spec, row, translation):
    import ai_edit_checks as ai
    if translation is None:
        return "needs_build_team", "no translation was produced"
    if translation.get("cannot"):
        return "needs_build_team", f"not expressible as a standard check: {translation['cannot']}"
    norm, why = ai._resolve(spec, translation)
    if why:
        return "needs_build_team", f"translation did not validate against the study ({why})"
    n = 1 + sum(1 for f in _forms(spec) for r in f.get("survey") or [] if isinstance(r, dict)
                for c in (r.get("edit_checks") or []) if str(c).startswith("DM."))
    prop = {**norm, "id": f"DM.{n:03d}", "rationale": f"DM request: {row.get('Plain-English Description')}"}
    if ai.apply_proposal(spec, prop, source="DM-Added"):
        return "applied", f"added as {prop['id']}: {norm['message']}"
    return "needs_build_team", "could not be built on the current study"

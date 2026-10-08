"""sdtm_mapping.py: SDTM Mapping Specification (XLSX + PDF) for a finished Study Spec. No AI.

One row per data field in the customer's build: customer field -> CDASH concept (+qualifier) ->
SDTM domain.variable, value-level detail, controlled terminology, with Confidence and Basis:

  High   / Customer alias        deterministic, from the customer's own alias list
  High   / CDASH variable name   deterministic, the field is already named with a CDASH variable
  Medium / QRS instrument        deterministic: instrument named on the form, items matched to CDISC test
                                 codes by item order; review
  Medium / AI (validated)        AI judgement, validated against CDASHIG and the field type; review
  None   / Not mapped            no CDASH equivalent identified

SDTM targets and notes come from CDASHIG v2.3 metadata (SDTMIG Target + Mapping Instructions).
This is a proposed mapping for review; derivations (e.g. --DTC from date + time, study day) are
described, not performed.
"""
import datetime, io, re

import cdisc_cdash
import cdisc_concepts
import cdisc_ct

COLUMNS = ["Form", "Field", "Label", "Type", "Choices", "CDASH Concept", "Qualifier", "SDTM Domain",
           "SDTM Variable", "Value-Level Detail", "Controlled Terminology", "Confidence", "Basis", "Notes"]

BASIS = {"customer_alias": ("High", "Customer alias"),
         "cdash_name": ("High", "CDASH variable name"),
         "qrs_instrument": ("Medium", "QRS instrument (item order)"),
         "claude": ("Medium", "AI (validated)"),
         None: ("None", "Not mapped")}

# concept -> (value-level variable template) for repeated uses told apart by a qualifier
_VLM = {"VSORRES": "VSTESTCD", "VSPERF": "VSTESTCD", "VSSTAT": "VSTESTCD", "SCORRES": "SCTESTCD",
        "LBORRES": "LBTESTCD", "EGORRES": "EGTESTCD", "IEORRES": "IECAT", "QSORRES": "QSTESTCD"}
_DIRECT = "Maps directly to the SDTMIG variable"
_QRS = ("QSORRES", "FTORRES", "RSORRES")


def _choices_text(form, row):
    t = str(row.get("type") or "")
    if " " not in t or not t.lower().startswith("select"):
        return ""
    ln = t.split(" ", 1)[1].strip()
    out = []
    for c in form.get("choices") or []:
        if isinstance(c, dict) and c.get("list_name") == ln:
            code = c.get("cdisc_submission_value") or c.get("name")
            out.append(f"{code}={c.get('label')}" if str(code) != str(c.get("label")) else str(code))
    return "; ".join(out[:12]) + (f"; ... (+{len(out) - 12})" if len(out) > 12 else "")


_BOILERPLATE = ("This does not map directly to an SDTMIG variable.", "This field is not an SDTM variable.")


def _first_sentence(text, limit=220):
    t = " ".join(str(text or "").split())
    if not t or t.startswith(_DIRECT):
        return ""
    for b in _BOILERPLATE:  # skip CDASHIG's opening boilerplate; keep the actual instruction
        if t.startswith(b) and len(t) > len(b) + 5:
            t = t[len(b):].strip()
    cut = t.split(". ")[0]
    cut = cut if cut.endswith(".") else cut + "."
    return cut if len(cut) <= limit else cut[: limit - 3] + "..."


def _value_level(concept, qual, sdtm_domain):
    if not qual:
        return ""
    if concept in _VLM:
        return f"{_VLM[concept]} = {qual}"
    if concept.startswith("SU"):
        return f"SUCAT = {qual}"
    if concept.endswith("OCCUR"):
        return f"{concept[:2]}TERM = {qual} (pre-specified)"
    return f"Qualifier: {qual}"


def build_rows(spec):
    """List of dicts keyed by COLUMNS, in form order. Uses concept tags already on the spec."""
    fields = cdisc_cdash.load_cdashig_fields()
    by_var = {}
    for (_d, v), rec in fields.items():
        by_var.setdefault(v, []).append(rec)
    rows = []
    for form in spec.get("forms") or []:
        if not isinstance(form, dict):
            continue
        fdom = str(form.get("cdash_domain") or "").strip().upper()
        for r in form.get("survey") or []:
            if not cdisc_concepts._is_data_field(r):
                continue
            concept = str(r.get("concept") or "").upper()
            qual = r.get("concept_qualifier") or ""
            source = r.get("concept_source") if concept else None
            conf, basis = BASIS.get(source, BASIS[None])
            ct = r.get("cdisc_ct") if isinstance(r.get("cdisc_ct"), dict) else {}
            row = {"Form": form.get("form_id"), "Field": r.get("name"), "Label": r.get("label") or "",
                   "Type": str(r.get("type") or "").split(" ")[0], "Choices": _choices_text(form, r),
                   "CDASH Concept": concept, "Qualifier": qual, "SDTM Domain": "", "SDTM Variable": "",
                   "Value-Level Detail": "", "Controlled Terminology":
                       f"{ct['codelist']} ({ct['codelist_code']})" if ct.get("codelist") else "",
                   "Confidence": conf, "Basis": basis, "Notes": ""}
            if concept in _QRS and qual and (isinstance(r.get("qrs"), dict) or concept not in by_var):
                # questionnaire / rating / scale item: --ORRES with --TESTCD = item code, --CAT = instrument
                q, dom = (r.get("qrs") if isinstance(r.get("qrs"), dict) else {}), concept[:2]
                row["SDTM Domain"], row["SDTM Variable"] = dom, concept
                row["Value-Level Detail"] = f"{dom}TESTCD = {qual}" + (
                    f"; {dom}CAT = {q['category']}" if q.get("category") else "")
                if q.get("orres_codelist"):
                    row["Controlled Terminology"] = f"{q['orres_codelist']} ({q.get('orres_codelist_code')})"
                notes = [f"{dom}TEST = {q['test']}." if q.get("test") else ""]
                if q.get("scores"):
                    notes.append(f"{dom}STRESN = response score ("
                                 + ", ".join(f"{k}={v}" for k, v in list(q["scores"].items())[:8]) + ").")
                if source in ("claude", "qrs_instrument"):
                    notes.insert(0, "Review: mapped by AI, validated against the instrument's CDISC test codes."
                                 if source == "claude" else "Review: matched to the instrument by item order.")
                row["Notes"] = " ".join(n for n in notes if n)
            elif concept:
                rec = cdisc_cdash._resolve(fields, by_var, fdom, concept) or {}
                dom = rec.get("domain") or fdom
                generic = bool(rec) and dom != fdom and not concept.startswith(dom)
                if generic:  # e.g. VISIT: a timing variable carried by every domain collected at the visit
                    dom = "(each domain)"
                target = rec.get("sdtm_target") or ""
                if target == "QVAL":
                    row["SDTM Domain"], row["SDTM Variable"] = f"SUPP{dom}", "QVAL"
                    row["Value-Level Detail"] = f"QNAM = {concept}"
                elif not target:
                    row["SDTM Variable"] = "(not submitted as its own variable)"
                else:
                    row["SDTM Domain"] = dom
                    row["SDTM Variable"] = target
                    row["Value-Level Detail"] = _value_level(concept, qual, dom)
                row["Notes"] = _first_sentence(rec.get("mapping"), 320 if not target else 220) or (
                    "Used to derive other SDTM variables (e.g. --DTC)." if not target else "")
                if generic and rec.get("domain"):  # domain-neutral wording: CPDTC -> --DTC
                    row["Notes"] = re.sub(r"\b" + rec["domain"] + r"([A-Z]{3,})\b", r"--\1", row["Notes"])
                if source == "claude":
                    row["Notes"] = ("Review: mapped by AI, validated against CDASHIG. " + row["Notes"]).strip()
            else:
                row["Notes"] = ("No CDASH equivalent identified. Candidate for "
                                + (f"SUPP{fdom}.QVAL" if fdom else "a supplemental qualifier")
                                + " or a custom domain.")
            rows.append(row)
    return rows


def summarize(rows):
    by = {}
    for r in rows:
        k = f"{r['Confidence']} / {r['Basis']}"
        by[k] = by.get(k, 0) + 1
    mapped = sum(1 for r in rows if r["Confidence"] != "None")
    return {"fields": len(rows), "mapped": mapped, "by_confidence_basis": by}


def _meta(spec):
    sm = spec.get("study_meta") or {}
    cs = sm.get("cdisc_standards") or {}
    m = {}
    try:
        m = cdisc_ct.manifest()
    except Exception:
        pass
    return {"protocol": sm.get("protocol_number") or "", "title": sm.get("study_title") or sm.get("title") or "",
            "sponsor": sm.get("sponsor") or "", "ct": cs.get("ct_version") or m.get("ct_version", ""),
            "cdashig": cs.get("cdashig") or m.get("cdashig_version", ""), "sdtmig": m.get("sdtmig_version", ""),
            "generated": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}


_FILL = {"High": "D5F5E3", "Medium": "FDEBD0", "None": "EEEEEE"}
_WIDTH = {"Form": 12, "Field": 16, "Label": 38, "Type": 14, "Choices": 34, "CDASH Concept": 14, "Qualifier": 14,
          "SDTM Domain": 11, "SDTM Variable": 18, "Value-Level Detail": 24, "Controlled Terminology": 20,
          "Confidence": 11, "Basis": 20, "Notes": 60}
LEGEND = [("High", "Customer alias", "Deterministic: from the customer's own alias list."),
          ("High", "CDASH variable name", "Deterministic: the field is already named with a CDASH variable."),
          ("Medium", "QRS instrument (item order)", "Deterministic: instrument named on the form, items matched to "
                                                    "CDISC test codes by item order. Review."),
          ("Medium", "AI (validated)", "AI judgement, validated against CDASHIG and the field type. Review."),
          ("None", "Not mapped", "No CDASH equivalent identified. Candidate for a supplemental qualifier or custom domain.")]
DISCLAIMER = ("Proposed mapping for review. SDTM targets and notes come from CDASHIG v{cdashig} metadata (SDTMIG v{sdtmig}). "
              "Derivations such as combining date and time into --DTC or computing study days are described, not performed.")


def build_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows, meta = build_rows(spec), _meta(spec)
    summ = summarize(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = Font(bold=True, color="FFFFFF")
    band = PatternFill("solid", fgColor="1B3A6B")
    ws.append(["SDTM Mapping Specification"]); ws["A1"].font = Font(bold=True, size=14, color="1B3A6B")
    for k, v in [("Protocol", meta["protocol"]), ("Study", meta["title"]), ("Sponsor", meta["sponsor"]),
                 ("CDISC CT", meta["ct"]), ("CDASHIG", meta["cdashig"]), ("SDTMIG", meta["sdtmig"]),
                 ("Generated", meta["generated"]), ("Fields", summ["fields"]), ("Mapped to SDTM", summ["mapped"])]:
        ws.append([k, v]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["Confidence", "Basis", "Fields", "Meaning"])
    for c in ws[ws.max_row]:
        c.font, c.fill = head, band
    for conf, basis, meaning in LEGEND:
        ws.append([conf, basis, summ["by_confidence_basis"].get(f"{conf} / {basis}", 0), meaning])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=_FILL[conf])
    ws.append([])
    ws.append([DISCLAIMER.format(**meta)])
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=4)
    ws.row_dimensions[ws.max_row].height = 45
    for col, w in zip("ABCD", (18, 24, 10, 90)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Mapping")
    ws2.append(COLUMNS)
    for c in ws2[1]:
        c.font, c.fill = head, band
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws2.append([r[c] for c in COLUMNS])
        fill = PatternFill("solid", fgColor=_FILL[r["Confidence"]])
        for c in (COLUMNS.index("Confidence") + 1, COLUMNS.index("Basis") + 1):
            ws2.cell(ws2.max_row, c).fill = fill
        for c in ws2[ws2.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, name in enumerate(COLUMNS, 1):
        ws2.column_dimensions[get_column_letter(i)].width = _WIDTH[name]
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws2.max_row, 2)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(spec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from xml.sax.saxutils import escape
    rows, meta = build_rows(spec), _meta(spec)
    summ = summarize(rows)
    dark, light = colors.HexColor("#1B3A6B"), colors.HexColor("#D6E4F0")
    fills = {k: colors.HexColor("#" + v) for k, v in _FILL.items()}
    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=dark, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.5, textColor=dark, spaceBefore=8, spaceAfter=3)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=8, leading=10)
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=6.6, leading=8)
    hcell = ParagraphStyle("hc", fontName="Helvetica-Bold", fontSize=6.8, leading=8, textColor=colors.white)
    P = lambda t, s=cell: Paragraph(escape(str(t or "")), s)

    def footer(c, doc):
        c.saveState(); c.setFont("Helvetica", 7); c.setFillColor(colors.HexColor("#555555"))
        c.drawString(1.2 * cm, 0.7 * cm, f"{meta['protocol']}  |  SDTM Mapping Specification  |  {meta['generated']}")
        c.drawRightString(landscape(A4)[0] - 1.2 * cm, 0.7 * cm, f"Page {doc.page}")
        c.restoreState()

    story = [Paragraph("SDTM Mapping Specification", title),
             Paragraph(escape(f"{meta['protocol']}  {meta['title']}".strip()), body),
             Paragraph(escape(f"CDISC CT {meta['ct']}  |  CDASHIG {meta['cdashig']}  |  SDTMIG {meta['sdtmig']}  |  "
                              f"{summ['mapped']} of {summ['fields']} fields mapped"), body), Spacer(1, 6)]
    leg = [[P("Confidence", hcell), P("Basis", hcell), P("Fields", hcell), P("Meaning", hcell)]]
    for conf, basis, meaning in LEGEND:
        leg.append([P(conf), P(basis), P(summ["by_confidence_basis"].get(f"{conf} / {basis}", 0)), P(meaning)])
    lt = Table(leg, colWidths=[2.2 * cm, 3.6 * cm, 1.5 * cm, 17 * cm])
    lt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), dark), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP")] +
                           [("BACKGROUND", (0, i), (0, i), fills[LEGEND[i - 1][0]]) for i in range(1, len(leg))]))
    story += [lt, Spacer(1, 4), Paragraph(escape(DISCLAIMER.format(**meta)), body)]

    cols = ["Field", "Label", "CDASH Concept", "SDTM Domain", "SDTM Variable", "Value-Level Detail",
            "Controlled Terminology", "Confidence", "Basis", "Notes"]
    widths = [2.3, 4.4, 2.6, 1.6, 2.6, 3.0, 2.5, 1.9, 2.4, 4.0]
    by_form = {}
    for r in rows:
        by_form.setdefault(r["Form"], []).append(r)
    for form_id, frs in by_form.items():
        mapped = sum(1 for r in frs if r["Confidence"] != "None")
        data = [[P(c, hcell) for c in cols]]
        for r in frs:
            concept = r["CDASH Concept"] + (f" + {r['Qualifier']}" if r["Qualifier"] else "")
            data.append([P(r["Field"]), P(r["Label"]), P(concept), P(r["SDTM Domain"]), P(r["SDTM Variable"]),
                         P(r["Value-Level Detail"]), P(r["Controlled Terminology"]), P(r["Confidence"]),
                         P(r["Basis"]), P(r["Notes"])])
        t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), dark), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (6, -1), [colors.white, light])]
        for i, r in enumerate(frs, 1):
            style.append(("BACKGROUND", (7, i), (8, i), fills[r["Confidence"]]))
        t.setStyle(TableStyle(style))
        story += [KeepTogether([Paragraph(escape(f"Form {form_id}: {mapped} of {len(frs)} fields mapped"), h2)]), t]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                      topMargin=1.2 * cm, bottomMargin=1.3 * cm,
                      title=f"{meta['protocol']} SDTM Mapping Specification").build(story, onFirstPage=footer,
                                                                                  onLaterPages=footer)
    return buf.getvalue()


def build_files(spec):
    """{'xlsx': bytes, 'pdf': bytes, 'summary': dict}."""
    rows = build_rows(spec)
    return {"xlsx": build_xlsx(spec), "pdf": build_pdf(spec), "summary": summarize(rows)}

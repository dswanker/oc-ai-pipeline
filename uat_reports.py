"""
uat_reports.py — UAT Traceability Matrix (XLSX) and UAT Validation Report (PDF)

Both documents are built from the UAT results workbook alone (bytes in, bytes out), so they can be regenerated
from a stored workbook and never disagree with it. Nothing here calls OpenClinica, monday.com or an AI model.

The documents report counts. They do not say a study "passed" or is "validated": that is a decision for the
people who sign the report, and a data-import Pass only shows that a value was stored and read back unchanged.
"""
import html
import io
import re
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESULTS = ("Pass", "Fail", "Blocked", "Not Run", "Skip")
METHODS = ("Data import (ODM)", "Browser (Playwright)", "Manual")
NOT_IN_BUILD_STATUSES = ("proposed", "needs build team")   # DVS rows that were never built, so cannot be tested

CASE_COLUMNS = (
    "UAT Case ID", "Status", "Related Check ID", "Scenario", "Preconditions", "Test Steps", "Input Data",
    "Expected Result", "Actual Result", "Test Result", "Tester", "Execution Date", "Notes", "Site_OID",
    "Participant_Key", "Study_Event_OID", "Form_OID", "Item_Group_OID", "Item_OID", "Participant_ID",
    "Load_Value", "Item_Name", "Setup_Steps", "Test_Value", "Test Method", "Evidence",
)
CHECK_COLUMNS = (
    "Check ID", "Check Type", "Severity", "Target Form OID", "Target Item Name", "Plain-English Description",
    "Expression / Calculation", "Constraint / Required / Relevant Message", "Protocol Reference", "Status",
)
META_ROWS = (
    ("Protocol number", "protocol_number"), ("Study name", "study_name"), ("Study UUID", "study_uuid"),
    ("Study OID", "study_oid"), ("Environment", "environment"), ("Environment URL", "environment_url"),
    ("Site OID", "site_oid"), ("Executed by", "executed_by"), ("Executed at (UTC)", "executed_at"),
    ("Browser step", "browser_status"),
)
METHOD_NOTE = (
    "Data import (ODM): the value was sent through the OpenClinica data import and read back. A Pass means the "
    "value was stored and read back unchanged; the import does not run form logic. "
    "Browser (Playwright): the case was entered in the form in a browser. "
    "Manual: the case could not be run automatically in this run and needs a check by hand."
)
_FALLBACK_LIMITS = (
    "What UAT Cannot Confirm",
    "A Pass for a data-import case means only that the value was stored and read back unchanged. The import does "
    "not run form logic, so calculated values and values pulled from another form must be checked by hand.",
)
_RESULT_ALIASES = {"pass": "Pass", "passed": "Pass", "fail": "Fail", "failed": "Fail", "blocked": "Blocked",
                   "not run": "Not Run", "notrun": "Not Run", "skip": "Skip", "skipped": "Skip"}
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ── Reading the workbook ─────────────────────────────────────────────────────

def _s(val) -> str:
    """Cell value as clean text. Control characters go: the XLSX writer refuses them."""
    return "" if val is None else _CONTROL.sub("", str(val)).strip()


def _sheet_rows(wb, name, key, scan):
    """Rows of a sheet as dicts, using the first of `scan` rows that has a cell equal to `key` as the header."""
    if name not in wb.sheetnames:
        return []
    hdr, out = None, []
    for i, row in enumerate(wb[name].iter_rows(values_only=True)):
        if hdr is None:
            if i >= scan:
                return []
            names = [_s(c) for c in row]
            if key in names:
                hdr = names
            continue
        d = {h: _s(v) for h, v in zip(hdr, row) if h}
        if d.get(key):
            out.append(d)
    return out


def _norm_result(val: str) -> str:
    """Blank means the case was not run. An unknown value is kept as written rather than guessed at."""
    return _RESULT_ALIASES.get(val.lower(), val) if val else "Not Run"


def _form_oid(form: str) -> str:
    """DVS_OC4 names a form without the F_ prefix that UAT_Cases uses; one spelling so the two line up."""
    return form if not form or form.upper().startswith("F_") else f"F_{form}"


def _check_ids(case: dict) -> list:
    return [c for c in re.split(r"[,;\s]+", case.get("Related Check ID", "") or "") if c]


def read_results(results_bytes: bytes) -> dict:
    """{'cases': [dict per UAT case], 'checks': {check_id: dict}}; tolerant of missing sheets/columns.

    Every known column is present on every dict ("" when the workbook lacks it), so callers never KeyError
    on an older workbook. Test Result is normalised (blank = "Not Run").
    """
    try:
        wb = load_workbook(io.BytesIO(results_bytes), read_only=True, data_only=True)
    except Exception:
        return {"cases": [], "checks": {}}
    cases = []
    for d in _sheet_rows(wb, "UAT_Cases", "UAT Case ID", 6):
        case = {**dict.fromkeys(CASE_COLUMNS, ""), **d}
        case["Test Result"] = _norm_result(case["Test Result"])
        cases.append(case)
    checks = {}
    for d in _sheet_rows(wb, "DVS_OC4", "Check ID", 8):
        checks.setdefault(d["Check ID"], {**dict.fromkeys(CHECK_COLUMNS, ""), **d})
    wb.close()
    return {"cases": cases, "checks": checks}


# ── Counting ─────────────────────────────────────────────────────────────────

def summarize(cases: list) -> dict:
    """Counts used by both documents: total; by_result; by_method {method: {result: n}}; by_form {form_oid: {result: n}};
    executed (Pass+Fail); pass_rate (Pass / executed, None when executed == 0)."""
    by_result = dict.fromkeys(RESULTS, 0)
    by_method, by_form = {}, {}
    for c in cases:
        res = _norm_result(_s(c.get("Test Result")))
        by_result[res] = by_result.get(res, 0) + 1
        for bucket, key in ((by_method, _s(c.get("Test Method")) or "Not recorded"),
                            (by_form, _s(c.get("Form_OID")) or "(no form)")):
            d = bucket.setdefault(key, {})
            d[res] = d.get(res, 0) + 1
    executed = by_result["Pass"] + by_result["Fail"]
    return {"total": len(cases), "by_result": by_result, "by_method": by_method, "by_form": by_form,
            "executed": executed, "pass_rate": (by_result["Pass"] / executed) if executed else None}


def _coverage(data: dict) -> list:
    """One dict per check: built checks in DVS order, then checks only UAT_Cases knows, then checks not in the build."""
    counts, first_case = {}, {}
    for c in data["cases"]:
        for cid in _check_ids(c):
            counts.setdefault(cid, Counter())[c["Test Result"]] += 1
            first_case.setdefault(cid, c)
    rows, parked = [], []
    for cid in list(data["checks"]) + [c for c in counts if c not in data["checks"]]:
        chk, case, n = data["checks"].get(cid, {}), first_case.get(cid, {}), counts.get(cid, Counter())
        row = {"check_id": cid, "form": case.get("Form_OID") or _form_oid(chk.get("Target Form OID", "")),
               "item": chk.get("Target Item Name") or case.get("Item_Name", ""),
               "rule_type": chk.get("Check Type", ""), "cases": sum(n.values()), "counts": n,
               "in_build": chk.get("Status", "").lower() not in NOT_IN_BUILD_STATUSES, "outcome": ""}
        if not row["in_build"]:
            row["covered"] = "Not in build"
            parked.append(row)
            continue
        if n["Pass"] or n["Fail"]:
            row["covered"], row["outcome"] = "Yes", ("Fail" if n["Fail"] else "Pass")
        else:
            row["covered"] = "No — not executed" if row["cases"] else "No — no test case"
        rows.append(row)
    return rows + parked


def _limits() -> tuple:
    """(topic, text) of the standing statement on what a UAT result shows. Imported late so this module loads alone."""
    try:
        from uat_loader import UAT_LIMITS_TOPIC, UAT_LIMITS_TEXT
        return UAT_LIMITS_TOPIC, UAT_LIMITS_TEXT
    except Exception:
        return _FALLBACK_LIMITS


def _meta_rows(meta: dict) -> list:
    rows = [(label, _s(meta.get(key)) or "not recorded") for label, key in META_ROWS]
    people = meta.get("participants")
    people = "; ".join(f"{k} = {v}" for k, v in people.items()) if isinstance(people, dict) and people else ""
    return rows + [("Participants (logical ID = OpenClinica ID)", people or "not recorded")]


def _pct(n, total) -> str:
    return f"{100.0 * n / total:.1f}%" if total else "n/a"


def _environment(meta: dict) -> str:
    return " ".join(x for x in (_s(meta.get("environment")), _s(meta.get("environment_url"))) if x)


# ── Traceability matrix (XLSX) ───────────────────────────────────────────────

_HEAD_FILL = PatternFill("solid", fgColor="1B3A6B")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_RESULT_FILL = {"Pass": "D5F5E3", "Fail": "FADBD8", "Blocked": "FFF3CD"}
_OTHER_FILL = "EEEEEE"

_TRACE_COLS = (   # (header, width, wrap)
    ("Form", 14, False), ("Item Name", 18, False), ("Item OID", 22, False), ("Rule Type", 16, False),
    ("Rule Description", 44, True), ("Expression", 36, True), ("DVS Check ID", 13, False),
    ("Protocol Reference", 24, True), ("UAT Case ID", 12, False), ("Test Scenario", 32, True),
    ("Input Data", 22, True), ("Expected Result", 36, True), ("Actual Result", 36, True), ("Test Result", 11, False),
    ("Test Method", 20, False), ("Evidence", 44, True), ("Participant", 26, False), ("Study Event", 16, False),
    ("Executed By", 20, False), ("Execution Date", 20, False), ("Environment", 28, True),
)
_COVER_COLS = (
    ("DVS Check ID", 13), ("Form", 16), ("Item Name", 20), ("Rule Type", 18), ("Cases", 8), ("Pass", 8),
    ("Fail", 8), ("Blocked", 9), ("Not Run", 9), ("Skip", 8), ("Covered", 20), ("Outcome", 10),
)


def _write_row(ws, r, values, wraps=None):
    for c, val in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=val)
        if isinstance(val, str):
            cell.data_type = "s"   # study text that starts with "=" is text, never a formula
        cell.alignment = Alignment(vertical="top", wrap_text=bool(wraps and wraps[c - 1]))


def _sheet(wb, title, cols, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for c, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=col[0])
        cell.fill, cell.font = _HEAD_FILL, _HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = col[1]
    ws.freeze_panes = "A2"
    return ws


def _paint_result(cell):
    cell.fill = PatternFill("solid", fgColor=_RESULT_FILL.get(cell.value, _OTHER_FILL))


def build_traceability_matrix(results_bytes: bytes, meta: dict) -> bytes:
    """XLSX bytes: Traceability (one row per UAT case, joined to its DVS check), Coverage (one row per check)
    and Run Info (who/where/when, the counts, and what the run does not show)."""
    meta = meta or {}
    data = read_results(results_bytes)
    cases, checks = data["cases"], data["checks"]
    wb = Workbook()

    ws = _sheet(wb, "Traceability", _TRACE_COLS, first=True)
    wraps = [c[2] for c in _TRACE_COLS]
    result_col = [c[0] for c in _TRACE_COLS].index("Test Result") + 1
    env = _environment(meta)
    for r, c in enumerate(cases, 2):
        ids = _check_ids(c)
        chk = next((checks[i] for i in ids if i in checks), {})
        get = lambda k: chk.get(k, "")
        _write_row(ws, r, [
            c["Form_OID"] or _form_oid(get("Target Form OID")), c["Item_Name"] or get("Target Item Name"),
            c["Item_OID"], get("Check Type"),
            get("Plain-English Description") or get("Constraint / Required / Relevant Message")
            or get("Expression / Calculation"),
            get("Expression / Calculation"), c["Related Check ID"], get("Protocol Reference"), c["UAT Case ID"],
            c["Scenario"], c["Input Data"] or c["Load_Value"], c["Expected Result"], c["Actual Result"],
            c["Test Result"], c["Test Method"], c["Evidence"], c["Participant_Key"] or c["Participant_ID"],
            c["Study_Event_OID"], c["Tester"] or _s(meta.get("executed_by")),
            c["Execution Date"] or _s(meta.get("executed_at")), env], wraps)
        _paint_result(ws.cell(row=r, column=result_col))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_TRACE_COLS))}{max(len(cases) + 1, 1)}"

    cov = _coverage(data)
    ws = _sheet(wb, "Coverage", _COVER_COLS)
    for r, row in enumerate(cov, 2):
        n = row["counts"]
        _write_row(ws, r, [row["check_id"], row["form"], row["item"], row["rule_type"], row["cases"]]
                   + [n[res] for res in RESULTS] + [row["covered"], row["outcome"]])
        if row["outcome"]:
            _paint_result(ws.cell(row=r, column=len(_COVER_COLS)))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COVER_COLS))}{max(len(cov) + 1, 1)}"

    ws = _sheet(wb, "Run Info", (("Item", 44), ("Value", 100)))
    summ = summarize(cases)
    built = [x for x in cov if x["in_build"]]
    rows = [("Document", "UAT Traceability Matrix")] + _meta_rows(meta) + [("", ""), ("Total UAT cases", summ["total"])]
    rows += [(res, n) for res, n in summ["by_result"].items()]
    rows += [("Executed (Pass + Fail)", summ["executed"]),
             ("Pass rate (Pass / executed)", _pct(summ["by_result"]["Pass"], summ["executed"]))]
    for method in [m for m in METHODS if m in summ["by_method"]] + sorted(set(summ["by_method"]) - set(METHODS)):
        d = summ["by_method"][method]
        rows.append((f"Method: {method}", f"{sum(d.values())} cases (" + ", ".join(
            f"{k} {d[k]}" for k in list(RESULTS) + sorted(set(d) - set(RESULTS)) if d.get(k)) + ")"))
    rows += [("Checks in the build", len(built)),
             ("Checks covered (at least one Pass or Fail)", sum(1 for x in built if x["covered"] == "Yes")),
             ("Checks not covered", sum(1 for x in built if x["covered"] != "Yes")),
             ("Checks not in the build (Proposed / Needs Build Team)", len(cov) - len(built)),
             ("", ""), ("What this run does and does not show", ""), ("Test methods", METHOD_NOTE), _limits()]
    for r, row in enumerate(rows, 2):
        _write_row(ws, r, list(row), [True, True])
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
        if row[0] and row[1] == "":
            ws.cell(row=r, column=1).font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Validation report (PDF) ──────────────────────────────────────────────────

_PDF_CELL_MAX = 220      # a table row must fit on one page, so cell text is shortened; the matrix has it in full
_REASON_MAX = 110
_MAX_NOT_EXECUTED_ROWS = 60
_MAX_UNCOVERED_IDS = 80
_PLAIN = {"≥": ">=", "≤": "<=", "≠": "!=", "→": "->", "←": "<-", "…": "...", " ": " ", "​": ""}
# Only real HTML tag names are stripped, so "a <b" or "x < 5 and y > 3" in a rule survives as written.
_TAG_NAMES = (r"(?:span|div|p|br|b|i|u|em|strong|font|a|sup|sub|ul|ol|li|table|thead|tbody|tr|td|th|h[1-6]|small|"
              r"big|code|pre|img|hr|label|output|para|style|script)")
_TAG = re.compile(rf"</?{_TAG_NAMES}(?:\s[^<>]*)?/?>", re.I)
_CUT_TAG = re.compile(rf"</?{_TAG_NAMES}(?:\s[^<>]*)?$", re.I)   # a tag cut in half by upstream truncation


def _plain(text, limit=0) -> str:
    """Study text as plain text the standard PDF fonts can draw: tags and entities removed, one line, shortened."""
    text = _s(text)
    for k, v in _PLAIN.items():
        text = text.replace(k, v)
    text = html.unescape(_CUT_TAG.sub(" ", _TAG.sub(" ", text)))
    text = re.sub(r"\s+", " ", text).strip().encode("cp1252", "replace").decode("cp1252")
    return text[:limit - 3].rstrip() + "..." if limit and len(text) > limit else text


def _reason(case: dict) -> str:
    return _plain(case["Actual Result"] or case["Evidence"] or case["Notes"], _REASON_MAX) or "No reason recorded"


def build_validation_report(results_bytes: bytes, meta: dict) -> bytes:
    """PDF bytes (A4). Reports what was run and the counts; the sign-off block is left for people to complete."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    meta = meta or {}
    data = read_results(results_bytes)
    cases, checks = data["cases"], data["checks"]
    summ, cov = summarize(cases), _coverage(data)
    built = [x for x in cov if x["in_build"]]
    width = A4[0] - 30 * mm

    base = getSampleStyleSheet()
    st_title = ParagraphStyle("t", parent=base["Title"], fontSize=17, leading=21, alignment=0, spaceAfter=2)
    st_sub = ParagraphStyle("s", parent=base["Normal"], fontSize=10, leading=13, textColor=colors.HexColor("#444444"))
    st_h = ParagraphStyle("h", parent=base["Heading2"], fontSize=12, leading=15, spaceBefore=12, spaceAfter=5,
                          textColor=colors.HexColor("#1B3A6B"))
    st_body = ParagraphStyle("b", parent=base["Normal"], fontSize=9, leading=12, spaceAfter=4)
    st_cell = ParagraphStyle("c", parent=base["Normal"], fontSize=7, leading=8.5)
    st_head = ParagraphStyle("ch", parent=st_cell, fontName="Helvetica-Bold", textColor=colors.white)

    def para(text, style=st_body, limit=0):
        """Study text can hold anything; it is escaped, and if the library still rejects it the cell is left blank
        rather than losing the document."""
        try:
            return Paragraph(html.escape(_plain(text, limit), quote=False), style)
        except Exception:
            return Paragraph("", style)

    def table(header, rows, widths):
        """Table of wrapped cells that splits across pages and repeats its header."""
        total = float(sum(widths))
        body = [[para(h, st_head) for h in header]]
        body += [[para(v, st_cell, _PDF_CELL_MAX) for v in row] for row in rows]
        t = Table(body, colWidths=[width * w / total for w in widths], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B3A6B")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        return t

    def form_item(c):
        return ".".join(x for x in (c["Form_OID"], c["Item_Name"] or c["Item_OID"]) if x)

    extra = sorted(set(summ["by_result"]) - set(RESULTS))          # results the workbook spells some other way
    result_cols = list(RESULTS) + extra
    count_widths = [12] * (len(result_cols) + 1)
    title = "UAT Validation Report"
    story = [Paragraph(title, st_title),
             para(" — ".join(x for x in (_s(meta.get("protocol_number")), _s(meta.get("study_name"))) if x)
                  or "Study not recorded", st_sub), Spacer(1, 6)]

    story.append(Paragraph("1. Study and run information", st_h))
    story.append(table(("Item", "Value"), _meta_rows(meta), (30, 70)))

    story.append(Paragraph("2. Scope", st_h))
    forms = sorted({c["Form_OID"] for c in cases if c["Form_OID"]} | {x["form"] for x in built if x["form"]})
    story.append(para(f"Forms in scope ({len(forms)}): " + (", ".join(forms) or "none recorded") + "."))
    if checks:
        story.append(para(f"Checks in the DVS: {len(cov)}, of which {len(built)} are in the build and "
                          f"{len(cov) - len(built)} are not (status Proposed or Needs Build Team)."))
    else:
        story.append(para(f"The workbook has no DVS check sheet. The {len(built)} checks counted here are the "
                          "check IDs named by the UAT cases."))
    story.append(para(f"UAT cases: {summ['total']}, of which {summ['executed']} were executed (Pass or Fail)."))

    story.append(Paragraph("3. Results summary", st_h))
    story.append(table(("Result", "Cases", "% of all cases"),
                       [(res, summ["by_result"][res], _pct(summ["by_result"][res], summ["total"]))
                        for res in result_cols] + [("Total", summ["total"], "")], (40, 20, 20)))
    story.append(Spacer(1, 4))
    story.append(para(f"Executed cases (Pass + Fail): {summ['executed']}. Pass rate: "
                      f"{_pct(summ['by_result']['Pass'], summ['executed'])} of executed cases "
                      f"({summ['by_result']['Pass']} of {summ['executed']}). Cases that were Blocked, Not Run or "
                      "skipped are not part of the pass rate."))
    story.append(para("By test method"))
    methods = [m for m in METHODS if m in summ["by_method"]] + sorted(set(summ["by_method"]) - set(METHODS))
    story.append(table(["Test method", "Cases"] + result_cols,
                       [[m, sum(summ["by_method"][m].values())] + [summ["by_method"][m].get(r, 0) for r in result_cols]
                        for m in methods] or [["No cases"] + [""] * (len(result_cols) + 1)], [34] + count_widths))
    story.append(Spacer(1, 4))
    story.append(para(METHOD_NOTE))
    story.append(para("By form"))
    story.append(table(["Form", "Cases"] + result_cols,
                       [[f, sum(d.values())] + [d.get(r, 0) for r in result_cols]
                        for f, d in sorted(summ["by_form"].items())]
                       or [["No cases"] + [""] * (len(result_cols) + 1)], [34] + count_widths))

    story.append(Paragraph("4. Failed cases", st_h))
    failed = [c for c in cases if c["Test Result"] == "Fail"]
    if failed:
        story.append(para(f"{len(failed)} cases have the result Fail. Long text is shortened in this report; "
                          "the traceability matrix has it in full."))
        story.append(table(("UAT Case ID", "Check ID", "Form.Item", "Scenario", "Expected", "Actual", "Method"),
                           [(c["UAT Case ID"], c["Related Check ID"], form_item(c), c["Scenario"],
                             c["Expected Result"], c["Actual Result"], c["Test Method"]) for c in failed],
                           (17, 16, 30, 30, 34, 34, 19)))
    else:
        story.append(para("None."))

    story.append(Paragraph("5. Blocked, Not Run and skipped cases", st_h))
    idle = [c for c in cases if c["Test Result"] not in ("Pass", "Fail")]
    if idle:
        story.append(para(f"{len(idle)} cases were not executed. Counts by recorded reason:"))
        reasons = Counter((c["Test Result"], _reason(c)) for c in idle)
        story.append(table(("Result", "Reason", "Cases"), [(res, why, n) for (res, why), n in reasons.most_common()],
                           (16, 90, 12)))
        story.append(Spacer(1, 6))
        more = len(idle) - _MAX_NOT_EXECUTED_ROWS
        story.append(para("Individual cases" + (f" (first {_MAX_NOT_EXECUTED_ROWS} of {len(idle)}; the other {more} "
                                                "are in the traceability matrix):" if more > 0 else ":")))
        story.append(table(("UAT Case ID", "Check ID", "Form.Item", "Result", "Reason", "Method"),
                           [(c["UAT Case ID"], c["Related Check ID"], form_item(c), c["Test Result"], _reason(c),
                             c["Test Method"]) for c in idle[:_MAX_NOT_EXECUTED_ROWS]], (17, 16, 32, 14, 70, 21)))
    else:
        story.append(para("None."))

    story.append(Paragraph("6. Coverage", st_h))
    covered = [x for x in built if x["covered"] == "Yes"]
    uncovered = [x for x in built if x["covered"] != "Yes"]
    story.append(para(f"A check counts as covered when at least one of its cases has the result Pass or Fail. "
                      f"Checks in the build: {len(built)}. Covered: {len(covered)} ({_pct(len(covered), len(built))}). "
                      f"Not covered: {len(uncovered)}, of which "
                      f"{sum(1 for x in uncovered if x['cases'])} have cases that were not executed and "
                      f"{sum(1 for x in uncovered if not x['cases'])} have no test case. "
                      f"Of the covered checks, {sum(1 for x in covered if x['outcome'] == 'Fail')} have at least one "
                      f"failed case. Checks not in the build: {len(cov) - len(built)}."))
    if uncovered:
        shown = ", ".join(x["check_id"] for x in uncovered[:_MAX_UNCOVERED_IDS])
        more = len(uncovered) - _MAX_UNCOVERED_IDS
        story.append(para("Checks not covered: " + shown + (f", and {more} more (see the Coverage sheet of the "
                                                            "traceability matrix)." if more > 0 else ".")))

    topic, text = _limits()
    story.append(Paragraph("7. Limits of this run", st_h))
    story.append(para(topic + ". " + text))

    sign = Table([[para(h, st_head) for h in ("Role", "Name", "Signature", "Date")]]
                 + [[para(role, st_body), "", "", ""] for role in ("Prepared by", "Reviewed by", "Approved by")],
                 colWidths=[width * w for w in (0.18, 0.30, 0.34, 0.18)], rowHeights=[None, 30, 30, 30])
    sign.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B3A6B")),
                              ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
                              ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.append(KeepTogether([Paragraph("8. Sign-off", st_h),
                               para("Signing records review of the counts above. This report does not itself "
                                    "state an outcome for the study."), sign]))

    footer = _plain(" — ".join(x for x in (title, _s(meta.get("protocol_number")), _s(meta.get("executed_at"))) if x), 110)

    def draw_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(15 * mm, 9 * mm, footer)
        canvas.drawRightString(A4[0] - 15 * mm, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=16 * mm, title=title, author=_plain(meta.get("executed_by"), 80))
    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return buf.getvalue()
